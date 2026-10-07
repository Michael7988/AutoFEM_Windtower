from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any


@dataclass
class SheetData:
    name: str
    rows: list[list[Any]]

    @property
    def nrows(self) -> int:
        return len(self.rows)

    @property
    def ncols(self) -> int:
        return max((len(row) for row in self.rows), default=0)

    def cell(self, row: int, col: int) -> Any:
        if row < 0 or col < 0 or row >= len(self.rows):
            return ""
        values = self.rows[row]
        if col >= len(values):
            return ""
        return "" if values[col] is None else values[col]


@dataclass
class WorkbookData:
    path: Path
    sheets: dict[str, SheetData]

    @property
    def sheet_names(self) -> list[str]:
        return list(self.sheets)

    def sheet_by_name(self, name: str) -> SheetData:
        return self.sheets[name]


def read_workbook_data(path: str | Path) -> WorkbookData:
    workbook_path = Path(path)
    suffix = workbook_path.suffix.lower()
    if suffix == ".xls":
        return _read_xls(workbook_path)
    if suffix in {".xlsx", ".xlsm"}:
        return _read_xlsx(workbook_path)
    raise ValueError(f"Unsupported workbook format: {workbook_path.suffix}")


def _trim_trailing_empty(rows: list[list[Any]]) -> list[list[Any]]:
    trimmed_rows: list[list[Any]] = []
    for row in rows:
        values = list(row)
        while values and _is_blank(values[-1]):
            values.pop()
        trimmed_rows.append(values)
    while trimmed_rows and not any(not _is_blank(value) for value in trimmed_rows[-1]):
        trimmed_rows.pop()
    return trimmed_rows


def _read_xls(path: Path) -> WorkbookData:
    import xlrd

    book = xlrd.open_workbook(str(path))
    sheets: dict[str, SheetData] = {}
    for name in book.sheet_names():
        sheet = book.sheet_by_name(name)
        rows = [[sheet.cell_value(row, col) for col in range(sheet.ncols)] for row in range(sheet.nrows)]
        sheets[name] = SheetData(name=name, rows=_trim_trailing_empty(rows))
    return WorkbookData(path=path, sheets=sheets)


def _read_xlsx(path: Path) -> WorkbookData:
    import openpyxl

    book = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheets: dict[str, SheetData] = {}
    for name in book.sheetnames:
        sheet = book[name]
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        sheets[name] = SheetData(name=name, rows=_trim_trailing_empty(rows))
    return WorkbookData(path=path, sheets=sheets)


def _is_blank(value: Any) -> bool:
    return value is None or str(value).strip() == ""

