from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import uuid
from pathlib import Path
from typing import Any, Callable

from ..capability_profile import EXPLORATORY_INPUT_EXTENSIONS, FORMAL_INPUT_EXTENSIONS
from .models import EvidenceFragment, EvidenceLocation


TABLE_SUFFIXES = {".xls", ".xlsx", ".xlsm", ".csv", ".tsv"}
TEXT_SUFFIXES = {".json", ".yaml", ".yml", ".txt", ".md", ".inp"}
DOCUMENT_SUFFIXES = {".pdf", ".docx", ".pptx"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}
SUPPORTED_SUFFIXES = TABLE_SUFFIXES | TEXT_SUFFIXES | DOCUMENT_SUFFIXES | IMAGE_SUFFIXES
PUBLIC_ENGINEERING_SUFFIXES = set(FORMAL_INPUT_EXTENSIONS) | set(EXPLORATORY_INPUT_EXTENSIONS)
IMAGE_SIZE_LIMIT = 20 * 1024 * 1024
FILE_SIZE_LIMIT = 50 * 1024 * 1024
BATCH_SIZE_LIMIT = 200 * 1024 * 1024
MAX_FILES = 20
EXTRACTOR_VERSION = "engineering_extractors_v2"


def classify_file(filename: str, media_type: str = "") -> tuple[str, str]:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(f"暂不支持文件格式：{suffix or '无扩展名'}")
    if suffix in {".xls", ".xlsx", ".xlsm", ".csv", ".tsv", ".json", ".yaml", ".yml", ".txt", ".md", ".inp"}:
        return "直接读取", "no_ocr"
    if suffix in {".docx", ".pptx"}:
        return "文字提取", "no_ocr"
    if suffix == ".pdf":
        return "逐页检测", "partial_ocr"
    return ("图纸识别", "complex_drawing") if "drawing" in filename.lower() or "图" in filename else ("需要OCR", "full_ocr")


def validate_upload(filename: str, content: bytes) -> None:
    suffix = Path(filename).suffix.lower()
    if not content:
        raise ValueError(f"{filename} 是空文件")
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(f"{filename} 的格式不受支持")
    limit = IMAGE_SIZE_LIMIT if suffix in IMAGE_SUFFIXES else FILE_SIZE_LIMIT
    if len(content) > limit:
        raise ValueError(f"{filename} 超过 {limit // 1024 // 1024} MB 限制")
    if content[:2] == b"MZ":
        raise ValueError(f"{filename} 疑似可执行文件")
    if suffix == ".pdf" and not content.startswith(b"%PDF"):
        raise ValueError(f"{filename} 的内容不是有效PDF")
    if suffix in {".png"} and not content.startswith(b"\x89PNG"):
        raise ValueError(f"{filename} 的内容与扩展名不一致")
    if suffix in {".jpg", ".jpeg"} and not content.startswith(b"\xff\xd8"):
        raise ValueError(f"{filename} 的内容与扩展名不一致")


def validate_engineering_upload(filename: str, content: bytes) -> None:
    validate_upload(filename, content)
    suffix = Path(filename).suffix.lower()
    if suffix not in PUBLIC_ENGINEERING_SUFFIXES:
        raise ValueError(f"{filename} 暂不纳入当前工程文档输入范围。")
    if suffix != ".pdf":
        return
    try:
        import fitz

        document = fitz.open(stream=content, filetype="pdf")
        has_text = any(page.get_text("text").strip() for page in document)
    except Exception as exc:
        raise ValueError(f"{filename} 无法读取文字层。") from exc
    # A scan is accepted as exploratory input. The processing stage must mark
    # OCR/VLM evidence as needs_review and may not promote it automatically.
    if not has_text:
        return


def extract_file(
    file_id: str,
    filename: str,
    content: bytes,
    ocr: Callable[[bytes, str, str], dict[str, Any]] | None = None,
) -> tuple[list[EvidenceFragment], dict[str, Any]]:
    suffix = Path(filename).suffix.lower()
    if suffix in {".xls", ".xlsx", ".xlsm"}:
        return [], {"method": "existing_excel_normalizer", "ocr_pages": 0}
    if suffix in {".csv", ".tsv"}:
        return _extract_delimited(file_id, filename, content, "\t" if suffix == ".tsv" else ","), {"method": "local_table", "ocr_pages": 0}
    if suffix == ".json":
        return _extract_json(file_id, filename, content), {"method": "local_json", "ocr_pages": 0}
    if suffix in {".yaml", ".yml"}:
        return _extract_yaml(file_id, filename, content), {"method": "local_yaml", "ocr_pages": 0}
    if suffix in {".txt", ".md", ".inp"}:
        return _extract_text(file_id, filename, content), {"method": "local_text", "ocr_pages": 0}
    if suffix == ".docx":
        return _extract_docx(file_id, filename, content), {"method": "local_docx", "ocr_pages": 0}
    if suffix == ".pptx":
        return _extract_pptx(file_id, filename, content), {"method": "local_pptx", "ocr_pages": 0}
    if suffix == ".pdf":
        return _extract_pdf(file_id, filename, content, ocr)
    return _extract_image(file_id, filename, content, ocr)


def _fragment(file_id: str, filename: str, kind: str, text: str, method: str, location: EvidenceLocation, **extra: Any) -> EvidenceFragment:
    digest = hashlib.sha256(f"{file_id}|{location.model_dump_json()}|{text}".encode("utf-8")).hexdigest()
    return EvidenceFragment(
        evidence_id=f"ev-{uuid.uuid4().hex[:12]}",
        file_id=file_id,
        filename=filename,
        kind=kind,
        text=text[:12000],
        location=location,
        extraction_method=method,
        content_hash=digest,
        **extra,
    )


def _decode(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "utf-16"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def _extract_delimited(file_id: str, filename: str, content: bytes, delimiter: str) -> list[EvidenceFragment]:
    rows = list(csv.reader(io.StringIO(_decode(content)), delimiter=delimiter))
    fragments = []
    for start in range(0, min(len(rows), 3000), 50):
        block = rows[start:start + 50]
        text = "\n".join(delimiter.join(str(cell) for cell in row) for row in block)
        fragments.append(_fragment(file_id, filename, "table", text, "local_delimited", EvidenceLocation(line=start + 1)))
    return fragments


def _extract_json(file_id: str, filename: str, content: bytes) -> list[EvidenceFragment]:
    data = json.loads(_decode(content))
    fragments = []
    for path, value in _flatten(data):
        fragments.append(_fragment(file_id, filename, "structured", json.dumps(value, ensure_ascii=False), "local_json", EvidenceLocation(json_path=path)))
    return fragments[:3000]


def _extract_yaml(file_id: str, filename: str, content: bytes) -> list[EvidenceFragment]:
    try:
        import yaml
    except ImportError as exc:
        raise RuntimeError("缺少 PyYAML，无法读取 YAML") from exc
    data = yaml.safe_load(_decode(content))
    fragments = []
    for path, value in _flatten(data):
        fragments.append(_fragment(file_id, filename, "structured", json.dumps(value, ensure_ascii=False), "local_yaml", EvidenceLocation(json_path=path)))
    return fragments[:3000]


def _flatten(value: Any, path: str = "$") -> list[tuple[str, Any]]:
    if isinstance(value, dict):
        rows = []
        for key, child in value.items():
            rows.extend(_flatten(child, f"{path}.{key}"))
        return rows
    if isinstance(value, list):
        if value and all(not isinstance(item, (dict, list)) for item in value):
            return [(path, value)]
        rows = []
        for index, child in enumerate(value):
            rows.extend(_flatten(child, f"{path}[{index}]"))
        return rows
    return [(path, value)]


def _extract_text(file_id: str, filename: str, content: bytes) -> list[EvidenceFragment]:
    text = _decode(content).replace("\x00", "")
    blocks = [item.strip() for item in re.split(r"\n\s*\n", text) if item.strip()]
    fragments = []
    line = 1
    for block in blocks:
        for offset in range(0, len(block), 6000):
            chunk = block[offset:offset + 6000]
            fragments.append(_fragment(file_id, filename, "text", chunk, "local_text", EvidenceLocation(line=line)))
        line += block.count("\n") + 2
    return fragments[:1000]


def _extract_docx(file_id: str, filename: str, content: bytes) -> list[EvidenceFragment]:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("缺少 python-docx，无法读取DOCX") from exc
    document = Document(io.BytesIO(content))
    fragments = []
    for index, paragraph in enumerate(document.paragraphs, start=1):
        if paragraph.text.strip():
            fragments.append(_fragment(file_id, filename, "text", paragraph.text.strip(), "local_docx", EvidenceLocation(line=index)))
    for table_index, table in enumerate(document.tables, start=1):
        text = "\n".join("\t".join(cell.text for cell in row.cells) for row in table.rows)
        fragments.append(_fragment(file_id, filename, "table", text, "local_docx_table", EvidenceLocation(json_path=f"$.tables[{table_index - 1}]")))
    return fragments


def _extract_pptx(file_id: str, filename: str, content: bytes) -> list[EvidenceFragment]:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise RuntimeError("缺少 python-pptx，无法读取PPTX") from exc
    presentation = Presentation(io.BytesIO(content))
    fragments = []
    for slide_index, slide in enumerate(presentation.slides, start=1):
        text = "\n".join(str(shape.text) for shape in slide.shapes if hasattr(shape, "text") and str(shape.text).strip())
        if text:
            fragments.append(_fragment(file_id, filename, "text", text, "local_pptx", EvidenceLocation(slide=slide_index)))
    return fragments


def _extract_pdf(file_id: str, filename: str, content: bytes, ocr: Callable[[bytes, str, str], dict[str, Any]] | None) -> tuple[list[EvidenceFragment], dict[str, Any]]:
    try:
        import fitz
    except ImportError as exc:
        raise RuntimeError("缺少 PyMuPDF，无法读取PDF") from exc
    document = fitz.open(stream=content, filetype="pdf")
    fragments = []
    ocr_pages = 0
    for page_index, page in enumerate(document, start=1):
        text = page.get_text("text").strip()
        if len(re.sub(r"\s+", "", text)) >= 40:
            fragments.append(_fragment(file_id, filename, "text", text, "local_pdf_text", EvidenceLocation(page=page_index)))
            continue
        ocr_pages += 1
        pixmap = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
        image = pixmap.tobytes("png")
        result = ocr(image, "image/png", "提取该工程资料页面中的文字、表格、尺寸和单位。") if ocr else {"status": "not_configured", "text": ""}
        recognized = str(result.get("text") or "").strip()
        fragments.append(_fragment(
            file_id,
            filename,
            "drawing" if not recognized else "text",
            recognized or "[该页需要OCR确认]",
            str(result.get("model") or "ocr_unavailable"),
            EvidenceLocation(page=page_index),
            confidence=float(result.get("confidence") or (0.75 if recognized else 0.0)),
            requires_ocr=True,
            from_ocr=bool(recognized),
            status="needs_review",
        ))
    return fragments, {"method": "hybrid_pdf", "ocr_pages": ocr_pages}


def _extract_image(file_id: str, filename: str, content: bytes, ocr: Callable[[bytes, str, str], dict[str, Any]] | None) -> tuple[list[EvidenceFragment], dict[str, Any]]:
    suffix = Path(filename).suffix.lower()
    media_type = "image/jpeg" if suffix in {".jpg", ".jpeg"} else f"image/{suffix.lstrip('.')}"
    result = ocr(content, media_type, "提取工程图片中的文字、尺寸、材料、图签和参数表；不推断完整几何。") if ocr else {"status": "not_configured", "text": ""}
    text = str(result.get("text") or "").strip()
    fragment = _fragment(
        file_id,
        filename,
        "drawing",
        text or "[图片需要视觉识别确认]",
        str(result.get("model") or "vision_unavailable"),
        EvidenceLocation(),
        confidence=float(result.get("confidence") or (0.7 if text else 0.0)),
        requires_ocr=True,
        from_ocr=bool(text),
        status="needs_review",
    )
    return [fragment], {"method": "vision", "ocr_pages": 1}
