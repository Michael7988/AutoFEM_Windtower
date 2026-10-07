"""Wind-tower project parsing and deterministic model implementation."""

from __future__ import annotations

import hashlib
import json
import re
import shutil
import subprocess
import unicodedata
import uuid
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from ..capability_profile import CANONICAL_DEFAULTS, L2_CAPABILITY_ID, get_capability
from ..abaqus_adapter import (
    element_center,
    find_flange_ranges,
    segments_from_element_centers,
    static_check,
    windward_surface_groups,
    write_generated_analysis_inp,
)
from ..config import STORAGE_ROOT, Settings
from ..facts import build_engineering_fact_package, write_fact_artifacts
from ..importers.windtower_adapter import parse_windtower_workbook
from ..meshing import build_windtower_quad_mesh
from ..meshing.abaqus_writer import write_mesh_inp
from ..meshing.models import normalize_tower_points
from ..output_contract import normalize_output_contract
from ..parameter_ir import apply_parameter_set_to_ir, l2_parameter_set, parameter_set_from_ir
from .repository import ArtifactRepository


@dataclass
class ProjectState:
    project_id: str
    root: Path
    excel_path: Path
    ir: dict[str, Any]
    artifacts: dict[str, Path]


PROJECTS: dict[str, ProjectState] = {}
ARTIFACT_REPOSITORY = ArtifactRepository()


def create_project(upload_name: str, content: bytes, modeling_goal: str | None = None, imported: dict[str, Any] | None = None) -> dict[str, Any]:
    project_id = f"wt-{uuid.uuid4().hex[:10]}"
    root = STORAGE_ROOT / project_id
    root.mkdir(parents=True, exist_ok=True)
    excel_path = root / sanitize_filename(upload_name)
    excel_path.write_bytes(content)
    imported = imported or parse_windtower_workbook(excel_path)
    imported.setdefault("source_file_original", upload_name)
    imported.setdefault("source_file_storage", excel_path.name)
    if imported.get("workbook_profile"):
        imported["workbook_profile"].setdefault("source_file_original", upload_name)
        imported["workbook_profile"].setdefault("source_file_storage", excel_path.name)
    fact_package = build_engineering_fact_package(imported, excel_path)
    imported["engineering_fact_package"] = fact_package.model_dump()
    ir = build_ir(
        project_id,
        upload_name,
        imported["tower"],
        imported["flanges"],
        imported["door"],
        imported["fem_params"],
        import_meta=imported,
        modeling_goal=modeling_goal,
    )
    parameter_set = parameter_set_from_ir(project_id, ir)
    ir = apply_parameter_set_to_ir(parameter_set, ir)
    state = ProjectState(project_id=project_id, root=root, excel_path=excel_path, ir=ir, artifacts={})
    PROJECTS[project_id] = state
    write_json(root / "ir.json", ir)
    state.artifacts["ir.json"] = root / "ir.json"
    for name, path in write_fact_artifacts(fact_package, root / "facts").items():
        state.artifacts[f"facts/{name}"] = path
    write_standard_input_artifacts(state, imported)
    parameter_path = root / "parameter_set.json"
    write_json(parameter_path, parameter_set.model_dump(mode="json"))
    state.artifacts["parameter_set.json"] = parameter_path
    standard_path = export_standard_excel(state)
    state.artifacts[standard_path.name] = standard_path
    ARTIFACT_REPOSITORY.replace(project_id, root, state.artifacts)
    return project_response(state)


def create_l2_project(
    modeling_goal: str,
    values: dict[str, Any] | None = None,
    *,
    source_name: str = "l2-engineering-intent.json",
    operator: str = "engineer",
) -> dict[str, Any]:
    """Create an isolated L2 project without coercing it through the L1 workbook IR."""
    project_id = f"wt-{uuid.uuid4().hex[:10]}"
    root = STORAGE_ROOT / project_id
    root.mkdir(parents=True, exist_ok=True)
    source_path = root / sanitize_filename(source_name)
    source_path.write_text(
        json.dumps({"modeling_goal": modeling_goal, "parameters": values or {}}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    parameter_set = l2_parameter_set(project_id, values, confirmed=True)
    parameter_set.confirmed_by = operator
    capability = get_capability(L2_CAPABILITY_ID)
    ir: dict[str, Any] = {
        "schema_version": "windtower_parameter_set_v2",
        "project_meta_ir": {
            "project_id": project_id,
            "source_files": [source_name],
            "modeling_goal": modeling_goal,
        },
        "raw_excel_ir": {
            "source_file_original": source_name,
            "source_file_storage": source_path.name,
            "parse_warnings": [],
            "input_format": "semantic_command_v1",
        },
        "semantic_ir": {
            "tower_sections": [], "flanges": [], "materials": [],
            "door_openings": [{"enabled": False, "shape": "none", "width_mm": 0, "height_mm": 0, "bottom_z_mm": 0}],
            "loads_inventory": {},
        },
        "modeling_capability": capability,
        "modeling_tasks": [],
        "modeling_intent_contract": {"goal": modeling_goal, "capability_id": L2_CAPABILITY_ID},
    }
    ir = apply_parameter_set_to_ir(parameter_set, ir)
    state = ProjectState(project_id=project_id, root=root, excel_path=source_path, ir=ir, artifacts={})
    PROJECTS[project_id] = state
    write_json(root / "ir.json", ir)
    write_json(root / "parameter_set.json", parameter_set.model_dump(mode="json"))
    state.artifacts.update({
        "ir.json": root / "ir.json",
        "parameter_set.json": root / "parameter_set.json",
        f"input/{source_path.name}": source_path,
    })
    ARTIFACT_REPOSITORY.replace(project_id, root, state.artifacts)
    return project_response(state)


def project_response(state: ProjectState) -> dict[str, Any]:
    sections = state.ir["semantic_ir"]["tower_sections"]
    capability_id = str((state.ir.get("modeling_capability") or {}).get("capability_id") or "")
    return {
        "project_id": state.project_id,
        "stage": "parsed",
        "summary": {
            "tower_height_m": sections[-1]["z_top_mm"] / 1000.0 if sections else 0,
            "section_count": len(sections),
            "flange_count": len(state.ir["semantic_ir"]["flanges"]),
            "material_count": len(state.ir["semantic_ir"]["materials"]),
            "door": door_summary(state.ir),
            "capability_id": capability_id,
        },
        "ir": state.ir,
        "warnings": state.ir["raw_excel_ir"]["parse_warnings"],
        "modeling_tasks": state.ir["modeling_tasks"],
        "input_format": state.ir["raw_excel_ir"].get("input_format") or state.ir["raw_excel_ir"].get("detected_format"),
        "source_file_original": state.ir["raw_excel_ir"].get("source_file_original"),
        "source_file_storage": state.ir["raw_excel_ir"].get("source_file_storage"),
        "format_confidence": state.ir["raw_excel_ir"].get("format_confidence"),
        "parse_strategy": state.ir["raw_excel_ir"].get("parse_strategy"),
        "llm_required": state.ir["raw_excel_ir"].get("llm_required"),
        "llm_status": (state.ir["raw_excel_ir"].get("llm_parse_assist") or {}).get("status"),
        "promotion_status": state.ir["raw_excel_ir"].get("promotion_status"),
        "standard_input_candidate": state.ir["raw_excel_ir"].get("standard_input_candidate"),
        "validation_report": state.ir["raw_excel_ir"].get("standard_input_validation"),
        "input_quality_scores": state.ir["raw_excel_ir"].get("input_quality_scores"),
        "requires_user_confirmation": state.ir["raw_excel_ir"].get("requires_user_confirmation", False),
        "llm_usage": state.ir["raw_excel_ir"].get("llm_parse_assist"),
        "llm_standard_input_candidate": state.ir["raw_excel_ir"].get("llm_standard_input_candidate"),
        "llm_standard_input_validation": state.ir["raw_excel_ir"].get("llm_standard_input_validation"),
        "llm_promotion_status": state.ir["raw_excel_ir"].get("llm_promotion_status"),
        "mapping_evidence": state.ir["raw_excel_ir"].get("parse_diagnostics"),
        "candidate_mapping_ir": state.ir["raw_excel_ir"].get("candidate_mapping_ir"),
        "mesh_automation_gate": state.ir.get("mesh_automation_gate"),
        "mesh_handoff_contract": state.ir.get("mesh_handoff_contract"),
        "load_contract": state.ir.get("load_contract"),
        "load_readiness": state.ir.get("load_readiness"),
        "analysis_readiness": state.ir.get("analysis_readiness"),
        "modal_policy": state.ir.get("modal_policy"),
        "engineering_review": state.ir.get("engineering_review"),
        "unsupported_candidates": state.ir.get("unsupported_candidates", []),
        "parameter_set": state.ir.get("parameter_set"),
    }


def write_standard_input_artifacts(state: ProjectState, imported: dict[str, Any]) -> None:
    input_dir = state.root / "input"
    input_dir.mkdir(parents=True, exist_ok=True)
    artifact_map = {
        "workbook_profile.json": imported.get("workbook_profile"),
        "standard_input_candidate.json": imported.get("standard_input_candidate") or imported.get("standard_input_package"),
        "confirmed_standard_input.json": imported.get("standard_input_package"),
        "standard_input_validation.json": imported.get("validation_report"),
        "input_quality_scores.json": imported.get("input_quality_scores"),
        "format_alias_registry_update.json": imported.get("format_alias_registry_update"),
        "llm_parse_audit.json": imported.get("llm_parse_assist") or imported.get("llm_suggestions"),
        "llm_standard_input_candidate.json": imported.get("llm_standard_input_candidate"),
        "llm_standard_input_validation.json": imported.get("llm_standard_input_validation"),
        "candidate_mapping_ir.json": imported.get("candidate_mapping_ir"),
    }
    for name, payload in artifact_map.items():
        if payload is None:
            continue
        path = input_dir / name
        write_json(path, payload)
        state.artifacts[f"input/{name}"] = path


def export_standard_excel(state: ProjectState) -> Path:
    import openpyxl

    path = state.root / "standardized_model.xlsx"
    ir = state.ir
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "塔节参数"
    ws.append(["section_id", "z_bottom_mm", "z_top_mm", "diameter_bottom_mm", "diameter_top_mm", "thickness_mm", "material", "source"])
    for section in ir["semantic_ir"]["tower_sections"]:
        ws.append([
            section.get("section_id"),
            section.get("z_bottom_mm"),
            section.get("z_top_mm"),
            section.get("diameter_bottom_mm"),
            section.get("diameter_top_mm"),
            section.get("thickness_mm"),
            section.get("material"),
            section.get("source"),
        ])

    ws = wb.create_sheet("法兰参数")
    ws.append(["flange_id", "elevation_mm", "mean_diameter_mm", "bolt_circle_mm", "inner_diameter_mm", "total_thickness_mm", "bolt_count", "bolt_diameter_mm", "source"])
    for flange in ir["semantic_ir"]["flanges"]:
        ws.append([
            flange.get("flange_id"),
            flange.get("elevation_mm"),
            flange.get("mean_diameter_mm"),
            flange.get("bolt_circle_mm"),
            flange.get("inner_diameter_mm"),
            flange.get("total_thickness_mm"),
            flange.get("bolt_count"),
            flange.get("bolt_diameter_mm"),
            flange.get("source"),
        ])

    ws = wb.create_sheet("门洞参数")
    ws.append(["key", "value", "unit", "source", "description"])
    door = ir["semantic_ir"]["door_openings"][0]
    ws.append(["enabled", door.get("enabled"), "", door.get("source"), "Whether complete door geometry was detected."])
    ws.append(["shape", door.get("shape"), "", door.get("source"), "Door opening shape."])
    ws.append(["bottom_z_mm", door.get("bottom_z_mm"), "mm", door.get("source"), "Door bottom elevation."])
    ws.append(["width_mm", door.get("width_mm"), "mm", door.get("source"), "Inner opening width X2."])
    ws.append(["height_mm", door.get("height_mm"), "mm", door.get("source"), "Inner opening height Y2."])

    ws = wb.create_sheet("载荷摘要")
    ws.append(["source_file", "sheet", "component", "extreme", "load_case", "Mx_kNm", "My_kNm", "Mxy_kNm", "Mz_kNm", "Fx_kN", "Fy_kN", "Fz_kN"])
    load_summary = ir["semantic_ir"].get("loads_inventory", {}).get("load_workbook_summary", {})
    for item in load_summary.get("ultimate_cases", []):
        ws.append([
            item.get("source_file"),
            item.get("sheet"),
            item.get("component"),
            item.get("extreme"),
            item.get("load_case"),
            item.get("Mx_kNm"),
            item.get("My_kNm"),
            item.get("Mxy_kNm"),
            item.get("Mz_kNm"),
            item.get("Fx_kN"),
            item.get("Fy_kN"),
            item.get("Fz_kN"),
        ])

    ws = wb.create_sheet("FEM模型参数")
    ws.append(["参数组", "key", "value", "unit", "source", "description"])
    material = ir["semantic_ir"]["materials"][0]
    fem_rows = [
        ("meta", "model_level", ir["modeling_decision_ir"].get("level"), "", "ir", "Computable shell model level."),
        ("meta", "unit_system", ir["semantic_ir"].get("unit_system"), "", "ir", "Abaqus consistent unit system."),
        ("material", "name", material.get("name"), "", material.get("source"), "Steel grade."),
        ("material", "elastic_modulus_mpa", material.get("elastic_modulus_mpa"), "MPa", material.get("source"), "Elastic modulus."),
        ("material", "poisson_ratio", material.get("poisson_ratio"), "", material.get("source"), "Poisson ratio."),
        ("material", "yield_strength_mpa", material.get("yield_strength_mpa"), "MPa", material.get("source"), "Nominal yield strength."),
        ("section", "type", ir["abaqus_assembly_ir"].get("sections", [{}])[0].get("type") if ir["abaqus_assembly_ir"].get("sections") else None, "", "ir", "Abaqus shell section."),
        ("element", "element_type", ir["abaqus_assembly_ir"].get("elements", [{}])[0].get("type") if ir["abaqus_assembly_ir"].get("elements") else None, "", "ir", "Abaqus element type."),
        ("mesh", "strategy", ir["model_spec_ir"].get("mesh", {}).get("strategy") if ir["model_spec_ir"].get("mesh") else None, "", "ir", "Mesh strategy."),
        ("mesh", "global_size_mm", ir["mesh_and_set_ir"].get("global_size_mm"), "mm", "ir", "Global mesh target."),
        ("mesh", "door_size_mm", ir["mesh_and_set_ir"].get("door_size_mm"), "mm", "ir", "Door local mesh target."),
        ("loads", "top_total_force_n", ir["semantic_ir"].get("loads_inventory", {}).get("top_total_force_n"), "N", "ir", "Computable top total force."),
        ("loads", "wind_pressure_mpa", ir["semantic_ir"].get("loads_inventory", {}).get("wind_pressure_mpa"), "MPa", "ir", "Computable wind pressure."),
        ("boundary", "name", ir["abaqus_assembly_ir"].get("boundary_conditions", [{}])[0].get("name") if ir["abaqus_assembly_ir"].get("boundary_conditions") else None, "", "ir", "Boundary condition name."),
    ]
    for row in fem_rows:
        ws.append(list(row))

    ws = wb.create_sheet("缺失项清单")
    ws.append(["item", "reason"])
    for item in ir.get("raw_excel_ir", {}).get("missing_items", []):
        ws.append([item.get("item"), item.get("reason")])

    for sheet in wb.worksheets:
        for column in sheet.columns:
            max_length = max((len(str(cell.value)) if cell.value is not None else 0 for cell in column), default=8)
            sheet.column_dimensions[column[0].column_letter].width = min(max(max_length + 2, 10), 42)
    wb.save(path)
    return path


def filter_flange_rows_for_range_matching(flanges: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[str]]:
    required = [
        "diameter_mm",
        "bolt_circle_mm",
        "inner_diameter_mm",
        "total_thickness_mm",
        "bolt_count",
        "bolt_diameter_mm",
    ]
    valid: list[dict[str, Any]] = []
    warnings: list[str] = []
    for index, flange in enumerate(flanges, start=1):
        missing = [key for key in required if not _is_numeric(flange.get(key))]
        if missing:
            warnings.append(
                f"Flange row {index} was skipped for flange-zone matching because these fields are not numeric: {', '.join(missing)}."
            )
            continue
        valid.append(flange)
    return valid, warnings


def _is_numeric(value: Any) -> bool:
    if value is None or value == "":
        return False
    try:
        float(value)
        return True
    except (TypeError, ValueError):
        return False


def build_ir(
    project_id: str,
    source_name: str,
    tower: list[dict[str, Any]],
    flanges: list[dict[str, Any]],
    door: dict[str, Any],
    fem_params: dict[str, Any] | None = None,
    import_meta: dict[str, Any] | None = None,
    modeling_goal: str | None = None,
) -> dict[str, Any]:
    fem_params = fem_params or {}
    import_meta = import_meta or {}
    missing_items = list(import_meta.get("missing_items", []))
    import_warnings = list(import_meta.get("warnings", []))
    flanges_for_ranges, flange_filter_warnings = filter_flange_rows_for_range_matching(flanges)
    import_warnings.extend(flange_filter_warnings)
    detected_format = import_meta.get("detected_format", "legacy_fixed_layout")
    sheet_map = import_meta.get("sheet_map", {})
    load_summary = import_meta.get("loads_inventory", {})
    fem_records = list(import_meta.get("fem_param_records", []))
    fem_record_map = build_fem_record_map(fem_records)
    modeling_goal_ir = build_modeling_goal_ir(modeling_goal)
    data_quality = build_data_quality_report(import_meta, missing_items, import_warnings)
    traceability_summary = build_traceability_summary(source_name, import_meta, tower, flanges, door)
    has_fem = bool(fem_params)
    material_name = str(fem_params.get("material.name", ""))
    elastic_modulus = fem_params.get("material.elastic_modulus_mpa")
    poisson_ratio = fem_params.get("material.poisson_ratio")
    yield_strength = fem_params.get("material.yield_strength_mpa")
    material_density = fem_params.get("material.density_tonne_mm3", fem_params.get("material.density"))
    top_force = fem_params.get("loads.top_total_force_n")
    top_force_direction = fem_params.get("loads.top_force_direction")
    wind_pressure = fem_params.get("loads.wind_pressure_mpa")
    wind_surface = fem_params.get("loads.wind_surface") or fem_params.get("sets_surfaces.wind_surface", "WINDWARD_SURFACE")
    part_name = str(fem_params.get("element.part_name", "WINDTOWER_SHELL_V3"))
    instance_name = str(fem_params.get("element.instance_name", f"{part_name}-1"))
    element_type = str(fem_params.get("element.element_type", ""))
    section_type = str(fem_params.get("section.type", ""))
    section_material = str(fem_params.get("section.material", material_name or "STEEL"))
    model_level = str(fem_params.get("meta.model_level", "L2"))
    base_set = str(fem_params.get("sets_surfaces.base_set", "BASE_FIXED"))
    top_set = str(fem_params.get("sets_surfaces.top_set", "TOP_LOAD_NODES"))
    door_set = str(fem_params.get("sets_surfaces.door_set", "DOOR_OPENING_NODES"))
    door_neighbor_set = str(fem_params.get("sets_surfaces.door_neighbor_set", "DOOR_NEIGHBOR_ELEMENTS"))
    flange_set = str(fem_params.get("sets_surfaces.flange_set", "FLANGE_ELEMENTS"))
    legacy_output_variables = fem_params.get("outputs.field", ["U", "RF", "S", "E"])
    output_contract = normalize_output_contract(legacy_field=legacy_output_variables)
    project_meta_ir = build_project_meta_ir(project_id, source_name, import_meta, modeling_goal_ir)
    sections = [
        {
            "section_id": f"SEG_{i:03d}",
            "z_bottom_mm": float(a["z_mm"]),
            "z_top_mm": float(b["z_mm"]),
            "diameter_bottom_mm": float(a["diameter_mm"]),
            "diameter_top_mm": float(b["diameter_mm"]),
            "thickness_mm": float(a["thickness_mm"]),
            "material": section_material or "STEEL",
            "source": "excel:tower",
            "source_trace": build_source_trace(source_name, sheet_map.get("tower") or "塔节参数", "auto", f"tower_sections.{i}", a, {"z_bottom_mm": float(a["z_mm"]), "z_top_mm": float(b["z_mm"]), "thickness_mm": float(a["thickness_mm"])}, "mm"),
        }
        for i, (a, b) in enumerate(zip(tower[:-1], tower[1:]), start=1)
    ]
    semantic_flanges = []
    ranges = find_flange_ranges(tower, flanges_for_ranges)
    for i, (flange, z0, z1) in enumerate(ranges, start=1):
        semantic_flanges.append(
            {
                "flange_id": f"FLANGE_{i:02d}",
                "elevation_mm": 0.5 * (z0 + z1),
                "mean_diameter_mm": float(flange["diameter_mm"]),
                "bolt_circle_mm": float(flange["bolt_circle_mm"]),
                "inner_diameter_mm": float(flange["inner_diameter_mm"]),
                "total_thickness_mm": float(flange["total_thickness_mm"]),
                "bolt_count": int(flange["bolt_count"]),
                "bolt_diameter_mm": float(flange["bolt_diameter_mm"]),
                "source": "excel:flange",
                "source_trace": build_source_trace(source_name, sheet_map.get("flange") or "法兰参数", "auto", f"flanges.{i}", flange, "normalized_flange", "mm"),
            }
        )
    def _door_number(*values: Any) -> float:
        return float(next((value for value in values if value is not None), 0.0))

    door_opening = {
        "enabled": bool(door),
        "shape": "ellipse",
        "center_z_mm": float(door.get("H", 0.0) + 0.5 * door.get("Y2", 0.0)) if door else None,
        "width_mm": float(door.get("X2", 0.0)) if door else None,
        "height_mm": float(door.get("Y2", 0.0)) if door else None,
        "bottom_z_mm": float(door.get("H", 0.0)) if door else None,
        "theta_center_deg": float(door.get("theta_center_deg", 0.0)) if door else 0.0,
        "local_reinforcement": bool(
            fem_params.get("door.reinforcement_enabled", door.get("local_reinforcement", False))
        ) if door else False,
        "reinforcement_enabled": bool(
            fem_params.get("door.reinforcement_enabled", door.get("local_reinforcement", False))
        ) if door else False,
        "outer_width_mm": _door_number(fem_params.get("door.outer_width_mm"), door.get("X1"), door.get("X2")) if door else None,
        "outer_height_mm": _door_number(fem_params.get("door.outer_height_mm"), door.get("Y1"), door.get("Y2")) if door else None,
        "frame_thickness_mm": _door_number(fem_params.get("door.frame_thickness_mm"), door.get("W")) if door else None,
        "reinforcement_material": str(fem_params.get("door.reinforcement_material", "Q355ND")),
        "source": "excel:door" if door else "missing",
        "source_trace": build_source_trace(source_name, sheet_map.get("door") or "门洞参数", "auto", "door_openings.0", door or None, "normalized_door_opening" if door else None, "mm"),
    }
    top_equivalent_system = {
        "mass_tonne": fem_params.get("loads.top_mass_tonne"),
        "center_of_mass_offset_mm": fem_params.get("loads.center_of_mass_offset_mm"),
        "rotary_inertia": {
            "Ixx": fem_params.get("loads.top_rotary_inertia_Ixx_tonne_mm2"),
            "Iyy": fem_params.get("loads.top_rotary_inertia_Iyy_tonne_mm2"),
            "Izz": fem_params.get("loads.top_rotary_inertia_Izz_tonne_mm2"),
            "enabled": bool(fem_params.get("loads.top_rotary_inertia_enabled")),
        },
        "coupling_type": fem_params.get("loads.top_coupling_type", "kinematic"),
        "top_ring_coupling_enabled": bool(fem_params.get("coupling.top_ring_enabled", True)),
        "coupled_region": top_set,
        "source_trace": source_trace_for_fem(fem_record_map, source_name, "loads.top_mass_tonne", fem_params.get("loads.top_mass_tonne"), "tonne"),
    }
    engineering_parameters = {
        "tower_geometry": {
            "tower_height_mm": float(tower[-1]["z_mm"]) if tower else 0.0,
            "tower_sections": sections,
            "source_trace": build_source_trace(source_name, sheet_map.get("tower") or "塔节参数", "auto", "engineering_parameters.tower_geometry", f"{len(tower)} tower points", len(sections), "mm"),
        },
        "material_mapping": [{"section_id": item["section_id"], "material": item["material"], "source_trace": item["source_trace"]} for item in sections],
        "door_geometry": door_opening,
        "flanges": semantic_flanges,
        "top_equivalent_system": top_equivalent_system,
    }
    fem_parameters = {
        "element": {"type": element_type or "S4R", "source_trace": source_trace_for_fem(fem_record_map, source_name, "element.element_type", element_type or "S4R", None)},
        "section": {"type": section_type or "Shell Section", "thickness_source": fem_params.get("section.thickness_source", "excel_tower_sections"), "source_trace": source_trace_for_fem(fem_record_map, source_name, "section.type", section_type or "Shell Section", None)},
        "mesh": {
            "global_size_mm": float(fem_params.get("mesh.global_size_mm", CANONICAL_DEFAULTS["mesh.global_size_mm"])),
            "door_size_mm": float(fem_params.get("mesh.door_size_mm", CANONICAL_DEFAULTS["mesh.door_size_mm"])),
            "flange_size_mm": float(fem_params.get("mesh.flange_size_mm", 500.0)),
            "hole_layers": int(fem_params.get("mesh.hole_layers", CANONICAL_DEFAULTS["mesh.hole_layers"])),
            "patch_margin_u_mm": float(fem_params.get("mesh.patch_margin_u_mm", CANONICAL_DEFAULTS["mesh.patch_margin_u_mm"])),
            "patch_margin_z_mm": float(fem_params.get("mesh.patch_margin_z_mm", CANONICAL_DEFAULTS["mesh.patch_margin_z_mm"])),
            "source_trace": source_trace_for_fem(fem_record_map, source_name, "mesh.global_size_mm", fem_params.get("mesh.global_size_mm", CANONICAL_DEFAULTS["mesh.global_size_mm"]), "mm"),
        },
        "sets_surfaces": {
            "base_set": base_set,
            "top_set": top_set,
            "door_boundary_set": door_set,
            "door_neighbor_elset": door_neighbor_set,
            "wind_surface": wind_surface,
            "flange_elset": flange_set,
            "source_trace": source_trace_for_fem(fem_record_map, source_name, "sets_surfaces.base_set", base_set, None),
        },
        "material": {
            "name": material_name,
            "elastic_modulus_mpa": elastic_modulus,
            "poisson_ratio": poisson_ratio,
            "density": material_density,
            "yield_strength_mpa": yield_strength,
            "source_trace": source_trace_for_fem(fem_record_map, source_name, "material.name", material_name, None),
        },
        "analysis_steps": build_fem_steps(fem_params),
        "loads": build_fem_loads(fem_params),
        "boundary_conditions": build_fem_boundaries(fem_params),
        "outputs": output_contract,
    }
    ir = {
        "project_meta_ir": project_meta_ir,
        "modeling_goal_ir": modeling_goal_ir,
        "data_quality": data_quality,
        "data_quality_ir": data_quality,
        "traceability_summary": traceability_summary,
        "source_trace": traceability_summary,
        "patch_history": [],
        "excel_facts_ir": {
            "source_file": source_name,
            "unit_system": "mm-N-MPa",
            "source_trace": traceability_summary,
            "engineering_parameters": engineering_parameters,
            "fem_parameters": fem_parameters,
            "tower_profile_points": [
                {
                    "z_mm": float(point["z_mm"]),
                    "diameter_mm": float(point["diameter_mm"]),
                    "thickness_mm": float(point["thickness_mm"]),
                    "source": "excel:tower",
                    "source_trace": build_source_trace(source_name, sheet_map.get("tower") or "塔节参数", "auto", "tower_profile_points", point, "normalized_tower_point", "mm"),
                }
                for point in tower
            ],
            "door": door_opening,
            "flanges": semantic_flanges,
        },
        "model_spec_ir": {
            "material": {"name": material_name, "source": "excel"} if material_name else None,
            "section": {"type": section_type, "thickness_source": fem_params.get("section.thickness_source"), "source": "excel"} if section_type else None,
            "element": {"type": element_type, "part_name": part_name, "instance_name": instance_name, "source": "excel"} if element_type else None,
            "mesh": {"strategy": fem_params.get("mesh.strategy", CANONICAL_DEFAULTS["mesh.strategy"]), "source": "excel"} if has_fem else None,
            "steps": build_fem_steps(fem_params),
            "loads": build_fem_loads(fem_params),
            "boundary_conditions": build_fem_boundaries(fem_params),
            "interactions": build_fem_interactions(fem_params),
            "outputs": output_contract,
        },
        "raw_excel_ir": {
            "source_files": import_meta.get("source_files") or [source_name],
            "source_file": import_meta.get("source_file", source_name),
            "source_file_original": import_meta.get("source_file_original", source_name),
            "source_file_storage": import_meta.get("source_file_storage"),
            "detected_format": detected_format,
            "input_format": import_meta.get("input_format") or detected_format,
            "format_confidence": import_meta.get("format_confidence"),
            "parse_strategy": import_meta.get("parse_strategy"),
            "llm_required": import_meta.get("llm_required"),
            "promotion_status": import_meta.get("promotion_status"),
            "sheet_map": sheet_map,
            "raw_tables": import_meta.get("raw_tables", {}),
            "workbook_profile": import_meta.get("workbook_profile", {}),
            "standard_input_candidate": import_meta.get("standard_input_candidate") or import_meta.get("standard_input_package"),
            "standard_input_package": import_meta.get("standard_input_package", {}),
            "standard_input_validation": import_meta.get("validation_report", {}),
            "input_quality_scores": import_meta.get("input_quality_scores", {}),
            "requires_user_confirmation": bool(import_meta.get("requires_user_confirmation", False)),
            "llm_parse_assist": import_meta.get("llm_parse_assist", {}),
            "llm_standard_input_candidate": import_meta.get("llm_standard_input_candidate"),
            "llm_standard_input_validation": import_meta.get("llm_standard_input_validation"),
            "llm_promotion_status": import_meta.get("llm_promotion_status"),
            "parse_diagnostics": import_meta.get("parse_diagnostics", {}),
            "parsed_facts": import_meta.get("parsed_facts", {}),
            "fem_param_records": import_meta.get("fem_param_records", []),
            "fem_model_facts": import_meta.get("fem_model_facts", {}),
            "derived_facts": import_meta.get("derived_facts", {}),
            "standard_references": import_meta.get("standard_references", []),
            "confidence_report": import_meta.get("confidence_report", {}),
            "llm_suggestions": import_meta.get("llm_suggestions", {}),
            "engineering_fact_package": import_meta.get("engineering_fact_package", {}),
            "candidate_mapping_ir": import_meta.get("candidate_mapping_ir", {}),
            "model_family_detection": import_meta.get("model_family_detection", {}),
            "missing_items": missing_items,
            "warnings": import_warnings,
            "tables": [
                {"sheet": sheet_map.get("tower") or "塔节参数", "range": "auto", "purpose": "tower profile and shell thickness"},
                {"sheet": sheet_map.get("flange") or "法兰参数", "range": "auto", "purpose": "flange and bolt parameters"},
                {"sheet": sheet_map.get("door") or "门洞参数", "range": "auto", "purpose": "door opening geometry"},
            ],
            "field_aliases": {"tower_height": "塔筒总高", "diameter": "中径/直径", "thickness": "壁厚"},
            "parse_warnings": import_warnings,
        },
        "semantic_ir": {
            "project_id": project_id,
            "unit_system": "mm-N-MPa",
            "tower_sections": sections,
            "materials": [
                {
                    "name": material_name,
                    "elastic_modulus_mpa": elastic_modulus,
                    "poisson_ratio": poisson_ratio,
                    "density": material_density,
                    "yield_strength_mpa": yield_strength,
                    "source": "excel" if material_name and elastic_modulus and poisson_ratio is not None else "missing",
                    "source_trace": source_trace_for_fem(fem_record_map, source_name, "material.name", material_name, None),
                }
            ],
            "door_openings": [door_opening],
            "flanges": semantic_flanges,
            "bolt_groups": [],
            "foundation": {"enabled": False, "material": "", "anchor_strategy": "none", "source": "missing"},
            "loads_inventory": {
                "top_total_force_n": top_force,
                "top_force_direction": top_force_direction,
                "top_horizontal_force_n": top_force,
                "wind_pressure_mpa": wind_pressure,
                "wind_pressure_profile": fem_params.get("loads.wind_pressure_profile"),
                "wind_surface": wind_surface,
                "wind_direction_deg": fem_params.get("loads.wind_direction_deg"),
                "wind_direction_vector": fem_params.get("loads.wind_direction_vector"),
                "top_force_parallel_n": fem_params.get("loads.top_force_parallel_n"),
                "top_force_transverse_n": fem_params.get("loads.top_force_transverse_n"),
                "top_force_vertical_n": fem_params.get("loads.top_force_vertical_n"),
                "top_moment_roll_nmm": fem_params.get("loads.top_moment_roll_nmm"),
                "top_moment_pitch_nmm": fem_params.get("loads.top_moment_pitch_nmm"),
                "top_moment_yaw_nmm": fem_params.get("loads.top_moment_yaw_nmm"),
                "hub_vertical_offset_mm": fem_params.get("loads.hub_vertical_offset_mm"),
                "gravity_enabled": bool(fem_params.get("loads.gravity_enabled", True)),
                "gravity_acceleration_mm_s2": fem_params.get("loads.gravity_acceleration_mm_s2", 9810.0),
                "top_mass_tonne": top_equivalent_system.get("mass_tonne"),
                "center_of_mass_offset_mm": top_equivalent_system.get("center_of_mass_offset_mm"),
                "load_workbook_summary": load_summary,
                "source": "excel" if top_force and wind_pressure is not None else load_summary.get("source", "missing"),
                "source_trace": source_trace_for_fem(fem_record_map, source_name, "loads.top_total_force_n", top_force, "N"),
            },
            "top_equivalent_system": top_equivalent_system,
            "provenance": {"generator": "windtower_backend", "source_excel": source_name, "modeling_goal_source": modeling_goal_ir["source"]},
        },
        "modeling_decision_ir": {
            "level": model_level,
            "tower_element_family": "shell",
            "door_strategy": str(fem_params.get("mesh.door_strategy", "cutout_with_local_refinement")),
            "flange_strategy": str(fem_params.get("interaction.flange_strategy", "ignore")),
            "bolt_strategy": str(fem_params.get("interaction.bolt_strategy", "ignore")),
            "contact_strategy": str(fem_params.get("interaction.contact_strategy", "none")),
            "pretension_workflow": str(fem_params.get("interaction.pretension_workflow", "none")),
            "assumptions": ["FEM parameters are read from Excel sheet FEM模型参数."] if has_fem else ["Excel currently supplies geometry facts only; engineering model choices require confirmation."],
            "risks": [] if has_fem else ["Do not generate computable INP until required modeling tasks are confirmed."],
        },
        "mesh_and_set_ir": {
            "global_size_mm": float(fem_params.get("mesh.global_size_mm", CANONICAL_DEFAULTS["mesh.global_size_mm"])),
            "door_size_mm": float(fem_params.get("mesh.door_size_mm", CANONICAL_DEFAULTS["mesh.door_size_mm"])),
            "flange_size_mm": float(fem_params.get("mesh.flange_size_mm", 500.0)),
            "hole_layers": int(fem_params.get("mesh.hole_layers", CANONICAL_DEFAULTS["mesh.hole_layers"])),
            "patch_margin_u_mm": float(fem_params.get("mesh.patch_margin_u_mm", CANONICAL_DEFAULTS["mesh.patch_margin_u_mm"])),
            "patch_margin_z_mm": float(fem_params.get("mesh.patch_margin_z_mm", CANONICAL_DEFAULTS["mesh.patch_margin_z_mm"])),
            "local_refinement_zones": [
                {
                    "name": "DOOR_PATCH",
                    "target": "door_opening",
                    "size_mm": float(fem_params.get("mesh.door_size_mm", CANONICAL_DEFAULTS["mesh.door_size_mm"])),
                    "reason": "V3 local refined conforming quadrilateral door mesh",
                    "source": "excel",
                }
            ] if fem_params.get("mesh.strategy") else [],
            "required_sets": [base_set, top_set, door_set, door_neighbor_set, flange_set],
            "required_surfaces": [str(fem_params.get("sets_surfaces.wind_surface", "WINDWARD_SURFACE"))],
            "naming_rules": {"part": part_name, "instance": instance_name},
            "quality_thresholds": {"min_positive_area": 1.0e-6, "max_aspect_ratio": 10.0, "door_boundary_closed": True},
            "mesh_quality": {"min_positive_area": 1.0e-6, "max_aspect_ratio": 10.0, "door_boundary_closed": True},
            "source_trace": source_trace_for_fem(fem_record_map, source_name, "mesh.global_size_mm", fem_params.get("mesh.global_size_mm", CANONICAL_DEFAULTS["mesh.global_size_mm"]), "mm"),
        },
        "abaqus_assembly_ir": {
            "part": {"name": part_name, "source": "excel" if has_fem else "strategy_default"},
            "elements": [{"type": element_type, "target": "tower_shell", "source": "excel"}] if element_type else [],
            "sections": [{"type": section_type, "material": section_material, "thickness": fem_params.get("section.thickness_source"), "source": "excel"}] if section_type else [],
            "materials": [
                {
                    "name": material_name or section_material,
                    "elastic_modulus_mpa": elastic_modulus,
                    "poisson_ratio": poisson_ratio,
                    "density": material_density,
                    "yield_strength_mpa": yield_strength,
                    "source": "excel" if material_name else "missing",
                }
            ],
            "steps": build_fem_steps(fem_params),
            "boundary_conditions": build_fem_boundaries(fem_params),
            "loads": build_fem_loads(fem_params),
            "interactions": [{"name": "NONE_L2", "type": "none", "status": "not_required", "source": "not_required"}],
            "outputs": output_contract,
            "datacheck_rules": ["required_keywords", "nonempty_sets", "top_force_sum", "wind_surface_spos_sneg"],
        },
        "top_equivalent_system": top_equivalent_system,
    }
    requested_analysis = list(modeling_goal_ir.get("analysis_intent") or [])
    if fem_params.get("analysis.static_enabled") is True and "static_strength" not in requested_analysis:
        requested_analysis.append("static_strength")
    if fem_params.get("analysis.static_enabled") is False:
        requested_analysis = [item for item in requested_analysis if item != "static_strength"]
    if fem_params.get("analysis.modal_enabled") is True and "modal" not in requested_analysis:
        requested_analysis.append("modal")
    if fem_params.get("analysis.modal_enabled") is False:
        requested_analysis = [item for item in requested_analysis if item != "modal"]
    ir["analysis_request"] = {"analysis_types": requested_analysis or ["static_strength", "modal"]}
    ir["analysis_cases"] = {
        "modal": {"modes": int(fem_params.get("modal.modes") or 20)},
    }
    ir["modeling_tasks"] = build_tasks(ir)
    return ir


def build_modeling_goal_ir(modeling_goal: str | None) -> dict[str, Any]:
    text = (modeling_goal or "").strip() or None
    lower = (text or "").lower()
    raw = text or ""
    analysis_intent = []
    if any(token in lower for token in ("static", "strength")) or any(token in raw for token in ("静力", "强度")):
        analysis_intent.append("static_strength")
    if any(token in lower for token in ("modal", "frequency")) or any(token in raw for token in ("模态", "频率")):
        analysis_intent.append("modal")
    if any(token in lower for token in ("buckling",)) or any(token in raw for token in ("屈曲", "稳定")):
        analysis_intent.append("buckling")
    if any(token in lower for token in ("fatigue",)) or "疲劳" in raw:
        analysis_intent.append("fatigue")
    if any(token in lower for token in ("door", "mises")) or any(token in raw for token in ("门洞", "应力")):
        analysis_intent.append("door_local_stress")
    output_preferences = []
    for key, tokens in {
        "displacement": ("displacement", "位移", "u"),
        "reaction_force": ("reaction", "反力", "rf"),
        "stress": ("stress", "mises", "应力", "s"),
        "modal_frequency": ("frequency", "模态频率", "频率"),
        "buckling_factor": ("buckling factor", "屈曲因子", "屈曲模态"),
    }.items():
        if any(token in lower for token in tokens) or any(token in raw for token in tokens):
            output_preferences.append(key)
    modeling_preferences = []
    if any(token in lower for token in ("shell", "s4r")) or "壳" in raw:
        modeling_preferences.append("shell_element")
    if any(token in lower for token in ("refine",)) or any(token in raw for token in ("加密", "门洞")):
        modeling_preferences.append("door_local_refinement")
    if any(token in lower for token in ("flange",)) or "法兰" in raw:
        modeling_preferences.append("flange_simplification_or_interaction")
    return {
        "text": text,
        "source": "user_optional",
        "created_at": datetime.now(timezone.utc).isoformat(),
        "role": "semantic_constraint_only",
        "usage": ["analysis_intent", "output_preference", "missing_item_context"],
        "analysis_intent": analysis_intent,
        "output_preferences": output_preferences,
        "modeling_preferences": modeling_preferences,
        "is_numeric_fact_source": False,
    }


def build_data_quality_report(import_meta: dict[str, Any], missing_items: list[dict[str, Any]], warnings: list[str]) -> dict[str, Any]:
    confidence = import_meta.get("confidence_report", {}) or {}
    checks = confidence.get("checks", {}) or {}
    parsed = import_meta.get("parsed_facts", {}) or {}
    fem_records = import_meta.get("fem_param_records", []) or []
    has_tower = bool(parsed.get("tower") or import_meta.get("tower"))
    has_fem = bool(import_meta.get("fem_params") or fem_records)
    ambiguous = []
    required_checks = {
        "geometry": has_tower,
        "material": bool(import_meta.get("fem_params", {}).get("material.name") or any(item.get("group") == "material" for item in fem_records)),
        "thickness": bool(parsed.get("tower") or import_meta.get("tower")),
        "door": bool(checks.get("door_fields_complete", True)),
        "loads": bool(import_meta.get("fem_params", {}).get("loads.top_total_force_n") or any(item.get("group") == "loads" for item in fem_records)),
        "boundary": bool(import_meta.get("fem_params", {}).get("boundary.name") or any(item.get("group") == "boundary" for item in fem_records)),
        "analysis_steps": bool(import_meta.get("fem_params", {}).get("steps.static_step_name") or any(item.get("group") == "steps" for item in fem_records)),
        "outputs": True,
    }
    data_missing_items = list(missing_items)
    for key, ok in required_checks.items():
        if not ok and not any(item.get("item") == key for item in data_missing_items):
            data_missing_items.append({"item": key, "reason": f"Required {key} data is missing or not traceable."})
    for item in missing_items:
        ambiguous.append({"field": item.get("item", "unknown"), "status": "requires_user_confirmation", "reason": item.get("reason", "")})
    for item in data_missing_items:
        if item not in missing_items:
            ambiguous.append({"field": item.get("item", "unknown"), "status": "requires_user_confirmation", "reason": item.get("reason", "")})
    for warning in warnings:
        ambiguous.append({"field": "import_warning", "status": "ambiguous", "reason": str(warning)})

    completeness_blocked = not has_tower
    completeness_status = "blocked" if completeness_blocked else ("warning" if data_missing_items else "pass")
    range_status = "warning" if warnings else "pass"
    unit_status = "pass" if checks.get("unit_system_declared", True) else "warning"
    semantic_status = "pass" if has_fem and not data_missing_items else ("warning" if has_fem else "blocked")
    blocked = completeness_blocked or semantic_status == "blocked"

    return {
        "summary": {"status": "blocked" if blocked else ("warning" if data_missing_items or warnings else "pass"), "blocked": blocked},
        "completeness": {
            "status": completeness_status,
            "required_groups": ["geometry", "material", "section", "loads", "boundary", "analysis_steps", "outputs"],
            "missing_items": data_missing_items,
            "checks": required_checks,
        },
        "range": {"status": range_status, "warnings": warnings},
        "unit": {"status": unit_status, "target_unit_system": "mm-N-MPa", "checks": checks},
        "semantic": {
            "status": semantic_status,
            "checks": {
                "engineering_and_fem_parameters_in_same_workbook": has_tower and has_fem,
                "fem_required_fields_complete": bool(checks.get("fem_required_fields_complete", not missing_items)),
                "source_trace_available": bool(fem_records or import_meta.get("sheet_map")),
            },
        },
        "failure_grades": ["auto_correctable", "rule_inferred", "requires_user_confirmation", "blocking"],
        "missing_items": data_missing_items,
        "ambiguous_fields": ambiguous,
        "confidence_report": confidence,
    }


def build_project_meta_ir(project_id: str, source_name: str, import_meta: dict[str, Any], modeling_goal_ir: dict[str, Any]) -> dict[str, Any]:
    source_files = import_meta.get("source_files") or [source_name]
    return {
        "project_id": project_id,
        "model_name": Path(source_name).stem,
        "unit_system": "mm-N-MPa",
        "coordinate_system": {
            "height_axis": "Z",
            "horizontal_axes": ["X", "Y"],
            "origin": "tower_base_center",
            "positive_z": "tower_top",
        },
        "source_files": source_files,
        "uploaded_at": datetime.now(timezone.utc).isoformat(),
        "modeling_goal_ref": {
            "source": modeling_goal_ir.get("source", "user_optional"),
            "is_numeric_fact_source": False,
            "has_text": bool(modeling_goal_ir.get("text")),
        },
    }


def build_fem_record_map(records: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    result: dict[str, dict[str, Any]] = {}
    for item in records:
        candidates = [
            item.get("path"),
            f"{item.get('group')}.{item.get('key')}" if item.get("group") and item.get("key") else None,
            item.get("key"),
        ]
        for key in candidates:
            if key:
                result[str(key)] = item
    return result


def source_trace_for_fem(
    record_map: dict[str, dict[str, Any]],
    source_name: str,
    field_name: str,
    normalized_value: Any,
    unit: str | None,
) -> dict[str, Any]:
    record = record_map.get(field_name) or {}
    source = str(record.get("source") or "")
    sheet, cell_range = split_source_ref(source)
    return build_source_trace(
        source_name,
        sheet or "FEM模型参数",
        cell_range or "auto",
        field_name,
        record.get("raw_value", normalized_value),
        record.get("value", normalized_value),
        str(record.get("unit") or unit or ""),
        confidence=str(record.get("confidence", "parsed" if record else "defaulted")),
        parse_rule="fem_param_record" if record else "framework_default",
    )


def split_source_ref(source: str) -> tuple[str | None, str | None]:
    if not source:
        return None, None
    if "!" in source:
        sheet, cell = source.split("!", 1)
        return sheet or None, cell or None
    if ":" in source:
        sheet, cell = source.split(":", 1)
        return sheet or None, cell or None
    return source, None


def build_source_trace(
    workbook: str,
    sheet: str,
    cell_range: str,
    field_name: str,
    raw_value: Any,
    normalized_value: Any,
    unit: str | None,
    confidence: str = "parsed",
    parse_rule: str = "deterministic_importer",
) -> dict[str, Any]:
    return {
        "workbook": workbook,
        "sheet": sheet,
        "cell_range": cell_range,
        "field_name": field_name,
        "raw_value": raw_value,
        "normalized_value": normalized_value,
        "unit": unit or "",
        "confidence": confidence,
        "parse_rule": parse_rule,
    }


def build_traceability_summary(
    source_name: str,
    import_meta: dict[str, Any],
    tower: list[dict[str, Any]],
    flanges: list[dict[str, Any]],
    door: dict[str, Any],
) -> dict[str, Any]:
    fem_records = import_meta.get("fem_param_records", []) or []
    source_files = import_meta.get("source_files") or [source_name]
    return {
        "workbook": source_name,
        "source_files": source_files,
        "sheet_map": import_meta.get("sheet_map", {}),
        "engineering_parameter_sources": {
            "tower_profile_points": {"count": len(tower), "source": "excel:tower"},
            "flanges": {"count": len(flanges), "source": "excel:flange"},
            "door": {"count": 1 if door else 0, "source": "excel:door" if door else "missing"},
        },
        "fem_parameter_sources": {
            "count": len(fem_records),
            "records": [
                {
                    "key": item.get("key"),
                    "group": item.get("group"),
                    "source": item.get("source"),
                    "raw_value": item.get("raw_value"),
                    "normalized_value": item.get("value"),
                    "unit": item.get("unit"),
                    "confidence": item.get("confidence", "parsed"),
                }
                for item in fem_records[:80]
            ],
        },
        "source_trace_fields": ["workbook", "sheet", "cell_range", "field_name", "raw_value", "normalized_value", "unit", "parse_rule", "confidence"],
        "missing_items_count": len(import_meta.get("missing_items", []) or []),
        "warnings_count": len(import_meta.get("warnings", []) or []),
    }


def _legacy_build_tasks_geometry_only(ir: dict[str, Any]) -> list[dict[str, Any]]:
    return [
        {"key": "geometry", "label": "几何", "name": "Tower profile + door + flanges", "status": "ready", "source": "excel", "message": "由 Excel 解析得到，只作为几何事实。"},
        {"key": "materials", "label": "材料", "name": "待定义", "status": "missing", "source": "engineer", "message": "请定义钢材牌号、弹性模量、泊松比，必要时定义密度。"},
        {"key": "sections", "label": "截面", "name": "待确认", "status": "missing", "source": "engineer", "message": "请确认采用 Shell Section，厚度是否使用 Excel 塔节壁厚。"},
        {"key": "elements", "label": "单元", "name": "待定义", "status": "missing", "source": "engineer", "message": "请确认塔筒主体单元类型，例如 S4R 壳单元。"},
        {"key": "mesh", "label": "网格", "name": "待定义", "status": "missing", "source": "engineer", "message": "请定义全局尺寸、门洞局部尺寸和 V3 conforming door mesh 策略。"},
        {"key": "steps", "label": "分析步", "name": "待定义", "status": "missing", "source": "engineer", "message": "请定义静力步、模态步或其他分析过程。"},
        {"key": "loads", "label": "荷载", "name": "待定义", "status": "missing", "source": "engineer", "message": "请定义顶部力、风压、重力等荷载。"},
        {"key": "boundary", "label": "边界条件", "name": "待定义", "status": "missing", "source": "engineer", "message": "请定义底部约束或基础等效边界。"},
        {"key": "interactions", "label": "相互作用", "name": "L1/L2 不需要", "status": "not_required", "source": "not_required", "message": "首版 L1/L2 壳模型不要求接触/预紧；切换 L3 后需要补全。"},
    ]


def read_fem_params(excel_path: Path) -> dict[str, Any]:
    try:
        import xlrd

        book = xlrd.open_workbook(str(excel_path))
        if "FEM模型参数" not in book.sheet_names():
            return {}
        sheet = book.sheet_by_name("FEM模型参数")
    except Exception:
        return {}

    params: dict[str, Any] = {}
    for row in range(1, sheet.nrows):
        group = str(sheet.cell_value(row, 0)).strip()
        key = str(sheet.cell_value(row, 1)).strip()
        if group and key:
            params[f"{group}.{key}"] = parse_fem_value(sheet.cell_value(row, 2))
    return params


def parse_fem_value(value: Any) -> Any:
    if isinstance(value, float):
        return int(value) if value.is_integer() else value
    text = str(value).strip()
    if text == "":
        return None
    if "," in text and not re.search(r"[A-Za-z_]+,", text):
        return [parse_fem_value(item) for item in text.split(",")]
    try:
        number = float(text)
        return int(number) if number.is_integer() else number
    except ValueError:
        return text


def build_fem_steps(params: dict[str, Any]) -> list[dict[str, Any]]:
    if not params.get("steps.STEP-1-TOP-REACTION") or not params.get("steps.STEP-2-WIND-PRESSURE"):
        return []
    return [
        {"name": "STEP-1-TOP-REACTION", "procedure": "static", "purpose": params["steps.STEP-1-TOP-REACTION"], "nlgeom": params.get("steps.nlgeom", "NO"), "source": "excel"},
        {"name": "STEP-2-WIND-PRESSURE", "procedure": "static", "purpose": params["steps.STEP-2-WIND-PRESSURE"], "nlgeom": params.get("steps.nlgeom", "NO"), "source": "excel"},
    ]


def build_fem_loads(params: dict[str, Any]) -> list[dict[str, Any]]:
    top_force = params.get("loads.top_total_force_n")
    wind_pressure = params.get("loads.wind_pressure_mpa")
    if top_force is None or wind_pressure is None:
        return []
    top_mass = params.get("loads.top_mass_tonne")
    eccentricity = params.get("loads.center_of_mass_offset_mm")
    gravity_enabled = bool(params.get("loads.gravity_enabled", True))
    gravity_acceleration = params.get("loads.gravity_acceleration_mm_s2", 9810.0)
    top_vertical = None
    eccentric_moment = None
    if top_mass is not None:
        top_vertical = -float(top_mass) * float(gravity_acceleration)
        if eccentricity is not None:
            eccentric_moment = abs(top_vertical) * float(eccentricity)
    loads = [
        {
            "name": "TOP_RP_THRUST",
            "type": "cload",
            "target": "TOP_RP_NSET" if params.get("coupling.top_ring_enabled", True) else params.get("sets_surfaces.top_set", "TOP_LOAD_NODES"),
            "step": "STEP-2-WIND-PRESSURE",
            "dof": params.get("loads.top_force_direction", 1),
            "magnitude": top_force,
            "source": "excel",
        },
        {
            "name": "WIND_PRESSURE",
            "type": "dsload",
            "target": params.get("loads.wind_surface", "WINDWARD_SURFACE"),
            "step": "STEP-2-WIND-PRESSURE",
            "magnitude": wind_pressure,
            "profile": params.get("loads.wind_pressure_profile"),
            "source": "excel",
        },
    ]
    if gravity_enabled:
        loads.append({
            "name": "TOWER_SELF_WEIGHT",
            "type": "dload_gravity",
            "target": "ALL_ELEMENTS",
            "step": "STEP-1-TOP-REACTION",
            "acceleration_mm_s2": gravity_acceleration,
            "direction": [0, 0, -1],
            "source": "excel",
        })
    if top_vertical is not None:
        loads.append({
            "name": "TOP_RP_VERTICAL_RNA_WEIGHT",
            "type": "cload",
            "target": "TOP_RP_NSET",
            "step": "STEP-1-TOP-REACTION",
            "dof": 3,
            "magnitude": top_vertical,
            "source": "excel",
        })
    if eccentric_moment is not None:
        loads.append({
            "name": "TOP_RP_ECCENTRIC_MOMENT",
            "type": "cload_moment",
            "target": "TOP_RP_NSET",
            "step": "STEP-2-WIND-PRESSURE",
            "dof": 5,
            "magnitude": eccentric_moment,
            "source": "excel",
        })
    return loads


def build_fem_boundaries(params: dict[str, Any]) -> list[dict[str, Any]]:
    name = params.get("boundary.name")
    target = params.get("boundary.target")
    dofs = params.get("boundary.dofs")
    if not name or not target or not dofs:
        return []
    return [{"name": name, "target": target, "dofs": dofs if isinstance(dofs, list) else [dofs], "source": "excel"}]


def build_fem_interactions(params: dict[str, Any]) -> list[dict[str, Any]]:
    level = params.get("interaction.level", params.get("meta.model_level", "L2"))
    if level == "L2":
        return [{"name": "NONE_L2", "type": "none", "status": "not_required", "source": "not_required"}]
    return [{"name": "L3_INTERACTION", "level": level, "source": "excel"}]


def _legacy_build_tasks_fem_defaults(ir: dict[str, Any]) -> list[dict[str, Any]]:
    material = ir["semantic_ir"]["materials"][0]
    sections = ir["abaqus_assembly_ir"]["sections"]
    elements = ir["abaqus_assembly_ir"]["elements"]
    steps = ir["abaqus_assembly_ir"]["steps"]
    loads = ir["abaqus_assembly_ir"]["loads"]
    boundaries = ir["abaqus_assembly_ir"]["boundary_conditions"]
    mesh = ir["model_spec_ir"].get("mesh")
    has_material = bool(material.get("name") and material.get("elastic_modulus_mpa") and material.get("poisson_ratio") is not None)
    has_section = bool(sections)
    has_element = bool(elements)
    has_mesh = bool(mesh and ir["mesh_and_set_ir"].get("door_size_mm"))
    has_steps = bool(steps)
    has_loads = bool(loads)
    has_boundary = bool(boundaries)

    def task(key: str, label: str, name: str, ok: bool, message: str) -> dict[str, Any]:
        return {"key": key, "label": label, "name": name if ok else "待定义", "status": "confirmed" if ok else "missing", "source": "excel" if ok else "engineer", "message": "Read from FEM模型参数." if ok else message}

    return [
        {"key": "geometry", "label": "几何", "name": "Tower profile + door + flanges", "status": "ready", "source": "excel", "message": "Read from Excel geometry sheets."},
        task("materials", "材料", material.get("name") or "Material", has_material, "Define steel grade, elastic modulus and Poisson ratio."),
        task("sections", "截面", sections[0]["type"] if sections else "Shell Section", has_section, "Confirm Shell Section and Excel thickness source."),
        task("elements", "单元", elements[0]["type"] if elements else "S4R", has_element, "Confirm tower shell element type."),
        task("mesh", "网格", mesh.get("strategy") if mesh else "V3 conforming door mesh", has_mesh, "Define global mesh size and door local refinement."),
        task("steps", "分析步", "STEP-1-TOP-REACTION + STEP-2-WIND-PRESSURE", has_steps, "Define static analysis steps."),
        task("loads", "荷载", "800kN top force + wind pressure", has_loads, "Define top force and wind pressure."),
        task("boundary", "边界条件", boundaries[0]["name"] if boundaries else "BASE_FIXED", has_boundary, "Define base fixed boundary condition."),
        {"key": "interactions", "label": "相互作用", "name": "L1/L2 not required", "status": "not_required", "source": "not_required", "message": "L2 shell model does not require contact or pretension."},
    ]


def build_tasks(ir: dict[str, Any]) -> list[dict[str, Any]]:
    material = ir["semantic_ir"]["materials"][0]
    sections = ir["abaqus_assembly_ir"]["sections"]
    elements = ir["abaqus_assembly_ir"]["elements"]
    steps = ir["abaqus_assembly_ir"]["steps"]
    loads = ir["abaqus_assembly_ir"]["loads"]
    boundaries = ir["abaqus_assembly_ir"]["boundary_conditions"]
    mesh = ir["model_spec_ir"].get("mesh")
    door = ir["semantic_ir"]["door_openings"][0]
    missing_items = ir.get("raw_excel_ir", {}).get("missing_items", [])
    loads_inventory = ir["semantic_ir"].get("loads_inventory", {})
    load_summary = loads_inventory.get("load_workbook_summary", {})
    has_load_inventory = bool(load_summary.get("ultimate_cases") or load_summary.get("fatigue_equivalent"))
    has_material = bool(material.get("name") and material.get("elastic_modulus_mpa") and material.get("poisson_ratio") is not None)
    has_section = bool(sections)
    has_element = bool(elements)
    has_mesh = bool(mesh and ir["mesh_and_set_ir"].get("door_size_mm"))
    has_steps = bool(steps)
    has_loads = bool(loads)
    has_boundary = bool(boundaries)

    def task(key: str, label: str, name: str, ok: bool, message: str) -> dict[str, Any]:
        return {
            "key": key,
            "label": label,
            "name": name if ok else "待定义",
            "status": "confirmed" if ok else "missing",
            "source": "excel" if ok else "engineer",
            "message": "Read from FEM model parameters." if ok else message,
        }

    return [
        {"key": "geometry_import", "label": "几何导入", "name": "Tower profile + flanges", "status": "ready", "source": "excel", "message": "Geometry facts were normalized by the importer."},
        {"key": "door_geometry", "label": "门洞", "name": "Elliptical door opening" if door.get("enabled") else "待补全", "status": "ready" if door.get("enabled") else "missing", "source": door.get("source", "missing"), "message": "Door values are available." if door.get("enabled") else "Door values are missing; a no-door draft model can still be previewed."},
        {"key": "load_inventory", "label": "载荷清单", "name": "Load workbook summary" if has_load_inventory else "待导入", "status": "ready" if has_load_inventory else "missing", "source": load_summary.get("source", "missing"), "message": "Load rows were summarized for review." if has_load_inventory else "No load workbook summary is attached."},
        {"key": "fem_defaults", "label": "FEM默认值", "name": "FEM model parameters" if mesh else "待补全", "status": "ready" if mesh else "missing", "source": "excel" if mesh else "engineer", "message": "FEM model parameters were read from Excel." if mesh else "Confirm material, section, element, mesh, steps, loads and boundary before build."},
        {"key": "manual_completion", "label": "人工补全", "name": "Missing items", "status": "missing" if missing_items else "not_required", "source": "engineer" if missing_items else "not_required", "message": "; ".join(item.get("reason", "") for item in missing_items) if missing_items else "No importer-level missing items."},
        task("materials", "材料", material.get("name") or "Material", has_material, "Define steel grade, elastic modulus and Poisson ratio."),
        task("sections", "截面", sections[0]["type"] if sections else "Shell Section", has_section, "Confirm Shell Section and Excel thickness source."),
        task("elements", "单元", elements[0]["type"] if elements else "S4R", has_element, "Confirm tower shell element type."),
        task("mesh", "网格", mesh.get("strategy") if mesh else "V3 conforming door mesh", has_mesh, "Define global mesh size and door local refinement."),
        task("steps", "分析步", "STEP-1-TOP-REACTION + STEP-2-WIND-PRESSURE", has_steps, "Define static analysis steps."),
        task("loads", "荷载", "top force + wind pressure", has_loads, "Define computable top force and wind pressure."),
        task("boundary", "边界条件", boundaries[0]["name"] if boundaries else "BASE_FIXED", has_boundary, "Define base fixed boundary condition."),
        {"key": "interactions", "label": "相互作用", "name": "L1/L2 not required", "status": "not_required", "source": "not_required", "message": "L2 shell model does not require contact or pretension."},
    ]


def geometry_from_ir(ir: dict[str, Any]) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, Any]]:
    # ParameterSet promotion updates semantic_ir.tower_sections.  It must be
    # the sole geometry authority after confirmation; falling back to the raw
    # Excel projection here previously caused reviewed parameter changes to be
    # ignored by meshing and INP compilation.
    semantic_sections = ir.get("semantic_ir", {}).get("tower_sections") or []
    if semantic_sections:
        tower = [
            {
                "index": index,
                "z_mm": point.z_mm,
                "diameter_mm": point.diameter_mm,
                "length_mm": None,
                "thickness_mm": point.thickness_mm,
                "source": point.source,
            }
            for index, point in enumerate(normalize_tower_points(semantic_sections))
        ]
    else:
        tower = [
            {
                "index": index,
                "z_mm": float(point["z_mm"]),
                "diameter_mm": float(point["diameter_mm"]),
                "length_mm": point.get("length_mm"),
                "thickness_mm": float(point["thickness_mm"]),
            }
            for index, point in enumerate(ir.get("excel_facts_ir", {}).get("tower_profile_points", []))
        ]
    flanges = []
    for index, flange in enumerate(ir.get("excel_facts_ir", {}).get("flanges", []), start=1):
        flanges.append(
            {
                "index": index,
                "diameter_mm": float(flange.get("mean_diameter_mm") or flange.get("diameter_mm") or 0.0),
                "bolt_circle_mm": float(flange.get("bolt_circle_mm") or 0.0),
                "inner_diameter_mm": float(flange.get("inner_diameter_mm") or 0.0),
                "clamp_thickness_mm": float(flange.get("clamp_thickness_mm") or 0.0),
                "neck_thickness_mm": float(flange.get("neck_thickness_mm") or 0.0),
                "total_thickness_mm": float(flange.get("total_thickness_mm") or 0.0),
                "bolt_count": int(flange.get("bolt_count") or 0),
                "bolt_diameter_mm": float(flange.get("bolt_diameter_mm") or 0.0),
            }
        )
    door_info = (ir.get("semantic_ir", {}).get("door_openings") or [{}])[0] or ir.get("excel_facts_ir", {}).get("door") or {}
    if door_info.get("enabled"):
        door = {
            "X2": door_info.get("width_mm"),
            "Y2": door_info.get("height_mm"),
            "H": door_info.get("bottom_z_mm"),
            "X1": door_info.get("outer_width_mm") or door_info.get("width_mm"),
            "Y1": door_info.get("outer_height_mm") or door_info.get("height_mm"),
            "shape": door_info.get("shape", "ellipse"),
            "theta_center_deg": door_info.get("theta_center_deg", 0.0),
            "W": door_info.get("frame_thickness_mm"),
            "reinforcement_enabled": door_info.get("reinforcement_enabled", door_info.get("local_reinforcement", False)),
            "reinforcement_material": door_info.get("reinforcement_material", "Q355ND"),
        }
    else:
        door = {}
    numeric_keys = {"X2", "Y2", "H", "X1", "Y1", "W", "theta_center_deg"}
    normalized_door: dict[str, Any] = {}
    for key, value in door.items():
        if value is None:
            continue
        normalized_door[key] = float(value) if key in numeric_keys else value
    return tower, flanges, normalized_door


def build_preview_mesh(tower: list[dict[str, Any]], door: dict[str, Any], mesh_cfg: dict[str, Any]) -> tuple[dict[int, tuple[float, float, float]], dict[int, list[int]], set[int]]:
    result = build_windtower_quad_mesh(tower_sections=tower, door=door, controls=mesh_cfg)
    return result.nodes, result.elements, set(result.sets.get("DOOR_OPENING_NODES", []))


def build_gmsh_mesh_result(state: ProjectState, ir: dict[str, Any] | None = None, write_artifacts: bool = False) -> dict[str, Any]:
    ir = ir or state.ir
    tower, flanges, door = geometry_from_ir(ir)
    mesh_cfg = dict(ir["mesh_and_set_ir"])
    msh_path = state.root / "mesh" / f"{state.project_id}_windtower_quad_gmsh_v1.msh" if write_artifacts else None
    result = build_windtower_quad_mesh(tower_sections=tower, door=door, flanges=flanges, controls=mesh_cfg, msh_path=msh_path)
    artifacts = dict(result.artifacts)
    if write_artifacts:
        mesh_dir = state.root / "mesh"
        mesh_dir.mkdir(parents=True, exist_ok=True)
        inp_path = mesh_dir / f"{state.project_id}_windtower_mesh_only.inp"
        write_mesh_inp(inp_path, result, normalize_tower_points(tower), part_name=mesh_cfg.get("naming_rules", {}).get("part", "WINDTOWER_SHELL_MESH"))
        artifacts["mesh_inp"] = str(inp_path)
        state.artifacts[inp_path.name] = inp_path
        if msh_path and msh_path.exists():
            state.artifacts[msh_path.name] = msh_path
    return mesh_result_response(result, artifacts)


def mesh_result_response(result: Any, artifacts: dict[str, str] | None = None) -> dict[str, Any]:
    quality = result.quality_report
    return {
        "strategy": result.strategy,
        "mesh_stats": {
            "nodes": len(result.nodes),
            "s4r_elements": len(result.elements),
            "door_boundary_nodes": len(result.sets.get("DOOR_OPENING_NODES", [])),
            "door_patch_elements": len(result.sets.get("DOOR_PATCH_ELEMENTS", [])),
            "door_neighbor_elements": len(result.sets.get("DOOR_NEIGHBOR_ELEMENTS", [])),
            "windward_spos_elements": len(result.sets.get("WINDWARD_SPOS_ELEMENTS", [])),
            "windward_sneg_elements": len(result.sets.get("WINDWARD_SNEG_ELEMENTS", [])),
            "element_family": "S4R shell",
        },
        "set_stats": [stat(name, "NSET" if name.endswith("NODES") or name == "BASE_FIXED" else "ELSET", len(values)) for name, values in result.sets.items() if isinstance(values, list)],
        "surface_stats": [{"name": name, "type": "SURFACE", "count": len(data.get("SPOS", [])) + len(data.get("SNEG", [])), "status": "ok" if data.get("SPOS") or data.get("SNEG") else "empty"} for name, data in result.surfaces.items()],
        "quality_report": {
            "status": quality.status,
            "blocked": quality.blocked,
            "errors": quality.errors,
            "warnings": quality.warnings,
            "metrics": quality.metrics,
        },
        "artifacts": artifacts or result.artifacts,
        "metadata": result.metadata,
    }


def mesh_preview(state: ProjectState, ir: dict[str, Any] | None = None) -> dict[str, Any]:
    ir = ir or state.ir
    tower, flanges, door = geometry_from_ir(ir)
    mesh_cfg = ir["mesh_and_set_ir"]
    nodes, elements, door_nodes = build_preview_mesh(tower, door, mesh_cfg)
    wind_spos, wind_sneg = windward_surface_groups(nodes, elements)
    min_z = min(xyz[2] for xyz in nodes.values())
    max_z = max(xyz[2] for xyz in nodes.values())
    base_nodes = [nid for nid, xyz in nodes.items() if abs(xyz[2] - min_z) < 1.0e-6]
    top_nodes = [nid for nid, xyz in nodes.items() if abs(xyz[2] - max_z) < 1.0e-6]
    flange_ranges = find_flange_ranges(tower, flanges)
    flange_count = 0
    for _flange, z0, z1 in flange_ranges:
        flange_count += sum(1 for conn in elements.values() if z0 <= element_center(nodes, conn)[2] <= z1)
    set_stats = [
        stat("BASE_FIXED", "NSET", len(base_nodes)),
        stat("TOP_LOAD_NODES", "NSET", len(top_nodes)),
        stat("DOOR_OPENING_NODES", "NSET", len(door_nodes)),
        stat("WINDWARD_SPOS_ELEMENTS", "ELSET", len(wind_spos)),
        stat("WINDWARD_SNEG_ELEMENTS", "ELSET", len(wind_sneg)),
        stat("WINDWARD_SURFACE", "SURFACE", len(set(wind_spos) | set(wind_sneg))),
        stat("FLANGE_ELEMENTS", "ELSET", flange_count),
    ]
    return {
        "mesh_stats": {
            "nodes": len(nodes),
            "s4r_elements": len(elements),
            "door_boundary_nodes": len(door_nodes),
            "windward_spos_elements": len(wind_spos),
            "windward_sneg_elements": len(wind_sneg),
            "global_size_mm": mesh_cfg["global_size_mm"],
            "door_size_mm": mesh_cfg["door_size_mm"],
            "flange_size_mm": mesh_cfg["flange_size_mm"],
            "element_family": "S4R shell",
        },
        "set_stats": set_stats,
        "preview": {
            "height_mm": max_z,
            "base_radius_mm": 0.5 * tower[0]["diameter_mm"],
            "top_radius_mm": 0.5 * tower[-1]["diameter_mm"],
            "door": {"z_mm": door.get("H", 0.0), "width_mm": door.get("X2", 0.0), "height_mm": door.get("Y2", 0.0)},
            "tower_sections": [
                {
                    "z_bottom_mm": float(a["z_mm"]),
                    "z_top_mm": float(b["z_mm"]),
                    "radius_bottom_mm": 0.5 * float(a["diameter_mm"]),
                    "radius_top_mm": 0.5 * float(b["diameter_mm"]),
                }
                for a, b in zip(tower[:-1], tower[1:])
            ],
            "flanges": [
                {"elevation_mm": 0.5 * (z0 + z1), "radius_mm": 0.5 * float(flange["diameter_mm"]), "label": f"FLANGE_{i:02d}"}
                for i, (flange, z0, z1) in enumerate(flange_ranges, start=1)
            ],
        },
        "warnings": ["L3 相互作用未启用；当前网格目标为 L1/L2 可计算壳模型。"] if ir["modeling_decision_ir"]["level"] == "L3" else [],
    }


def build_inp(state: ProjectState, ir: dict[str, Any] | None = None) -> dict[str, Any]:
    ir = ir or state.ir
    missing = missing_required_tasks(ir)
    if missing:
        return {
            "artifact_name": "",
            "inp_text": "",
            "report": {
                "status": "blocked",
                "errors": [
                    {"category": "unconfirmed_task", "message": f"{task['label']}: {task['message']}"}
                    for task in missing
                ],
                "warnings": ["请先通过用户建模目标/补充约束或表单确认 L1/L2 必需建模任务。"],
                "checks": {f"task_{task['key']}": False for task in missing},
                "stats": {},
            },
        }
    tower, flanges, door = geometry_from_ir(ir)
    mesh_cfg = ir["mesh_and_set_ir"]
    nodes, elements, _door_nodes = build_preview_mesh(tower, door, mesh_cfg)
    flange_ranges = find_flange_ranges(tower, flanges)
    segments = segments_from_element_centers(nodes, elements, tower)
    out = state.root / f"{state.project_id}_windtower_shell_local_refined_analysis.inp"
    material = ir["semantic_ir"]["materials"][0]
    loads_inventory = ir["semantic_ir"].get("loads_inventory", {})
    naming = ir["mesh_and_set_ir"].get("naming_rules", {})
    part_name = naming.get("part") or ir["abaqus_assembly_ir"].get("part", {}).get("name", "WINDTOWER_SHELL_V3")
    instance_name = naming.get("instance") or f"{part_name}-1"
    wind_pressure = float(loads_inventory.get("wind_pressure_mpa") or 0.0012)
    top_total_force = float(loads_inventory.get("top_total_force_n") or 800000.0)
    stats = write_generated_analysis_inp(
        out,
        part_name,
        instance_name,
        nodes,
        elements,
        segments,
        tower,
        flange_ranges,
        material_name=material.get("name") or "STEEL",
        elastic_modulus=float(material.get("elastic_modulus_mpa") or 206000.0),
        poisson_ratio=float(material.get("poisson_ratio") if material.get("poisson_ratio") is not None else 0.3),
        top_total_force_n=top_total_force,
        wind_pressure=wind_pressure,
    )
    checks = static_check(out, stats["top_total_force"], expected_wind_pressure=wind_pressure)
    report = build_static_report(out, checks, stats)
    report_path = state.root / "build_report.json"
    write_json(report_path, report)
    state.artifacts[out.name] = out
    state.artifacts["build_report.json"] = report_path
    state.artifacts["ir.json"] = state.root / "ir.json"
    return {
        "artifact_name": out.name,
        "inp_text": out.read_text(encoding="utf-8"),
        "report": {
            "status": "ready" if all(checks.values()) else "blocked",
            "errors": [{"category": "inp_static", "message": key} for key, ok in checks.items() if not ok],
            "warnings": [] if all(checks.values()) else ["静态检查存在失败项，请修复后再导入 Abaqus。"],
            "checks": checks,
            "stats": stats,
        },
    }


def datacheck(state: ProjectState, settings: Settings, run_abaqus: bool = True) -> dict[str, Any]:
    build = build_inp(state)
    if build["report"]["status"] == "blocked":
        return {
            "status": "blocked",
            "categories": [
                {"category": error["category"], "severity": "fatal", "message": error["message"]}
                for error in build["report"]["errors"]
            ],
            "diagnostic_hints": [],
        }
    checks = build["report"]["checks"]
    categories = [
        {"category": "inp_static", "severity": "info" if ok else "fatal", "message": f"{key}: {'PASS' if ok else 'FAIL'}"}
        for key, ok in checks.items()
    ]
    abaqus_available = Path(settings.abaqus_cmd).exists() or shutil.which(settings.abaqus_cmd) is not None
    if run_abaqus and abaqus_available and build["report"]["status"] == "ready":
        inp_path = state.artifacts[build["artifact_name"]]
        job = f"{state.project_id}_datacheck"
        result = subprocess.run(
            [settings.abaqus_cmd, f"job={job}", f"input={inp_path}", "datacheck", "interactive"],
            cwd=state.root,
            capture_output=True,
            text=True,
            timeout=300,
        )
        log_text = (result.stdout or "") + "\n" + (result.stderr or "")
        msg_path = state.root / f"{job}.msg"
        dat_path = state.root / f"{job}.dat"
        if msg_path.exists():
            log_text += "\n" + msg_path.read_text(errors="ignore")
        if dat_path.exists():
            log_text += "\n" + dat_path.read_text(errors="ignore")
        has_error = bool(re.search(r"\b(ERROR|Abaqus Error)\b", log_text, re.IGNORECASE)) or result.returncode != 0
        categories.append({"category": "abaqus_datacheck", "severity": "fatal" if has_error else "info", "message": "Abaqus datacheck failed." if has_error else "Abaqus datacheck completed without detected errors."})
        (state.root / "datacheck.log").write_text(log_text, encoding="utf-8", errors="ignore")
        state.artifacts["datacheck.log"] = state.root / "datacheck.log"
    else:
        message = "未检测到 Abaqus 命令，已完成静态检查。" if not abaqus_available else "静态检查未通过，跳过 Abaqus datacheck。"
        categories.append({"category": "abaqus_datacheck", "severity": "warning", "message": message})
    status = "blocked" if any(item["severity"] == "fatal" for item in categories) else ("warning" if any(item["severity"] == "warning" for item in categories) else "pass")
    return {"status": status, "categories": categories, "diagnostic_hints": []}


def build_static_report(path: Path, checks: dict[str, bool], stats: dict[str, Any]) -> dict[str, Any]:
    text = path.read_text(encoding="utf-8")
    keyword_checks = {
        "has_part": "*Part" in text,
        "has_s4r": "*Element, type=S4R" in text,
        "has_shell_section": "*Shell Section" in text,
        "has_assembly": "*Assembly" in text,
        "has_surface": "*Surface, type=ELEMENT, name=WINDWARD_SURFACE" in text,
        "has_material": bool(re.search(r"^\*Material,\s*name=", text, flags=re.MULTILINE | re.IGNORECASE)),
        "has_two_steps": text.count("*Step,") == 2,
        "has_boundary": "*Boundary" in text,
        "has_cload": "*Cload" in text,
        "has_dsload": "*Dsload" in text,
    }
    return {"inp": str(path), "checks": checks, "keyword_checks": keyword_checks, "stats": stats}


def missing_required_tasks(ir: dict[str, Any]) -> list[dict[str, Any]]:
    level = ir.get("modeling_decision_ir", {}).get("level", "L2")
    required = {"materials", "sections", "elements", "mesh", "steps", "loads", "boundary"}
    if level == "L3":
        required.add("interactions")
    missing = []
    for task in ir.get("modeling_tasks", []):
        if task.get("key") in required and task.get("status") not in {"confirmed", "ready"}:
            missing.append(task)
    return missing


def get_state(project_id: str) -> ProjectState:
    if project_id not in PROJECTS:
        root = STORAGE_ROOT / project_id
        ir_path = root / "ir.json"
        if not ir_path.exists():
            raise KeyError(f"Project not found: {project_id}")
        ir = json.loads(ir_path.read_text(encoding="utf-8"))
        raw = ir.get("raw_excel_ir") or {}
        stored_name = raw.get("source_file_storage")
        excel_path = root / str(stored_name) if stored_name else _find_project_source_file(root)
        state = ProjectState(
            project_id=project_id,
            root=root,
            excel_path=excel_path,
            ir=ir,
            artifacts=ARTIFACT_REPOSITORY.load(project_id, root),
        )
        PROJECTS[project_id] = state
    return PROJECTS[project_id]


def _find_project_source_file(root: Path) -> Path:
    ignored = {"ir.json", "workflow_manifest.json"}
    candidates = [
        path
        for path in root.iterdir()
        if path.is_file() and path.name not in ignored and not path.name.startswith(".")
    ]
    if candidates:
        return sorted(candidates, key=lambda path: path.name.lower())[0]
    # Historical projects may no longer retain the upload.  Keep a stable path
    # so read-only project and solver-status APIs remain available after restart.
    return root / "source_missing"


def stat(name: str, kind: str, count: int) -> dict[str, Any]:
    return {"name": name, "type": kind, "count": count, "status": "ok" if count > 0 else "empty"}


def door_summary(ir: dict[str, Any]) -> str:
    door = ir["semantic_ir"]["door_openings"][0]
    if not door["enabled"]:
        return "missing"
    return f"{door['shape']} {door['width_mm']} x {door['height_mm']} mm, bottom z = {door['bottom_z_mm']} mm"


def sanitize_filename(name: str) -> str:
    original = Path(name).name
    suffix = Path(original).suffix
    stem = Path(original).stem or "upload"
    ascii_stem = unicodedata.normalize("NFKD", stem).encode("ascii", "ignore").decode("ascii")
    ascii_stem = re.sub(r'[<>:"/\\\\|?*\s]+', "_", ascii_stem).strip("._")
    if not ascii_stem:
        ascii_stem = "upload"
    digest = hashlib.sha256(original.encode("utf-8")).hexdigest()[:10]
    safe_suffix = re.sub(r'[^A-Za-z0-9.]', "", suffix)[:16]
    return f"{ascii_stem[:60]}-{digest}{safe_suffix}"


def write_json(path: Path, data: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.{uuid.uuid4().hex}.tmp")
    temporary.write_text(
        json.dumps(data, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    temporary.replace(path)
