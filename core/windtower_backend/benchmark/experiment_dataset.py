from __future__ import annotations

import csv
import hashlib
import json
import os
from copy import deepcopy
from pathlib import Path
from typing import Any

from PIL import Image, ImageDraw
from docx import Document
from pptx import Presentation

from ..config import EXAMPLE_ROOT, ROOT
from .dataset import DEFAULT_BENCHMARK_ROOT, DISCLAIMER, generate_llm_benchmark


CROSS_FORMATS = (
    "json",
    "csv",
    "txt",
    "md",
    "docx",
    "pptx",
    "text_pdf",
    "scan_pdf",
    "png",
)

DATASET_VERSION = "1.0.0"
GENERATION_SEED = 20260729


def generate_experiment_datasets(
    root: Path | None = None,
    *,
    case_ids: tuple[str, ...] | None = None,
    cross_format_case_ids: tuple[str, ...] = ("S01", "S05", "S10"),
    include_internal_engineering: bool = False,
) -> dict[str, Any]:
    """Generate versioned G1-G6 manifests without calling LLM or Abaqus."""

    root = Path(root or DEFAULT_BENCHMARK_ROOT)
    base_manifest_path = root / "benchmark_manifest.json"
    if not base_manifest_path.exists():
        generate_llm_benchmark(root)
    base = _read_json(base_manifest_path)
    all_cases = {
        item["case_id"]: _resolve_case_paths(root, item)
        for item in base["cases"]
    }
    selected = {
        case_id: case
        for case_id, case in all_cases.items()
        if case_ids is None or case_id in set(case_ids)
    }
    if case_ids:
        missing = sorted(set(case_ids) - set(selected))
        if missing:
            raise ValueError(f"Unknown benchmark case IDs: {missing}")

    experiment_root = root / "experiment_datasets"
    cross_root = experiment_root / "G3_cross_format"
    fault_root = experiment_root / "G4_fault_injection"
    cross_cases: list[dict[str, Any]] = []
    cross_format_set = set(cross_format_case_ids)
    for case_id in sorted(cross_format_set):
        case = all_cases.get(case_id)
        if case is None:
            continue
        truth = _read_json(Path(case["ground_truth"]))
        case_root = cross_root / case_id
        artifacts = _write_cross_format_case(case_root, truth)
        cross_cases.append(
            {
                "case_id": case_id,
                "truth": case["ground_truth"],
                "truth_sha256": _sha256(Path(case["ground_truth"])),
                "artifacts": artifacts,
            }
        )

    fault_cases: list[dict[str, Any]] = []
    if selected:
        anchor_id = sorted(selected)[0]
        anchor = selected[anchor_id]
        fault_truth = _read_json(Path(anchor["ground_truth"]))
        fault_fem = fault_truth.setdefault("fem_settings", {})
        if float(fault_fem.get("loads.top_mass_tonne") or 0) > 0:
            fault_fem.setdefault("loads.center_of_mass_offset_mm", 0.0)
        fault_cases = _write_fault_cases(
            fault_root,
            fault_truth,
            source_case_id=anchor_id,
        )

    g1_inputs = _reference_inputs(
        base,
        include_internal_engineering=include_internal_engineering,
    )
    modification_cases = _modification_cases()
    solver_cases = _solver_validation_cases(g1_inputs)
    payload = {
        "schema_version": "wind_tower_experiment_dataset_manifest_v2",
        "dataset_version": DATASET_VERSION,
        "generation_seed": GENERATION_SEED,
        "declaration": DISCLAIMER,
        "data_policy": {
            "engineering_references": "internal_checksum_reference_only",
            "public_artifacts": "deidentified_synthetic",
            "immutable_release": True,
            "safety_claim": "not_for_engineering_design_or_safety_assessment",
        },
        "groups": {
            "G1": {
                "name": "golden_engineering_cases",
                "purpose": "compiler, mesh, Datacheck and result baselines",
                "cases": g1_inputs,
            },
            "G2": {
                "name": "heterogeneous_expression_cases",
                "purpose": "layout, mixed-unit and LLM mapping robustness",
                "cases": [selected[key] for key in sorted(selected)],
            },
            "G3": {
                "name": "cross_format_same_truth",
                "purpose": "compare carriers while holding engineering truth fixed",
                "formats": list(CROSS_FORMATS),
                "cases": cross_cases,
            },
            "G4": {
                "name": "fault_injection",
                "purpose": "missing, invalid, unsupported, duplicate and corrupt input handling",
                "cases": fault_cases,
            },
            "G5": {
                "name": "model_modification_equivalence",
                "purpose": "compare registered ParameterChange and natural-language modification paths",
                "base_case": "synthetic_standard_tower.xlsx",
                "case_count": len(modification_cases),
                "single_parameter_cases": sum(
                    len(item["changes"]) == 1 for item in modification_cases
                ),
                "compound_cases": sum(
                    len(item["changes"]) > 1 for item in modification_cases
                ),
                "natural_language_variants_per_case": 3,
                "cases": modification_cases,
            },
            "G6": {
                "name": "solver_validation_matrix",
                "purpose": "INP static checks, Abaqus Datacheck, solve and ODB sanity checks",
                **solver_cases,
            },
        },
    }
    manifest_path = experiment_root / "experiment_dataset_manifest.json"
    _write_json(manifest_path, payload)
    payload["manifest"] = str(manifest_path)
    payload["manifest_sha256"] = _sha256(manifest_path)
    return payload


def _write_cross_format_case(root: Path, truth: dict[str, Any]) -> list[dict[str, Any]]:
    root.mkdir(parents=True, exist_ok=True)
    artifacts: list[dict[str, Any]] = []
    canonical_text = _engineering_text(truth)

    json_path = root / "parameters.json"
    _write_json(json_path, {"standard_input_package": truth})
    artifacts.append(_artifact(json_path, "structured", False))

    csv_path = root / "engineering_parameters.csv"
    with csv_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=("record_type", "path", "value", "unit"),
        )
        writer.writeheader()
        writer.writerows(_parameter_rows(truth))
    artifacts.append(_artifact(csv_path, "structured", False))

    txt_path = root / "engineering_parameters.txt"
    txt_path.write_text(canonical_text, encoding="utf-8")
    artifacts.append(_artifact(txt_path, "text", False))

    md_path = root / "engineering_parameters.md"
    md_path.write_text("# Synthetic wind-tower parameters\n\n" + canonical_text, encoding="utf-8")
    artifacts.append(_artifact(md_path, "text", False))

    docx_path = root / "engineering_parameters.docx"
    document = Document()
    document.add_heading("Synthetic wind-tower parameters", level=1)
    for line in canonical_text.splitlines():
        document.add_paragraph(line)
    table = document.add_table(rows=1, cols=6)
    headings = (
        "segment_id",
        "z_bottom_mm",
        "z_top_mm",
        "diameter_bottom_mm",
        "diameter_top_mm",
        "thickness_mm",
    )
    for index, heading in enumerate(headings):
        table.rows[0].cells[index].text = heading
    for segment in truth.get("tower_segments") or []:
        cells = table.add_row().cells
        for index, heading in enumerate(headings):
            cells[index].text = str(segment.get(heading, ""))
    document.save(docx_path)
    artifacts.append(_artifact(docx_path, "document", False))

    pptx_path = root / "engineering_parameters.pptx"
    presentation = Presentation()
    for title, body in (
        ("Model intent", "\n".join(canonical_text.splitlines()[:12])),
        ("Tower segments", "\n".join(_segment_lines(truth))),
        ("Loads and analysis", "\n".join(canonical_text.splitlines()[12:])),
    ):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    presentation.save(pptx_path)
    artifacts.append(_artifact(pptx_path, "presentation", False))

    text_pdf_path = root / "engineering_parameters_text.pdf"
    _write_text_pdf(text_pdf_path, canonical_text)
    artifacts.append(_artifact(text_pdf_path, "text_pdf", False))

    image_path = root / "engineering_parameters_scan.png"
    image = _render_text_image(canonical_text)
    image.save(image_path, format="PNG")
    artifacts.append(_artifact(image_path, "image", True))

    scan_pdf_path = root / "engineering_parameters_scan.pdf"
    image.convert("RGB").save(scan_pdf_path, format="PDF", resolution=144.0)
    artifacts.append(_artifact(scan_pdf_path, "scanned_pdf", True))
    return artifacts


def _write_fault_cases(
    root: Path,
    truth: dict[str, Any],
    *,
    source_case_id: str,
) -> list[dict[str, Any]]:
    root.mkdir(parents=True, exist_ok=True)
    cases: list[dict[str, Any]] = []

    def add(
        fault_id: str,
        payload: dict[str, Any] | bytes,
        expected: str,
        *,
        suffix: str = ".json",
    ) -> None:
        path = root / f"{fault_id}{suffix}"
        if isinstance(payload, bytes):
            path.write_bytes(payload)
        else:
            _write_json(path, {"standard_input_package": payload})
        cases.append(
            {
                "fault_id": fault_id,
                "source_case_id": source_case_id,
                "path": str(path),
                "sha256": _sha256(path),
                "expected_behavior": expected,
            }
        )

    missing_thickness = deepcopy(truth)
    missing_thickness["tower_segments"][0].pop("thickness_mm", None)
    add("missing_required_thickness", missing_thickness, "block_canonical_or_compilation")

    invalid_geometry = deepcopy(truth)
    invalid_geometry["tower_segments"][0]["thickness_mm"] = -10
    add("invalid_negative_thickness", invalid_geometry, "block_engineering_validation")

    invalid_direction = deepcopy(truth)
    invalid_direction["fem_settings"]["loads.top_force_direction"] = "north-east"
    add("invalid_load_direction", invalid_direction, "require_resolution_or_block_static")

    unknown_feature = deepcopy(truth)
    unknown_feature["door_reinforcement"] = {"ring_plate_thickness_mm": 30}
    add("unknown_reinforcement_feature", unknown_feature, "isolate_as_unsupported")

    optional_removed = deepcopy(truth)
    optional_removed["fem_settings"].pop("outputs.field", None)
    add("missing_defaultable_output", optional_removed, "default_or_warn_without_crash")

    duplicate = deepcopy(truth)
    add("duplicate_source_a", duplicate, "deduplicate_by_content_hash")
    add("duplicate_source_b", duplicate, "deduplicate_by_content_hash")

    add("corrupt_document", b"not-a-valid-office-or-pdf-file\x00\x01", "structured_file_error", suffix=".pdf")

    missing_elasticity = deepcopy(truth)
    missing_elasticity["materials"][0].pop("elastic_modulus_mpa", None)
    add("missing_required_elastic_modulus", missing_elasticity, "block_active_analysis")

    missing_density = deepcopy(truth)
    missing_density["materials"][0].pop("density_tonne_mm3", None)
    add("missing_required_density", missing_density, "block_modal_or_mass_validation")

    missing_boundary = deepcopy(truth)
    missing_boundary["fem_settings"].pop("boundary.name", None)
    add("missing_required_boundary", missing_boundary, "block_solver")

    missing_modal_modes = deepcopy(truth)
    missing_modal_modes.setdefault("analysis", {})["modal"] = True
    missing_modal_modes["analysis"].pop("modal_modes", None)
    missing_modal_modes["fem_settings"].pop("modal.modes", None)
    add("missing_conditional_modal_modes", missing_modal_modes, "block_modal_only")

    inactive_modal = deepcopy(truth)
    inactive_modal.setdefault("analysis", {})["modal"] = False
    inactive_modal["analysis"].pop("modal_modes", None)
    inactive_modal["fem_settings"].pop("modal.modes", None)
    add("missing_inactive_modal_modes", inactive_modal, "do_not_report_as_missing")

    disabled_door = deepcopy(truth)
    disabled_door["door_opening"] = {"enabled": False}
    add("missing_inactive_door_dimensions", disabled_door, "disable_door_without_blocking")

    noncontiguous = deepcopy(truth)
    noncontiguous["tower_segments"][1]["z_bottom_mm"] += 250
    add("noncontiguous_tower_segments", noncontiguous, "block_engineering_validation")

    negative_diameter = deepcopy(truth)
    negative_diameter["tower_segments"][0]["diameter_bottom_mm"] = -100
    add("invalid_negative_diameter", negative_diameter, "block_engineering_validation")

    invalid_poisson = deepcopy(truth)
    invalid_poisson["materials"][0]["poisson_ratio"] = 0.7
    add("invalid_poisson_ratio", invalid_poisson, "block_material_validation")

    invalid_mesh = deepcopy(truth)
    invalid_mesh["fem_settings"]["mesh.global_size_mm"] = 0
    add("invalid_zero_mesh_size", invalid_mesh, "block_mesh_generation")

    door_outside = deepcopy(truth)
    door_outside["door_opening"]["z_bottom_mm"] = (
        truth["tower_segments"][-1]["z_top_mm"] + 1000
    )
    add("invalid_door_outside_tower", door_outside, "block_geometry_validation")

    ambiguous_unit = deepcopy(truth)
    ambiguous_unit["fem_settings"]["loads.top_total_force_n"] = "250 unknown_force_unit"
    add("ambiguous_load_unit", ambiguous_unit, "require_resolution")

    direction_alias = deepcopy(truth)
    direction_alias["fem_settings"]["loads.top_force_direction"] = 1
    add("normalizable_direction_alias", direction_alias, "normalize_to_plus_x_with_trace")

    conflicting_force = deepcopy(truth)
    conflicting_force["conflicting_candidates"] = [
        {"path": "loads.top_horizontal_force_n", "value": 200000, "source": "file_a"},
        {"path": "loads.top_horizontal_force_n", "value": 300000, "source": "file_b"},
    ]
    add("conflicting_top_force_sources", conflicting_force, "detect_conflict_and_require_resolution")

    unknown_solver_keyword = deepcopy(truth)
    unknown_solver_keyword["solver_keyword"] = "*UNREGISTERED USER KEYWORD"
    add("unknown_solver_keyword", unknown_solver_keyword, "isolate_as_unsupported")

    add(
        "truncated_json_document",
        b'{"standard_input_package":{"tower_segments":[',
        "structured_file_error",
        suffix=".json",
    )
    return cases


def _modification_cases() -> list[dict[str, Any]]:
    cases = [
        _modification("M01", "geometry", [("geometry.tower_sections.0.thickness_mm", 40.0)], "第一段塔筒壁厚改为40 mm", "把最底下一段钢板厚度调成40毫米", "Set the first tower course thickness to 40 mm."),
        _modification("M02", "geometry", [("geometry.door.width_mm", 1200.0)], "门洞净宽调整为1200 mm", "把门洞加宽到1.2米", "Set the door opening width to 1200 mm."),
        _modification("M03", "material", [("material.elastic_modulus_mpa", 205000.0)], "材料弹性模量设为205000 MPa", "钢材的E改成205 GPa", "Set Young's modulus to 205000 MPa."),
        _modification("M04", "material", [("material.yield_strength_mpa", 345.0)], "材料屈服强度设为345 MPa", "换成屈服强度345兆帕的钢材", "Set material yield strength to 345 MPa."),
        _modification("M05", "mesh", [("mesh.global_size_mm", 500.0)], "全局网格尺寸调整为500 mm", "整体网格细一点，用0.5米", "Set the global mesh target size to 500 mm."),
        _modification("M06", "mesh", [("mesh.door_size_mm", 100.0)], "门洞局部网格尺寸调整为100 mm", "门洞附近网格用10厘米", "Set the door local mesh size to 100 mm."),
        _modification("M07", "loads", [("loads.top_mass_tonne", 900.0)], "顶部等效质量调整为900 tonne", "机舱和叶轮等效质量改成900吨", "Set the top equivalent mass to 900 tonnes."),
        _modification("M08", "loads", [("loads.top_force_direction", "+X")], "顶部水平力方向调整为+X", "塔顶推力朝X正方向", "Set the top horizontal force direction to +X."),
        _modification("M09", "analysis", [("analysis.modal_modes", 15)], "模态提取阶数设为15", "计算前15阶振型", "Extract 15 modal modes."),
        _modification("M10", "analysis", [("analysis.static_enabled", True)], "启用静力分析", "这次要算静力", "Enable static analysis."),
        _modification("M11", "door", [("geometry.door.height_mm", 2400.0)], "门洞高度调整为2400 mm", "门洞做成2.4米高", "Set the door opening height to 2400 mm."),
        _modification("M12", "door", [("geometry.door.theta_center_deg", 90.0)], "门洞环向角度调整为90 deg", "把门洞转到90度方向", "Set the door azimuth to 90 degrees."),
        _modification("M13", "geometry", [("geometry.tower_sections.0.z_top_mm", 12000.0), ("geometry.tower_sections.1.z_bottom_mm", 12000.0)], "第一、二塔段交界标高统一调整为12000 mm", "把前两段的接缝一起移到12米，保持连续", "Move the shared boundary of tower courses 1 and 2 to 12000 mm."),
        _modification("M14", "door", [("geometry.door.width_mm", 1150.0), ("geometry.door.height_mm", 2350.0), ("geometry.door.bottom_z_mm", 600.0)], "门洞宽1150 mm、高2350 mm、底标高600 mm", "门洞改成1.15米宽、2.35米高，底边离地0.6米", "Set door width 1150 mm, height 2350 mm and bottom elevation 600 mm."),
        _modification("M15", "material", [("material.elastic_modulus_mpa", 206000.0), ("material.density", 7.85e-9), ("material.yield_strength_mpa", 355.0)], "材料参数调整为E=206000 MPa、密度7.85e-9 tonne/mm3、屈服355 MPa", "钢材用E206 GPa、密度7850千克每立方米、屈服355兆帕", "Use E=206000 MPa, density=7.85e-9 tonne/mm3 and yield strength=355 MPa."),
        _modification("M16", "mesh", [("mesh.global_size_mm", 550.0), ("mesh.door_size_mm", 90.0)], "全局网格550 mm，门洞局部网格90 mm", "整体网格0.55米，门洞附近加密到9厘米", "Use 550 mm globally and 90 mm around the door."),
        _modification("M17", "loads", [("loads.top_horizontal_force_n", 250000.0), ("loads.top_force_direction", "-Y"), ("loads.wind_pressure_mpa", 0.0005)], "顶部水平力250000 N、方向-Y、风压0.0005 MPa", "塔顶向Y负方向推250千牛，同时风压设0.5千帕", "Set top force to 250000 N in -Y and wind pressure to 0.0005 MPa."),
        _modification("M18", "analysis", [("analysis.static_enabled", True), ("analysis.modal_enabled", True), ("analysis.modal_modes", 20)], "启用静力与模态分析并提取20阶模态", "静力和振动都算，振型取20阶", "Enable static and modal analyses and extract 20 modes."),
    ]
    return cases


def _modification(
    case_id: str,
    group: str,
    changes: list[tuple[str, Any]],
    professional: str,
    plain: str,
    english: str,
) -> dict[str, Any]:
    return {
        "case_id": case_id,
        "group": group,
        "changes": [{"key": key, "value": value} for key, value in changes],
        "natural_language": {
            "professional_zh": professional,
            "plain_zh": plain,
            "engineering_en": english,
        },
        "expected": {
            "parent_immutable": True,
            "stale_hash_rejected": True,
            "normalized_changes_equivalent": True,
            "canonical_equivalent_when_changes_match": True,
            "build_equivalent_when_changes_match": True,
            "inp_semantically_equivalent_when_changes_match": True,
        },
    }


def _solver_validation_cases(g1_inputs: list[dict[str, Any]]) -> dict[str, Any]:
    filenames = [item["filename"] for item in g1_inputs]
    standard_case = next(
        (name for name in filenames if name == "synthetic_standard_tower.xlsx"),
        filenames[0],
    )
    engineering = [name for name in filenames if name != standard_case]
    representative = [
        engineering[0] if engineering else standard_case,
        engineering[len(engineering) // 2] if engineering else standard_case,
        standard_case,
    ]
    return {
        "datacheck": {
            "cases": filenames,
            "analyses": ["static_strength", "modal"],
            "repetitions": 3,
        },
        "solve": {
            "cases": list(dict.fromkeys(representative)),
            "analyses": ["static_strength", "modal"],
            "repetitions": 3,
        },
        "mesh_sensitivity": {
            "case": standard_case,
            "global_mesh_factors": [1.5, 1.0, 0.75],
            "door_mesh_factors": [1.5, 1.0, 0.75],
            "interpretation": "Numerical sensitivity only; not a structural safety assessment.",
        },
        "acceptance": {
            "datacheck_pass_rate": 1.0,
            "odb_valid_rate": 1.0,
            "mass_relative_error_max": 0.02,
            "reaction_balance_residual_max": 0.01,
            "repeat_scalar_relative_difference_max": 1e-6,
            "last_two_mesh_level_difference_target": 0.05,
        },
    }


def _engineering_text(truth: dict[str, Any]) -> str:
    material = (truth.get("materials") or [{}])[0]
    door = truth.get("door_opening") or {}
    fem = truth.get("fem_settings") or {}
    analysis = truth.get("analysis") or {}
    lines = [
        "Model family: steel wind-tower shell",
        "Element type: S4R",
        f"Material name: {material.get('name')}",
        f"Elastic modulus: {material.get('elastic_modulus_mpa')} MPa",
        f"Poisson ratio: {material.get('poisson_ratio')}",
        f"Density: {material.get('density_tonne_mm3')} tonne/mm^3",
        f"Yield strength: {material.get('yield_strength_mpa')} MPa",
        f"Door enabled: {door.get('enabled')}",
        f"Door width: {door.get('width_mm')} mm",
        f"Door height: {door.get('height_mm')} mm",
        f"Door bottom elevation: {door.get('z_bottom_mm')} mm",
        f"Global mesh size: {fem.get('mesh.global_size_mm')} mm",
        f"Door mesh size: {fem.get('mesh.door_size_mm')} mm",
        f"Top mass: {fem.get('loads.top_mass_tonne')} tonne",
        f"Top horizontal force: {fem.get('loads.top_total_force_n')} N",
        f"Top force direction: {fem.get('loads.top_force_direction')}",
        f"Wind pressure: {fem.get('loads.wind_pressure_mpa')} MPa",
        f"Static analysis enabled: {analysis.get('static_strength')}",
        f"Modal analysis enabled: {analysis.get('modal')}",
        f"Modal extraction count: {analysis.get('modal_modes')}",
        "Boundary condition: tower base fully fixed in DOF 1,2,3,4,5,6",
        "Tower segments:",
        *_segment_lines(truth),
    ]
    return "\n".join(lines) + "\n"


def _segment_lines(truth: dict[str, Any]) -> list[str]:
    return [
        (
            f"{item.get('segment_id', index + 1)}: "
            f"z={item.get('z_bottom_mm')}..{item.get('z_top_mm')} mm, "
            f"OD={item.get('diameter_bottom_mm')}..{item.get('diameter_top_mm')} mm, "
            f"t={item.get('thickness_mm')} mm"
        )
        for index, item in enumerate(truth.get("tower_segments") or [])
    ]


def _parameter_rows(truth: dict[str, Any]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for index, segment in enumerate(truth.get("tower_segments") or []):
        for field in (
            "z_bottom_mm",
            "z_top_mm",
            "diameter_bottom_mm",
            "diameter_top_mm",
            "thickness_mm",
        ):
            rows.append(
                {
                    "record_type": "tower_segment",
                    "path": f"tower_segments[{index}].{field}",
                    "value": segment.get(field),
                    "unit": "mm",
                }
            )
    material = (truth.get("materials") or [{}])[0]
    for field, unit in (
        ("name", ""),
        ("elastic_modulus_mpa", "MPa"),
        ("poisson_ratio", ""),
        ("density_tonne_mm3", "tonne/mm^3"),
        ("yield_strength_mpa", "MPa"),
    ):
        rows.append(
            {
                "record_type": "material",
                "path": f"materials[0].{field}",
                "value": material.get(field),
                "unit": unit,
            }
        )
    door = truth.get("door_opening") or {}
    for field, unit in (
        ("enabled", ""),
        ("width_mm", "mm"),
        ("height_mm", "mm"),
        ("z_bottom_mm", "mm"),
        ("azimuth_deg", "deg"),
    ):
        rows.append(
            {
                "record_type": "door",
                "path": f"door_opening.{field}",
                "value": door.get(field),
                "unit": unit,
            }
        )
    for key, value in sorted((truth.get("fem_settings") or {}).items()):
        rows.append(
            {
                "record_type": "fem_setting",
                "path": f"fem_settings.{key}",
                "value": json.dumps(value, ensure_ascii=False) if isinstance(value, (list, dict)) else value,
                "unit": _unit_for_parameter(key),
            }
        )
    for key, value in sorted((truth.get("analysis") or {}).items()):
        rows.append(
            {
                "record_type": "analysis",
                "path": f"analysis.{key}",
                "value": value,
                "unit": "",
            }
        )
    return rows


def _unit_for_parameter(key: str) -> str:
    if key.endswith("_mm"):
        return "mm"
    if key.endswith("_n"):
        return "N"
    if key.endswith("_mpa"):
        return "MPa"
    if key.endswith("_tonne"):
        return "tonne"
    if "density" in key:
        return "tonne/mm^3"
    return ""


def _write_text_pdf(path: Path, text: str) -> None:
    import fitz

    document = fitz.open()
    page = document.new_page(width=595, height=842)
    y = 48
    for line in text.splitlines():
        if y > 800:
            page = document.new_page(width=595, height=842)
            y = 48
        page.insert_text((40, y), line[:110], fontsize=8)
        y += 14
    document.save(path)
    document.close()


def _render_text_image(text: str) -> Image.Image:
    lines = [line[:105] for line in text.splitlines()]
    width = 1400
    line_height = 18
    height = max(900, 50 + line_height * len(lines))
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    y = 24
    for line in lines:
        draw.text((24, y), line, fill="black")
        y += line_height
    return image


def _reference_inputs(
    _base: dict[str, Any],
    *,
    include_internal_engineering: bool,
) -> list[dict[str, Any]]:
    if include_internal_engineering:
        configured_root = os.getenv("WINDTOWER_PRIVATE_REFERENCE_ROOT", "").strip()
        if not configured_root:
            raise FileNotFoundError(
                "Internal G1 requires WINDTOWER_PRIVATE_REFERENCE_ROOT to be configured"
            )
        engineering_root = Path(configured_root).expanduser()
        paths = sorted(
            path
            for pattern in ("*.xls", "*.xlsx")
            for path in engineering_root.glob(pattern)
            if path.is_file()
        )[:5]
        if not paths:
            raise FileNotFoundError(
                f"Internal G1 found no workbooks under {engineering_root}"
            )
        return [
            {
                "filename": path.name,
                "source_path": path.name,
                "sha256": _sha256(path),
                "visibility": "internal_engineering_reference",
            }
            for path in paths
        ]
    engineering_root = EXAMPLE_ROOT
    paths = [
        engineering_root / "synthetic_standard_tower.xlsx",
        engineering_root / "synthetic_mechanical_interface_FEM.xls",
    ]
    paths = [path for path in paths if path.is_file()]
    if len(paths) != 2:
        raise FileNotFoundError(f"G1 requires two synthetic reference workbooks; found {len(paths)}")
    return [
        {
            "filename": path.name,
            "source_path": path.relative_to(ROOT).as_posix(),
            "sha256": _sha256(path),
            "visibility": "public_synthetic_fixture",
        }
        for path in paths
    ]


def _artifact(path: Path, content_kind: str, requires_ocr: bool) -> dict[str, Any]:
    return {
        "format": path.suffix.lower().lstrip("."),
        "content_kind": content_kind,
        "path": str(path),
        "sha256": _sha256(path),
        "size_bytes": path.stat().st_size,
        "requires_ocr": requires_ocr,
    }


def _read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def _resolve_case_paths(root: Path, item: dict[str, Any]) -> dict[str, Any]:
    resolved = dict(item)
    for key in ("input", "ground_truth"):
        path = Path(str(item[key]))
        if not path.is_absolute():
            path = root / path
        resolved[key] = str(path.resolve())
    return resolved


def _write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    encoded = json.dumps(payload, ensure_ascii=False, indent=2, default=str)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(encoded, encoding="utf-8")
    temporary.replace(path)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()
