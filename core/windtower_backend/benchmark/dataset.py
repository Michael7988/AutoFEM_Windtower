from __future__ import annotations

import hashlib
import json
from copy import deepcopy
from pathlib import Path
from typing import Any, Callable

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from ..config import EXAMPLE_ROOT


DEFAULT_BENCHMARK_ROOT = EXAMPLE_ROOT / "llm_benchmark"
REFERENCE_ROOT = EXAMPLE_ROOT
ALPHAS = tuple(round(0.05 + 0.10 * index, 2) for index in range(10))
DISCLAIMER = "SYNTHETIC BENCHMARK / 非真实设计 / NOT FOR ENGINEERING DESIGN OR SAFETY ASSESSMENT"
STATIONS = tuple(index / 10.0 for index in range(11))
SYNTHETIC_ANCHOR_A = {
    "height_mm": 90000.0,
    "diameter_mm": [6800.0, 6500.0, 6200.0, 5900.0, 5600.0, 5300.0, 5000.0, 4700.0, 4400.0, 4100.0, 3800.0],
    "thickness_mm": [32.0, 30.0, 28.0, 26.0, 24.0, 22.0, 20.0, 18.0, 16.0, 14.0],
    "door": {"width_mm": 1100.0, "height_mm": 2100.0, "z_bottom_mm": 700.0},
    "top_force_n": 500000.0,
    "top_mass_tonne": 250.0,
    "wind_pressure_mpa": 0.0009,
}
SYNTHETIC_ANCHOR_B = {
    "height_mm": 110000.0,
    "diameter_mm": [7600.0, 7200.0, 6800.0, 6400.0, 6000.0, 5600.0, 5200.0, 4800.0, 4400.0, 4000.0, 3600.0],
    "thickness_mm": [40.0, 37.0, 34.0, 31.0, 28.0, 25.0, 22.0, 20.0, 18.0, 16.0],
    "door": {"width_mm": 1300.0, "height_mm": 2300.0, "z_bottom_mm": 900.0},
    "top_force_n": 900000.0,
    "top_mass_tonne": 400.0,
    "wind_pressure_mpa": 0.0014,
}
# Keep the public declaration UTF-8 clean even when older worktrees contain a
# damaged legacy literal above.
DISCLAIMER = (
    "SYNTHETIC BENCHMARK / 非真实设计 / "
    "NOT FOR ENGINEERING DESIGN OR SAFETY ASSESSMENT"
)
# Use Unicode escapes so the declaration is stable across Windows code pages.
DISCLAIMER = (
    "SYNTHETIC BENCHMARK / \u975e\u771f\u5b9e\u8bbe\u8ba1 / "
    "NOT FOR ENGINEERING DESIGN OR SAFETY ASSESSMENT"
)


def generate_llm_benchmark(out_root: Path | None = None) -> dict[str, Any]:
    out_root = Path(out_root or DEFAULT_BENCHMARK_ROOT)
    input_dir = out_root / "inputs"
    truth_dir = out_root / "ground_truth"
    input_dir.mkdir(parents=True, exist_ok=True)
    truth_dir.mkdir(parents=True, exist_ok=True)

    references = _reference_paths()
    reference_hashes = {path.name: _sha256(path) for path in references}
    anchor_a, anchor_b, stations = _build_anchors(references)
    cases = []
    persisted_cases = []
    writers: tuple[Callable[[Path, dict[str, Any]], None], ...] = (
        _write_s01,
        _write_s02,
        _write_s03,
        _write_s04,
        _write_s05,
        _write_s06,
        _write_s07,
        _write_s08,
        _write_s09,
        _write_s10,
    )
    for index, (alpha, writer) in enumerate(zip(ALPHAS, writers), start=1):
        case_id = f"S{index:02d}"
        truth = _interpolate_case(case_id, alpha, anchor_a, anchor_b, stations)
        truth["benchmark"] = {
            "case_id": case_id,
            "alpha": alpha,
            "declaration": DISCLAIMER,
            "data_origin": "fixed_fictional_synthetic_anchors_v1",
            "basis": "Convex interpolation between two fixed fictional tower profiles created only for software testing.",
            "reference_sha256": reference_hashes,
            "formula": "P_i=(1-alpha_i)*AnchorA+alpha_i*AnchorB; alpha_i=0.05+0.10*(i-1)",
        }
        workbook_path = input_dir / f"{case_id}_heterogeneous_wind_tower.xlsx"
        truth_path = truth_dir / f"{case_id}_ground_truth.json"
        writer(workbook_path, truth)
        truth["benchmark"]["workbook"] = workbook_path.relative_to(out_root).as_posix()
        truth["benchmark"]["workbook_sha256"] = _sha256(workbook_path)
        truth_path.write_text(json.dumps(truth, ensure_ascii=False, indent=2), encoding="utf-8")
        persisted_case = {
            "case_id": case_id,
            "alpha": alpha,
            "input": workbook_path.relative_to(out_root).as_posix(),
            "ground_truth": truth_path.relative_to(out_root).as_posix(),
            "input_sha256": truth["benchmark"]["workbook_sha256"],
            "format_challenge": _format_challenge(index),
        }
        persisted_cases.append(persisted_case)
        cases.append(
            {
                **persisted_case,
                "input": str(workbook_path),
                "ground_truth": str(truth_path),
            }
        )

    manifest = {
        "schema_version": "llm_windtower_benchmark_manifest_v1",
        "declaration": DISCLAIMER,
        "case_count": len(cases),
        "reference_files": reference_hashes,
        "alphas": list(ALPHAS),
        "data_origin": "fixed_fictional_synthetic_anchors_v1",
        "cases": persisted_cases,
    }
    manifest_path = out_root / "benchmark_manifest.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return {**manifest, "cases": cases, "manifest": str(manifest_path), "output_root": str(out_root)}


def _reference_paths() -> list[Path]:
    paths = [
        REFERENCE_ROOT / "synthetic_standard_tower.xlsx",
        REFERENCE_ROOT / "synthetic_mechanical_interface_FEM.xls",
    ]
    missing = [str(path) for path in paths if not path.is_file()]
    if missing:
        raise FileNotFoundError(f"Synthetic fixture references are missing: {missing}")
    return paths


def _build_anchors(_references: list[Path] | None = None) -> tuple[dict[str, Any], dict[str, Any], list[float]]:
    """Return fixed fictional anchors without reading engineering workbooks."""

    return deepcopy(SYNTHETIC_ANCHOR_A), deepcopy(SYNTHETIC_ANCHOR_B), list(STATIONS)


def _plain_case(normalized: dict[str, Any]) -> dict[str, Any]:
    package = normalized.get("standard_input_package") or normalized.get("standard_input_candidate") or {}
    segments = []
    for segment in package.get("tower_segments", []) or []:
        segments.append({key: float(_value(segment.get(key))) for key in ("z_bottom_mm", "z_top_mm", "diameter_bottom_mm", "diameter_top_mm", "thickness_mm")})
    door = package.get("door_opening") or {}
    fem = dict(normalized.get("fem_params") or {})
    return {
        "height_mm": max(item["z_top_mm"] for item in segments),
        "tower_segments": segments,
        "door": {key: _value(door.get(key)) for key in ("width_mm", "height_mm", "z_bottom_mm")},
        "fem_settings": fem,
    }


def _sample_profile(case: dict[str, Any], u: float, field: str) -> float:
    z = max(0.0, min(1.0, u)) * float(case["height_mm"])
    segments = case["tower_segments"]
    segment = next((item for item in segments if item["z_bottom_mm"] - 1e-9 <= z <= item["z_top_mm"] + 1e-9), segments[-1])
    if field == "thickness_mm":
        return float(segment[field])
    dz = max(segment["z_top_mm"] - segment["z_bottom_mm"], 1e-12)
    ratio = max(0.0, min(1.0, (z - segment["z_bottom_mm"]) / dz))
    return float(segment["diameter_bottom_mm"] + ratio * (segment["diameter_top_mm"] - segment["diameter_bottom_mm"]))


def _interpolate_case(case_id: str, alpha: float, anchor_a: dict[str, Any], anchor_b: dict[str, Any], stations: list[float]) -> dict[str, Any]:
    def lerp(a: float, b: float) -> float:
        return (1.0 - alpha) * float(a) + alpha * float(b)
    height = round(lerp(anchor_a["height_mm"], anchor_b["height_mm"]), 1)
    diameters = [round(lerp(a, b), 1) for a, b in zip(anchor_a["diameter_mm"], anchor_b["diameter_mm"])]
    diameters = [min(diameters[: index + 1]) for index in range(len(diameters))]
    thicknesses = [round(max(1.0, lerp(a, b)), 1) for a, b in zip(anchor_a["thickness_mm"], anchor_b["thickness_mm"])]
    segments = []
    for index in range(10):
        segments.append(
            {
                "segment_id": f"SEG-{index + 1:02d}",
                "z_bottom_mm": round(height * stations[index], 1),
                "z_top_mm": round(height * stations[index + 1], 1),
                "diameter_bottom_mm": diameters[index],
                "diameter_top_mm": diameters[index + 1],
                "thickness_mm": thicknesses[index],
            }
        )
    material = _material(alpha)
    door = {
        "enabled": True,
        "shape": "rounded_rectangle",
        "width_mm": round(lerp(anchor_a["door"]["width_mm"], anchor_b["door"]["width_mm"]), 1),
        "height_mm": round(lerp(anchor_a["door"]["height_mm"], anchor_b["door"]["height_mm"]), 1),
        "z_bottom_mm": round(lerp(anchor_a["door"]["z_bottom_mm"], anchor_b["door"]["z_bottom_mm"]), 1),
        "corner_radius_mm": 250.0,
        "azimuth_deg": 0.0,
    }
    fem = {
        "material.name": material["name"],
        "material.elastic_modulus_mpa": material["elastic_modulus_mpa"],
        "material.poisson_ratio": material["poisson_ratio"],
        "material.density_tonne_mm3": material["density_tonne_mm3"],
        "material.yield_strength_mpa": material["yield_strength_mpa"],
        "section.type": "Shell Section",
        "section.material": material["name"],
        "section.thickness_source": "tower_segments",
        "element.element_type": "S4R",
        "element.part_name": f"SYNTHETIC_{case_id}_TOWER",
        "element.instance_name": f"SYNTHETIC_{case_id}_TOWER-1",
        "mesh.strategy": "V3_CONFORMING_DOOR_MESH",
        "mesh.global_size_mm": 650.0,
        "mesh.door_size_mm": 120.0,
        "sets_surfaces.base_set": "BASE_FIXED",
        "sets_surfaces.top_set": "TOP_LOAD_NODES",
        "sets_surfaces.door_set": "DOOR_OPENING_NODES",
        "sets_surfaces.wind_surface": "WINDWARD_SURFACE",
        "coupling.top_ring_enabled": True,
        "boundary.name": "BASE_FIXED",
        "boundary.target": "BASE_FIXED",
        "boundary.dofs": "1,2,3,4,5,6",
        "loads.top_total_force_n": round(lerp(anchor_a["top_force_n"], anchor_b["top_force_n"]), 1),
        "loads.top_force_direction": 1,
        "loads.top_mass_tonne": round(lerp(anchor_a["top_mass_tonne"], anchor_b["top_mass_tonne"]), 3),
        "loads.gravity_enabled": True,
        "loads.gravity_acceleration_mm_s2": 9810.0,
        "loads.wind_pressure_mpa": round(lerp(anchor_a["wind_pressure_mpa"], anchor_b["wind_pressure_mpa"]), 9),
        "loads.wind_surface": "WINDWARD_SURFACE",
        "steps.STEP-1-TOP-REACTION": "Apply gravity and top equivalent mass reaction preload",
        "steps.STEP-2-WIND-PRESSURE": "Apply tower wind pressure and top horizontal thrust",
        "steps.nlgeom": "NO",
        "modal.modes": 20,
        "outputs.field": ["U", "RF", "S", "E"],
    }
    return {
        "schema_version": "standard_input_v1",
        "source_file": f"{case_id}_heterogeneous_wind_tower.xlsx",
        "input_format": "synthetic_llm_benchmark",
        "unit_system": "mm-N-MPa-tonne",
        "tower_segments": segments,
        "door_opening": door,
        "materials": [material],
        "fem_settings": fem,
        "flanges": [],
        "analysis": {"static_strength": True, "modal": True, "modal_modes": 20, "buckling": False},
    }


def _material(alpha: float) -> dict[str, Any]:
    if alpha < 0.5:
        return {"name": "Q345", "elastic_modulus_mpa": 206000.0, "poisson_ratio": 0.3, "density_tonne_mm3": 7.85e-9, "yield_strength_mpa": 345.0}
    return {"name": "S355", "elastic_modulus_mpa": 210000.0, "poisson_ratio": 0.3, "density_tonne_mm3": 7.85e-9, "yield_strength_mpa": 355.0}


def _new_workbook(title: str) -> tuple[Workbook, Any]:
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws["A1"] = DISCLAIMER
    ws["A1"].font = Font(bold=True, color="9C0006")
    ws["A1"].fill = PatternFill("solid", fgColor="FFC7CE")
    return wb, ws


def _save(wb: Workbook, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)


def _write_segment_rows(ws: Any, start_row: int, segments: list[dict[str, Any]], headers: list[str], values: Callable[[dict[str, Any]], list[Any]]) -> None:
    for col, value in enumerate(headers, start=1):
        ws.cell(start_row, col, value)
        ws.cell(start_row, col).font = Font(bold=True)
    for row, segment in enumerate(segments, start=start_row + 1):
        for col, value in enumerate(values(segment), start=1):
            ws.cell(row, col, value)


def _write_key_values(ws: Any, start_row: int, rows: list[tuple[str, Any, str]]) -> None:
    ws.cell(start_row, 1, "Parameter / 参数")
    ws.cell(start_row, 2, "Value / 数值")
    ws.cell(start_row, 3, "Unit / 单位")
    for cell in ws[start_row]:
        cell.font = Font(bold=True)
    for row, (key, value, unit) in enumerate(rows, start=start_row + 1):
        ws.cell(row, 1, key)
        ws.cell(row, 2, value)
        ws.cell(row, 3, unit)


def _common_rows(case: dict[str, Any], mixed_units: bool = False, pressure_kpa: bool = False) -> list[tuple[str, Any, str]]:
    door = case["door_opening"]
    material = case["materials"][0]
    fem = case["fem_settings"]
    rows = [
        ("Steel grade / 钢材", material["name"], ""),
        ("Elastic modulus / 弹性模量", material["elastic_modulus_mpa"], "MPa"),
        ("Poisson ratio / 泊松比", material["poisson_ratio"], ""),
        ("Density / 密度", material["density_tonne_mm3"], "tonne/mm^3"),
        ("Yield strength / 屈服强度", material["yield_strength_mpa"], "MPa"),
        ("Door enabled / 门洞启用", "yes", ""),
        ("Door opening shape / 门洞形状", door["shape"], ""),
        ("Door clear width / 门洞净宽", door["width_mm"] / 10.0 if mixed_units else door["width_mm"], "cm" if mixed_units else "mm"),
        ("Door clear height / 门洞净高", door["height_mm"] / 1000.0 if mixed_units else door["height_mm"], "m" if mixed_units else "mm"),
        ("Door bottom elevation / 门洞底标高", door["z_bottom_mm"] / 1000.0 if mixed_units else door["z_bottom_mm"], "m" if mixed_units else "mm"),
        ("Element formulation / 单元", fem["element.element_type"], ""),
        ("Section formulation / 截面", fem["section.type"], ""),
        ("Section material / 截面材料", fem["section.material"], ""),
        ("Section thickness source / 厚度来源", fem["section.thickness_source"], ""),
        ("Mesh strategy / 网格策略", fem["mesh.strategy"], ""),
        ("Global mesh target", fem["mesh.global_size_mm"], "mm"),
        ("Door mesh target", fem["mesh.door_size_mm"], "mm"),
        ("Base node set / 塔底节点集", fem["sets_surfaces.base_set"], ""),
        ("Top load node set / 塔顶荷载节点集", fem["sets_surfaces.top_set"], ""),
        ("Door node set / 门洞节点集", fem["sets_surfaces.door_set"], ""),
        ("Wind surface / 迎风面", fem["sets_surfaces.wind_surface"], ""),
        ("Top horizontal force / 塔顶水平力", fem["loads.top_total_force_n"] / 1000.0 if mixed_units else fem["loads.top_total_force_n"], "kN" if mixed_units else "N"),
        ("Top force direction / 塔顶力方向", fem["loads.top_force_direction"], "global DOF"),
        ("Top equivalent mass / 顶部等效质量", fem["loads.top_mass_tonne"], "tonne"),
        ("Wind pressure / 风压", fem["loads.wind_pressure_mpa"] * 1000.0 if pressure_kpa else fem["loads.wind_pressure_mpa"], "kPa" if pressure_kpa else "MPa"),
        ("Wind load target / 风荷载面", fem["loads.wind_surface"], ""),
        ("Boundary condition name / 边界名称", fem["boundary.name"], ""),
        ("Base restraint target / 塔底约束对象", fem["boundary.target"], ""),
        ("Base restraint DOFs / 塔底约束自由度", fem["boundary.dofs"], ""),
        ("Gravity enabled / 重力", "yes" if fem["loads.gravity_enabled"] else "no", ""),
        ("Static step 1 / 静力步1", fem["steps.STEP-1-TOP-REACTION"], ""),
        ("Static step 2 / 静力步2", fem["steps.STEP-2-WIND-PRESSURE"], ""),
        ("Geometric nonlinearity / 几何非线性", fem["steps.nlgeom"], ""),
        ("Modal extraction count / 模态阶数", 20, "modes"),
        ("Requested field output", "U, RF, S, E", ""),
    ]
    return rows


def _write_s01(path: Path, case: dict[str, Any]) -> None:
    wb, ws = _new_workbook("轮廓清单_Profile")
    _write_segment_rows(
        ws,
        3,
        case["tower_segments"],
        ["构件编号", "下端标高(m)", "上端标高(m)", "下口外径(cm)", "上口外径(cm)", "板厚(mm)"],
        lambda s: [s["segment_id"], s["z_bottom_mm"] / 1000.0, s["z_top_mm"] / 1000.0, s["diameter_bottom_mm"] / 10.0, s["diameter_top_mm"] / 10.0, s["thickness_mm"]],
    )
    ws2 = wb.create_sheet("设计意图_DesignIntent")
    ws2["A1"] = DISCLAIMER
    _write_key_values(ws2, 3, _common_rows(case, mixed_units=True, pressure_kpa=False))
    _save(wb, path)


def _write_s02(path: Path, case: dict[str, Any]) -> None:
    wb, ws = _new_workbook("Shell Schedule")
    _write_segment_rows(ws, 4, case["tower_segments"], ["Wall t [mm]", "Upper OD [mm]", "Station end [mm]", "Member", "Lower OD [mm]", "Station start [mm]"], lambda s: [s["thickness_mm"], s["diameter_top_mm"], s["z_top_mm"], s["segment_id"], s["diameter_bottom_mm"], s["z_bottom_mm"]])
    ws2 = wb.create_sheet("Design Basis")
    _write_key_values(ws2, 2, list(reversed(_common_rows(case))))
    _save(wb, path)


def _write_s03(path: Path, case: dict[str, Any]) -> None:
    wb, ws = _new_workbook("Parameter Ledger")
    _write_key_values(ws, 3, _common_rows(case))
    ws2 = wb.create_sheet("Course Geometry")
    _write_segment_rows(ws2, 2, case["tower_segments"], ["Course", "From", "To", "OD lower", "OD upper", "Plate"], lambda s: [s["segment_id"], s["z_bottom_mm"], s["z_top_mm"], s["diameter_bottom_mm"], s["diameter_top_mm"], s["thickness_mm"]])
    _save(wb, path)


def _write_s04(path: Path, case: dict[str, Any]) -> None:
    wb, ws = _new_workbook("Rotated Geometry")
    labels = ["quantity", *[s["segment_id"] for s in case["tower_segments"]]]
    rows = [
        labels,
        ["bottom elevation mm", *[s["z_bottom_mm"] for s in case["tower_segments"]]],
        ["top elevation mm", *[s["z_top_mm"] for s in case["tower_segments"]]],
        ["bottom outside diameter mm", *[s["diameter_bottom_mm"] for s in case["tower_segments"]]],
        ["top outside diameter mm", *[s["diameter_top_mm"] for s in case["tower_segments"]]],
        ["wall thickness mm", *[s["thickness_mm"] for s in case["tower_segments"]]],
    ]
    for r, values in enumerate(rows, start=3):
        for c, value in enumerate(values, start=1):
            ws.cell(r, c, value)
    ws2 = wb.create_sheet("Model Conditions")
    _write_key_values(ws2, 2, _common_rows(case))
    _save(wb, path)


def _write_s05(path: Path, case: dict[str, Any]) -> None:
    wb, ws = _new_workbook("Lower Courses")
    first, second = case["tower_segments"][:5], case["tower_segments"][5:]
    headers = ["Course", "z0 mm", "z1 mm", "D0 mm", "D1 mm", "t mm"]
    def values(segment: dict[str, Any]) -> list[Any]:
        return [
            segment["segment_id"],
            segment["z_bottom_mm"],
            segment["z_top_mm"],
            segment["diameter_bottom_mm"],
            segment["diameter_top_mm"],
            segment["thickness_mm"],
        ]
    _write_segment_rows(ws, 3, first, headers, values)
    ws2 = wb.create_sheet("Upper Courses")
    _write_segment_rows(ws2, 2, second, headers, values)
    ws3 = wb.create_sheet("Analysis Intent")
    _write_key_values(ws3, 2, _common_rows(case))
    _save(wb, path)


def _write_s06(path: Path, case: dict[str, Any]) -> None:
    wb, ws = _new_workbook("Fabrication Data")
    ws.merge_cells("A2:F2")
    ws["A2"] = "Synthetic shell can schedule — headers intentionally offset"
    _write_segment_rows(ws, 7, case["tower_segments"], ["ID", "Start mm", "End mm", "OD-start mm", "OD-end mm", "Gauge mm"], lambda s: [s["segment_id"], s["z_bottom_mm"], s["z_top_mm"], s["diameter_bottom_mm"], s["diameter_top_mm"], s["thickness_mm"]])
    ws2 = wb.create_sheet("Assumptions")
    _write_key_values(ws2, 6, _common_rows(case))
    _save(wb, path)


def _write_s07(path: Path, case: dict[str, Any]) -> None:
    wb, ws = _new_workbook("Mixed Unit Courses")
    _write_segment_rows(ws, 3, case["tower_segments"], ["Course", "z-start m", "z-end m", "OD-start cm", "OD-end m", "plate mm"], lambda s: [s["segment_id"], s["z_bottom_mm"] / 1000.0, s["z_top_mm"] / 1000.0, s["diameter_bottom_mm"] / 10.0, s["diameter_top_mm"] / 1000.0, s["thickness_mm"]])
    ws2 = wb.create_sheet("Mixed Unit Actions")
    _write_key_values(ws2, 2, _common_rows(case, mixed_units=True, pressure_kpa=True))
    _save(wb, path)


def _write_s08(path: Path, case: dict[str, Any]) -> None:
    wb, ws = _new_workbook("Engineer Notes")
    fem = case["fem_settings"]
    door = case["door_opening"]
    notes = [
        "Use a ten-course tapered steel shell tower represented with S4R shell elements.",
        f"The steel grade is {case['materials'][0]['name']}; the base is fully fixed in degrees of freedom 1 through 6.",
        f"Provide a rounded rectangular door {door['width_mm']} mm wide and {door['height_mm']} mm high, beginning at {door['z_bottom_mm']} mm.",
        f"Transfer a {fem['loads.top_total_force_n']} N horizontal top force and {fem['loads.top_mass_tonne']} tonne equivalent mass through the top ring.",
        f"Apply {fem['loads.wind_pressure_mpa']} MPa pressure on the windward shell surface; request static and 20-mode frequency analyses.",
    ]
    for row, note in enumerate(notes, start=3):
        ws.cell(row, 1, note)
    ws2 = wb.create_sheet("Compact Schedule")
    _write_segment_rows(ws2, 2, case["tower_segments"], ["No.", "z0", "z1", "D0", "D1", "t"], lambda s: [s["segment_id"], s["z_bottom_mm"], s["z_top_mm"], s["diameter_bottom_mm"], s["diameter_top_mm"], s["thickness_mm"]])
    ws3 = wb.create_sheet("Constants")
    _write_key_values(ws3, 2, _common_rows(case)[:5] + _common_rows(case)[10:13])
    _save(wb, path)


def _write_s09(path: Path, case: dict[str, Any]) -> None:
    wb, ws = _new_workbook("Summary Copy")
    _write_key_values(ws, 3, _common_rows(case))
    ws2 = wb.create_sheet("Detailed Copy")
    _write_key_values(ws2, 2, _common_rows(case))
    ws3 = wb.create_sheet("Plate Courses")
    _write_segment_rows(ws3, 2, case["tower_segments"], ["Course", "Lower z mm", "Upper z mm", "Lower OD mm", "Upper OD mm", "Thickness mm"], lambda s: [s["segment_id"], s["z_bottom_mm"], s["z_top_mm"], s["diameter_bottom_mm"], s["diameter_top_mm"], s["thickness_mm"]])
    _save(wb, path)


def _write_s10(path: Path, case: dict[str, Any]) -> None:
    wb, ws = _new_workbook("Minimum Required Data")
    _write_segment_rows(ws, 3, case["tower_segments"], ["Ring", "Elevation low mm", "Elevation high mm", "OD low mm", "OD high mm", "Plate mm"], lambda s: [s["segment_id"], s["z_bottom_mm"], s["z_top_mm"], s["diameter_bottom_mm"], s["diameter_top_mm"], s["thickness_mm"]])
    rows = [item for item in _common_rows(case) if item[0] not in {"Global mesh target", "Door mesh target", "Requested field output"}]
    ws2 = wb.create_sheet("Required Intent")
    _write_key_values(ws2, 2, rows)
    _save(wb, path)


def _format_challenge(index: int) -> str:
    return (
        "bilingual_headers_m_cm_mm_kn",
        "english_aliases_shuffled_columns",
        "long_form_parameter_ledger",
        "transposed_segment_table",
        "split_geometry_sheets",
        "merged_titles_offset_headers",
        "mixed_units_by_column",
        "narrative_notes_and_tables",
        "consistent_duplicate_summary",
        "missing_defaultable_optional_fields",
    )[index - 1]


def _value(value: Any) -> Any:
    return value.get("value") if isinstance(value, dict) and "value" in value else value


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()
