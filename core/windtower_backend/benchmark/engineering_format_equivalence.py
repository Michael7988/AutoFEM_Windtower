from __future__ import annotations

import csv
from copy import deepcopy
import hashlib
import json
import shutil
import time
from dataclasses import replace
from pathlib import Path
from typing import Any, Iterable

import fitz
import yaml
from docx import Document
from docx.enum.section import WD_SECTION
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt

from ..config import load_settings
from ..capability_profile import CANONICAL_DEFAULTS
from ..core.pipeline import WindTowerPipeline
from ..importers.windtower_adapter import parse_windtower_workbook
from ..input_batches.service import EngineeringInputBatchService
from ..l0 import build_l0
from ..modeling.implementation import ProjectState
from ..modeling.implementation import build_ir
from ..parameter_normalization import normalize_parameter_value
from ..parameter_ir import (
    ParameterSet,
    apply_parameter_set_to_ir,
    parameter_set_from_ir,
    validate_parameter_set,
)
from ..reproducibility import normalized_inp_sha256
from ..reproducibility import stable_hash
from .metrics import aggregate_binary_metric, compare_parameter_payloads, stability_metrics


DATASET_SCHEMA_VERSION = "g3e_engineering_format_dataset_v1"
RESULT_SCHEMA_VERSION = "g3e_engineering_format_results_v1"
DEFAULT_PRIVATE_DIRECTORY = "private_llm_format_equivalence"
DEFAULT_MODELING_GOAL = (
    "建立带门洞钢制风电塔筒整体S4R壳模型，执行线弹性静力与模态分析。"
)
ENGINEERING_CARRIERS = (
    "xlsx",
    "csv",
    "json",
    "yaml",
    "txt",
    "docx",
    "pptx",
    "text_pdf",
    "scan_pdf",
    "png",
)
DEFAULT_PARAMETER_KEYS = {
    "analysis.modal_modes",
    "output_requests",
}


def discover_engineering_case_pairs(
    source_root: str | Path,
    *,
    expected_cases: int | None = 4,
) -> list[dict[str, Path | str]]:
    root = Path(source_root)
    fem_files = sorted(root.glob("*_FEM参数.xls"))
    pairs: list[dict[str, Path | str]] = []
    for index, fem_path in enumerate(fem_files, start=1):
        raw_path = fem_path.with_name(fem_path.name.replace("_FEM参数.xls", ".xls"))
        if not raw_path.is_file():
            raise FileNotFoundError(f"缺少机械接口原表：{raw_path.name}")
        pairs.append(
            {
                "case_id": f"CASE{index:02d}",
                "raw": raw_path,
                "fem": fem_path,
            }
        )
    if expected_cases is not None and len(pairs) != expected_cases:
        raise ValueError(
            f"期望发现 {expected_cases} 个工程案例，实际发现 {len(pairs)} 个。"
        )
    return pairs


def generate_engineering_format_dataset(
    source_root: str | Path,
    output_root: str | Path,
    *,
    expected_cases: int | None = 4,
    force: bool = False,
    compile_reference_models: bool = True,
) -> dict[str, Any]:
    source = Path(source_root).resolve()
    output = Path(output_root).resolve()
    pairs = discover_engineering_case_pairs(source, expected_cases=expected_cases)
    if output.exists() and any(output.iterdir()):
        if not force:
            raise FileExistsError(f"输出目录非空：{output}；如需重建请使用 force=True。")
        if output == source or source not in output.parents:
            raise ValueError("仅允许清理 source_root 下的专用实验目录。")
        shutil.rmtree(output)
    output.mkdir(parents=True, exist_ok=True)

    manifest_cases = []
    for pair in pairs:
        manifest_cases.append(
            _generate_case(
                pair,
                output,
                compile_reference_models=compile_reference_models,
            )
        )

    manifest = {
        "schema_version": DATASET_SCHEMA_VERSION,
        "experiment_id": "G3-E",
        "generated_at": _utc_now(),
        "modeling_goal": DEFAULT_MODELING_GOAL,
        "truth_policy": (
            "Each FEM workbook is normalized by the current production parser. "
            "The raw mechanical workbook is retained only as source context."
        ),
        "truth_isolation": (
            "Ground-truth files are stored outside every carrier directory and "
            "are never uploaded in an input batch."
        ),
        "case_count": len(manifest_cases),
        "carrier_count": len(ENGINEERING_CARRIERS),
        "condition_count": len(manifest_cases) * len(ENGINEERING_CARRIERS),
        "carriers": list(ENGINEERING_CARRIERS),
        "cases": manifest_cases,
    }
    _write_json(output / "manifest.json", manifest)
    _write_result_templates(output / "results")
    (output / "runs" / "deterministic").mkdir(parents=True, exist_ok=True)
    (output / "runs" / "llm").mkdir(parents=True, exist_ok=True)
    return manifest


def run_engineering_format_experiment(
    dataset_root: str | Path,
    *,
    run_llm: bool = False,
    repetitions: int = 3,
    poll_seconds: float = 0.2,
    timeout_seconds: float = 300.0,
) -> dict[str, Any]:
    root = Path(dataset_root).resolve()
    manifest = json.loads((root / "manifest.json").read_text(encoding="utf-8"))
    if manifest.get("schema_version") != DATASET_SCHEMA_VERSION:
        raise ValueError("数据集清单版本不受支持。")
    if repetitions < 1:
        raise ValueError("repetitions 必须大于或等于 1。")

    mode = "llm" if run_llm else "deterministic"
    run_root = root / "runs" / mode
    run_root.mkdir(parents=True, exist_ok=True)
    settings = load_settings()
    settings = replace(settings, metadata_db=run_root / "metadata.sqlite3")
    service = EngineeringInputBatchService(WindTowerPipeline(settings))
    service.root = run_root / "input_batches"
    service.cache_root = service.root / "_cache"
    service.root.mkdir(parents=True, exist_ok=True)
    service.cache_root.mkdir(parents=True, exist_ok=True)
    if run_llm and not service.capabilities().get("configured"):
        raise RuntimeError("未配置LLM；未显式可用时不得启动真实LLM实验。")

    rows = []
    compile_cache: dict[str, dict[str, Any]] = {}
    repeat_count = repetitions if run_llm else 1
    for case in manifest["cases"]:
        truth = json.loads((root / case["ground_truth"]["parameter_set"]).read_text(encoding="utf-8"))
        truth_parameter_set = ParameterSet.model_validate(truth)
        truth_payload = parameter_set_active_payload(truth_parameter_set)
        reference_ir = json.loads(
            (root / case["ground_truth"]["reference_ir"]).read_text(encoding="utf-8")
        )
        truth_signature = case["ground_truth"]["model_signature_sha256"]
        for carrier in case["carriers"]:
            for repeat in range(1, repeat_count + 1):
                rows.append(
                    _run_condition(
                        service,
                        root,
                        case["case_id"],
                        carrier,
                        repeat,
                        mode,
                        truth_parameter_set,
                        truth_payload,
                        reference_ir,
                        truth_signature,
                        compile_cache,
                        poll_seconds,
                        timeout_seconds,
                    )
                )

    result = _summarize_results(rows, mode=mode, repetitions=repeat_count)
    _write_json(root / "results" / f"{mode}_summary.json", result)
    _write_rows_csv(root / "results" / f"{mode}_runs.csv", rows)
    _write_format_report(root / "results" / f"{mode}_format_report.md", result)
    return result


def parameter_set_active_payload(parameter_set: ParameterSet) -> dict[str, Any]:
    sections: dict[int, dict[str, Any]] = {}
    groups: dict[str, dict[str, Any]] = {
        "door_opening": {},
        "material": {},
        "mesh": {},
        "loads": {},
        "analysis": {},
        "boundary": {},
    }
    for key, record in sorted(parameter_set.parameters.items()):
        if key.startswith("geometry.tower_sections."):
            _, _, index, field = key.split(".", 3)
            sections.setdefault(int(index), {})[field] = record.value
        elif key.startswith("geometry.door."):
            groups["door_opening"][key[len("geometry.door.") :]] = record.value
        else:
            group, _, field = key.partition(".")
            if group in groups:
                groups[group][field] = record.value
    return {
        "schema_version": "g3e_active_parameter_payload_v1",
        "structure_family": parameter_set.structure_family,
        "tower_segments": [
            {"segment_id": f"SEG_{index + 1:03d}", **sections[index]}
            for index in sorted(sections)
        ],
        **groups,
        "analysis_tasks": parameter_set.analysis_tasks,
        "output_requests": parameter_set.output_requests,
    }


def carrier_artifacts(dataset_root: str | Path, carrier: dict[str, Any]) -> list[Path]:
    root = Path(dataset_root)
    return [root / item for item in carrier["artifacts"]]


def _generate_case(
    pair: dict[str, Path | str],
    output: Path,
    *,
    compile_reference_models: bool,
) -> dict[str, Any]:
    case_id = str(pair["case_id"])
    raw_path = Path(pair["raw"])
    fem_path = Path(pair["fem"])
    source_dir = output / "source" / case_id
    source_dir.mkdir(parents=True, exist_ok=True)
    raw_copy = source_dir / raw_path.name
    fem_copy = source_dir / fem_path.name
    shutil.copy2(raw_path, raw_copy)
    shutil.copy2(fem_path, fem_copy)

    parsed = parse_windtower_workbook(fem_path)
    ir = build_ir(
        f"g3e-{case_id.lower()}",
        fem_path.name,
        parsed.get("tower") or [],
        parsed.get("flanges") or [],
        parsed.get("door") or {},
        parsed.get("fem_params") or {},
        parsed,
        DEFAULT_MODELING_GOAL,
    )
    parameter_set = parameter_set_from_ir(f"g3e-{case_id.lower()}", ir)
    direction = parameter_set.parameters.get("loads.top_force_direction")
    if direction is not None:
        direction.value, conversion = normalize_parameter_value(
            direction.key, direction.value
        )
        if conversion:
            direction.note = f"Canonical experiment normalization: {conversion}"
    _mark_experiment_defaults(parameter_set, parsed)
    validation = validate_parameter_set(parameter_set)
    if validation["status"] != "pass":
        raise ValueError(f"{case_id} 正式参数集未通过参数门禁：{validation['issues']}")
    active_payload = parameter_set_active_payload(parameter_set)
    active_payload_sha256 = _semantic_model_signature(active_payload)
    truth_dir = output / "ground_truth"
    truth_dir.mkdir(parents=True, exist_ok=True)
    reference_model_root = truth_dir / "reference_models" / case_id
    reference_compile = (
        _compile_model(
            ir,
            parameter_set,
            reference_model_root,
            project_id=f"g3e-reference-{case_id.lower()}",
        )
        if compile_reference_models
        else {
            "status": "skipped",
            "signature": active_payload_sha256,
            "job_hashes": {},
            "node_count": None,
            "element_count": None,
            "generator": "not_run",
        }
    )
    model_signature = reference_compile["signature"]

    parameter_path = truth_dir / f"{case_id}_parameter_set.json"
    active_path = truth_dir / f"{case_id}_active_parameters.json"
    source_status_path = truth_dir / f"{case_id}_source_status.json"
    signature_path = truth_dir / f"{case_id}_reference_model_signature.json"
    ir_path = truth_dir / f"{case_id}_reference_ir.json"
    _write_json(parameter_path, parameter_set.model_dump(mode="json"))
    _write_json(active_path, active_payload)
    _write_json(ir_path, ir)
    _write_json(
        source_status_path,
        {
            "schema_version": "g3e_source_status_v1",
            "case_id": case_id,
            "observed_source": fem_path.name,
            "context_source": raw_path.name,
            "experiment_defaults": sorted(DEFAULT_PARAMETER_KEYS),
            "validation": validation,
        },
    )
    _write_json(
        signature_path,
        {
            "schema_version": "g3e_reference_model_signature_v1",
            "case_id": case_id,
            "signature_scope": (
                "Active parameters consumed by the deterministic wind-tower "
                "adapter; volatile project metadata and source filenames excluded."
            ),
            "sha256": model_signature,
            "compile": reference_compile,
        },
    )

    carriers = _write_case_carriers(output, case_id, active_payload)
    return {
        "case_id": case_id,
        "source": {
            "raw": _relative(output, raw_copy),
            "fem": _relative(output, fem_copy),
            "raw_sha256": _sha256(raw_copy),
            "fem_sha256": _sha256(fem_copy),
        },
        "ground_truth": {
            "parameter_set": _relative(output, parameter_path),
            "active_parameters": _relative(output, active_path),
            "source_status": _relative(output, source_status_path),
            "reference_model_signature": _relative(output, signature_path),
            "reference_ir": _relative(output, ir_path),
            "active_payload_sha256": active_payload_sha256,
            "model_signature_sha256": model_signature,
        },
        "tower_segment_count": len(active_payload["tower_segments"]),
        "carriers": carriers,
    }


def _mark_experiment_defaults(
    parameter_set: ParameterSet,
    parsed: dict[str, Any],
) -> None:
    modal = parameter_set.parameters.get("analysis.modal_modes")
    if modal is not None:
        modal.status = "defaulted"
        modal.note = "Fixed G3-E experiment default; not observed in the source workbook."
        modal.sources = []
    offset = parameter_set.parameters.get("loads.center_of_mass_offset_mm")
    if offset is not None and offset.confirmation_status == "pending":
        offset.status = "confirmed"
        offset.confirmation_status = "confirmed"
        offset.requires_confirmation = False
        offset.note = "Frozen G3-E ground-truth assumption confirmed by the experiment protocol."
    parameter_set.output_requests = ["U", "RF", "S", "E", "FREQUENCY"]
    parameter_set.confirmed_by = "G3-E deterministic truth generation"
    parameter_set.confirmed_at = _utc_now()
    _ = parsed


def _write_case_carriers(
    output: Path,
    case_id: str,
    payload: dict[str, Any],
) -> list[dict[str, Any]]:
    case_root = output / "inputs" / case_id
    rows = _payload_rows(payload)
    page_blocks = _page_blocks(payload)
    active_payload_sha256 = _semantic_model_signature(payload)
    writers = {
        "xlsx": lambda directory: _write_xlsx(directory / f"{case_id}.xlsx", payload),
        "csv": lambda directory: _write_csv(directory / f"{case_id}.csv", rows),
        "json": lambda directory: _write_json(
            directory / f"{case_id}.json", {"standard_input_package": payload}
        ),
        "yaml": lambda directory: _write_yaml(
            directory / f"{case_id}.yaml", {"standard_input_package": payload}
        ),
        "txt": lambda directory: _write_txt(directory / f"{case_id}.txt", page_blocks),
        "docx": lambda directory: _write_docx(directory / f"{case_id}.docx", page_blocks),
        "pptx": lambda directory: _write_pptx(directory / f"{case_id}.pptx", page_blocks),
        "text_pdf": lambda directory: _write_text_pdf(directory / f"{case_id}.pdf", page_blocks),
        "scan_pdf": lambda directory: _write_scan_pdf(directory / f"{case_id}.pdf", page_blocks),
        "png": lambda directory: _write_png_pages(directory, case_id, page_blocks),
    }
    carriers = []
    for carrier in ENGINEERING_CARRIERS:
        directory = case_root / carrier
        directory.mkdir(parents=True, exist_ok=True)
        artifacts = writers[carrier](directory)
        if isinstance(artifacts, Path):
            artifact_paths = [artifacts]
        else:
            artifact_paths = list(artifacts)
        carriers.append(
            {
                "carrier": carrier,
                "condition_id": f"{case_id}-{carrier}",
                "artifacts": [_relative(output, item) for item in artifact_paths],
                "checksums": {_relative(output, item): _sha256(item) for item in artifact_paths},
                "tower_segment_coverage": list(range(1, len(payload["tower_segments"]) + 1)),
                "active_payload_sha256": active_payload_sha256,
                "input_only": True,
            }
        )
    return carriers


def _payload_rows(payload: dict[str, Any]) -> list[dict[str, Any]]:
    units = {
        "z_bottom_mm": "mm",
        "z_top_mm": "mm",
        "diameter_bottom_mm": "mm",
        "diameter_top_mm": "mm",
        "thickness_mm": "mm",
        "width_mm": "mm",
        "height_mm": "mm",
        "bottom_z_mm": "mm",
        "theta_center_deg": "deg",
        "elastic_modulus_mpa": "MPa",
        "yield_strength_mpa": "MPa",
        "density": "tonne/mm^3",
        "global_size_mm": "mm",
        "door_size_mm": "mm",
        "patch_margin_u_mm": "mm",
        "patch_margin_z_mm": "mm",
        "top_mass_tonne": "tonne",
        "top_horizontal_force_n": "N",
        "wind_pressure_mpa": "MPa",
    }
    rows = []
    for path, value in _flatten(payload):
        field = path.rsplit(".", 1)[-1]
        status = "experiment_default" if path in {
            "analysis.modal_modes",
            "output_requests",
        } else "observed_or_contract"
        rows.append(
            {
                "path": path,
                "value": json.dumps(value, ensure_ascii=False) if isinstance(value, list) else value,
                "unit": units.get(field, ""),
                "status": status,
            }
        )
    return rows


def _page_blocks(payload: dict[str, Any], per_page: int = 10) -> list[list[str]]:
    segments = payload["tower_segments"]
    blocks = [
        [
            "G3-E engineering input",
            f"structure_family = {payload['structure_family']}",
            "All dimensions use the units shown. Do not infer missing values.",
        ]
    ]
    for start in range(0, len(segments), per_page):
        lines = [
            (
                "segment_id | z_bottom_mm | z_top_mm | diameter_bottom_mm | "
                "diameter_top_mm | thickness_mm"
            )
        ]
        for segment in segments[start : start + per_page]:
            lines.append(
                " | ".join(
                    str(segment[key])
                    for key in (
                        "segment_id",
                        "z_bottom_mm",
                        "z_top_mm",
                        "diameter_bottom_mm",
                        "diameter_top_mm",
                        "thickness_mm",
                    )
                )
            )
        blocks.append(lines)
    blocks.append(
        [
            f"{group}.{key} = {json.dumps(value, ensure_ascii=False)}"
            for group in (
                "door_opening",
                "material",
                "mesh",
                "loads",
                "analysis",
                "boundary",
            )
            for key, value in payload[group].items()
        ]
        + [
            f"analysis_tasks = {json.dumps(payload['analysis_tasks'], ensure_ascii=False)}",
            (
                "output_requests = "
                f"{json.dumps(payload['output_requests'], ensure_ascii=False)} "
                "[experiment_default]"
            ),
        ]
    )
    return blocks


def _write_xlsx(path: Path, payload: dict[str, Any]) -> Path:
    workbook = Workbook()
    workbook.remove(workbook.active)
    global_sheet = workbook.create_sheet("01_Global")
    global_sheet.append(["Category", "Parameter", "key", "value", "unit", "source"])
    global_sheet.append(
        ["model", "structure family", "structure_family", payload["structure_family"], "", "G3-E"]
    )

    segment_sheet = workbook.create_sheet("02_Segments")
    segment_sheet.append(
        [
            "塔节ID",
            "底标高 mm",
            "顶标高 mm",
            "高度 mm",
            "底外径 mm",
            "顶外径 mm",
            "厚度 mm",
        ]
    )
    for segment in payload["tower_segments"]:
        segment_sheet.append(
            [
                segment["segment_id"],
                segment["z_bottom_mm"],
                segment["z_top_mm"],
                segment["z_top_mm"] - segment["z_bottom_mm"],
                segment["diameter_bottom_mm"],
                segment["diameter_top_mm"],
                segment["thickness_mm"],
            ]
        )

    door_sheet = workbook.create_sheet("03_Door")
    door_sheet.append(["Category", "Parameter", "key", "value", "unit"])
    door_key_map = {
        "enabled": ("door_enabled", ""),
        "theta_center_deg": ("door_azimuth", "deg"),
        "bottom_z_mm": ("door_z_bottom", "mm"),
        "width_mm": ("door_width", "mm"),
        "height_mm": ("door_height", "mm"),
    }
    for key, (target, unit) in door_key_map.items():
        door_sheet.append(["geometry", key, target, payload["door_opening"].get(key), unit])

    material_sheet = workbook.create_sheet("04_Materials")
    material = payload["material"]
    material_sheet.append(
        [
            "材料名",
            "用途",
            "E MPa",
            "nu",
            "密度 tonne/mm3",
            "fy MPa",
            "fu MPa",
            "",
            "",
        ]
    )
    material_sheet.append(
        [
            material.get("name"),
            "tower shell",
            material.get("elastic_modulus_mpa"),
            material.get("poisson_ratio"),
            material.get("density"),
            material.get("yield_strength_mpa"),
            None,
            "",
            "",
        ]
    )
    material_sheet.append([])
    material_sheet.append(["壳截面定义"])
    material_sheet.append([])
    material_sheet.append(
        ["截面名", "塔节", "厚度 mm", "材料", "备注", "", "", "参数", "值"]
    )
    for index, segment in enumerate(payload["tower_segments"], start=1):
        material_sheet.append(
            [
                f"SEC_{index:03d}",
                segment["segment_id"],
                segment["thickness_mm"],
                material.get("name"),
                "S4R",
                "",
                "",
                "RP_TOP_MASS" if index == 1 else "",
                payload["loads"].get("top_mass_tonne") if index == 1 else "",
            ]
        )

    fem_sheet = workbook.create_sheet("05_FEM")
    fem_sheet.append(["模块", "参数", "Abaqus/关键词", "推荐值", "单位", "范围", "说明"])
    fem_rows = [
        ("model", "element", "element_type", payload["mesh"].get("element_type"), "", "all"),
        ("mesh", "global size", "mesh_global", payload["mesh"].get("global_size_mm"), "mm", "all"),
        ("mesh", "door size", "mesh_door", payload["mesh"].get("door_size_mm"), "mm", "door"),
    ]
    for group in ("mesh", "loads", "analysis", "boundary"):
        for key, value in payload[group].items():
            path_key = f"{group}.{key}"
            if path_key in {
                "mesh.element_type",
                "mesh.global_size_mm",
                "mesh.door_size_mm",
            }:
                continue
            fem_rows.append(
                (
                    group,
                    key,
                    path_key,
                    json.dumps(value, ensure_ascii=False)
                    if isinstance(value, list)
                    else value,
                    "",
                    "all",
                )
            )
    for row in fem_rows:
        fem_sheet.append([*row, "G3-E"])

    load_sheet = workbook.create_sheet("06_LoadCases")
    load_sheet.append(
        [
            "工况ID",
            "Abaqus Fx N",
            "Abaqus Fy N",
            "Abaqus Fz N",
            "Abaqus Mx Nmm",
            "Abaqus My Nmm",
            "Abaqus Mz Nmm",
        ]
    )
    force = payload["loads"].get("top_horizontal_force_n") or 0
    direction = payload["loads"].get("top_force_direction")
    fx = force if direction == "+X" else -force if direction == "-X" else 0
    fy = force if direction == "+Y" else -force if direction == "-Y" else 0
    load_sheet.append(["LC_G3E_01", fx, fy, 0, 0, 0, 0])
    load_sheet.append([])
    load_sheet.append(["塔节", "面压力 N/mm²", "节段风力 N"])
    load_sheet.append(["ALL", payload["loads"].get("wind_pressure_mpa"), 0])

    modal_sheet = workbook.create_sheet("07_Modal")
    modal_sheet.append(["设置", "值", "", "", "", "", "特征值方法", "值"])
    modal_sheet.append(
        [
            "Requested modes",
            payload["analysis"].get("modal_modes"),
            "",
            "",
            "",
            "",
            "特征值方法",
            "LANCZOS",
        ]
    )
    modal_sheet.append(
        [
            f"{payload['analysis'].get('modal_modes')} 阶 [experiment_default]",
        ]
    )
    output_sheet = workbook.create_sheet("08_Outputs")
    output_sheet.append(["分析", "目标", "Abaqus输出", "分量"])
    for request in payload["output_requests"]:
        output_sheet.append(["static/modal", "model", request, request])
    workbook.save(path)
    return path


def _write_csv(path: Path, rows: list[dict[str, Any]]) -> Path:
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=["path", "value", "unit", "status"])
        writer.writeheader()
        writer.writerows(rows)
    return path


def _write_yaml(path: Path, payload: dict[str, Any]) -> Path:
    path.write_text(
        yaml.safe_dump(payload, allow_unicode=True, sort_keys=False),
        encoding="utf-8",
    )
    return path


def _write_txt(path: Path, blocks: list[list[str]]) -> Path:
    path.write_text(
        "\n\n".join(
            f"[PAGE {index}]\n" + "\n".join(block)
            for index, block in enumerate(blocks, start=1)
        ),
        encoding="utf-8",
    )
    return path


def _write_docx(path: Path, blocks: list[list[str]]) -> Path:
    document = Document()
    for page_index, block in enumerate(blocks):
        if page_index:
            document.add_section(WD_SECTION.NEW_PAGE)
        document.add_heading(f"G3-E Engineering Input — Page {page_index + 1}", level=1)
        if block and block[0].startswith("segment_id"):
            headers = [item.strip() for item in block[0].split("|")]
            table = document.add_table(rows=1, cols=len(headers))
            for index, header in enumerate(headers):
                table.rows[0].cells[index].text = header
            for line in block[1:]:
                cells = [item.strip() for item in line.split("|")]
                row = table.add_row()
                for index, value in enumerate(cells):
                    row.cells[index].text = value
        else:
            for line in block:
                document.add_paragraph(line)
    document.save(path)
    return path


def _write_pptx(path: Path, blocks: list[list[str]]) -> Path:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    for page_index, block in enumerate(blocks):
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        title = slide.shapes.title
        title.text = f"G3-E Engineering Input — Page {page_index + 1}"
        text_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.2), Inches(12.3), Inches(5.8))
        frame = text_box.text_frame
        frame.word_wrap = True
        frame.text = "\n".join(block)
        for paragraph in frame.paragraphs:
            paragraph.font.name = "Consolas"
            paragraph.font.size = Pt(13)
    presentation.save(path)
    return path


def _write_text_pdf(path: Path, blocks: list[list[str]]) -> Path:
    document = fitz.open()
    for block in blocks:
        page = document.new_page(width=842, height=595)
        page.insert_textbox(
            fitz.Rect(35, 35, 807, 560),
            "\n".join(block),
            fontsize=8.5,
            fontname="cour",
            lineheight=1.25,
        )
    document.save(path)
    document.close()
    return path


def _write_scan_pdf(path: Path, blocks: list[list[str]]) -> Path:
    images = [_render_page(block) for block in blocks]
    document = fitz.open()
    for image in images:
        page = document.new_page(width=image.width, height=image.height)
        page.insert_image(page.rect, stream=_image_bytes(image))
    document.save(path)
    document.close()
    return path


def _write_png_pages(
    directory: Path,
    case_id: str,
    blocks: list[list[str]],
) -> list[Path]:
    paths = []
    for index, block in enumerate(blocks, start=1):
        path = directory / f"{case_id}_page_{index:02d}.png"
        _render_page(block).save(path)
        paths.append(path)
    return paths


def _render_page(lines: list[str]) -> Image.Image:
    image = Image.new("RGB", (1800, 1270), "white")
    draw = ImageDraw.Draw(image)
    font = _monospace_font(24)
    y = 45
    for raw_line in lines:
        wrapped = _wrap_line(raw_line, 112)
        for line in wrapped:
            draw.text((45, y), line, fill="black", font=font)
            y += 34
        y += 4
    return image


def _monospace_font(size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = (
        Path("C:/Windows/Fonts/consola.ttf"),
        Path("C:/Windows/Fonts/arial.ttf"),
    )
    for candidate in candidates:
        if candidate.is_file():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


def _image_bytes(image: Image.Image) -> bytes:
    import io

    stream = io.BytesIO()
    image.save(stream, format="PNG")
    return stream.getvalue()


def _wrap_line(text: str, width: int) -> list[str]:
    if len(text) <= width:
        return [text]
    return [text[index : index + width] for index in range(0, len(text), width)]


def _run_condition(
    service: EngineeringInputBatchService,
    root: Path,
    case_id: str,
    carrier: dict[str, Any],
    repeat: int,
    mode: str,
    truth_parameter_set: ParameterSet,
    truth_payload: dict[str, Any],
    reference_ir: dict[str, Any],
    truth_signature: str,
    compile_cache: dict[str, dict[str, Any]],
    poll_seconds: float,
    timeout_seconds: float,
) -> dict[str, Any]:
    artifacts = carrier_artifacts(root, carrier)
    uploads = [
        (path.name, _media_type(path), path.read_bytes())
        for path in artifacts
    ]
    started = time.perf_counter()
    usage_before = service.provider.usage()
    batch = service.create_batch(
        uploads,
        modeling_goal=DEFAULT_MODELING_GOAL,
        parse_mode=mode,
    )
    deadline = time.monotonic() + timeout_seconds
    while batch["status"] in {"queued", "processing"}:
        if time.monotonic() >= deadline:
            raise TimeoutError(f"{carrier['condition_id']} 解析超时。")
        time.sleep(poll_seconds)
        batch = service.get_batch(batch["batch_id"])
    predicted = _resolved_payload(batch.get("resolved_parameters") or {})
    metrics = compare_parameter_payloads(predicted, truth_payload)
    parameter_equivalent = metrics["exact_field_accuracy"] == 1.0
    compilable_payload = _apply_public_compiler_defaults(predicted)
    model_metrics = compare_parameter_payloads(compilable_payload, truth_payload)
    model_gate_passed = model_metrics["exact_field_accuracy"] == 1.0
    condition_dir = (
        root / "runs" / mode / case_id / carrier["carrier"] / f"repeat_{repeat:02d}"
    )
    condition_dir.mkdir(parents=True, exist_ok=True)
    compile_result: dict[str, Any] | None = None
    if model_gate_passed:
        cache_key = f"{case_id}:{_semantic_model_signature(compilable_payload)}"
        compile_result = compile_cache.get(cache_key)
        if compile_result is None:
            predicted_parameter_set = _parameter_set_from_active_payload(
                truth_parameter_set,
                compilable_payload,
                project_id=f"g3e-{mode}-{case_id.lower()}-{carrier['carrier']}-{repeat}",
            )
            compile_result = _compile_model(
                reference_ir,
                predicted_parameter_set,
                condition_dir / "model",
                project_id=predicted_parameter_set.project_id,
            )
            compile_cache[cache_key] = compile_result
    predicted_signature = compile_result["signature"] if compile_result else None
    usage = _usage_delta(usage_before, batch.get("usage") or {})
    capabilities = service.capabilities()
    row = {
        "case_id": case_id,
        "carrier": carrier["carrier"],
        "repeat": repeat,
        "mode": mode,
        "provider": batch.get("provider") or capabilities.get("provider"),
        "model": batch.get("model") or capabilities.get("main_model"),
        "temperature": getattr(service.pipeline.settings, "qwen_temperature", None),
        "prompt_version": batch.get("prompt_hash") or "input_batch_mapping_v1",
        "batch_status": batch["status"],
        "parameter_precision": metrics["precision"],
        "parameter_recall": metrics["recall"],
        "parameter_f1": metrics["f1"],
        "parameter_exact_accuracy": metrics["exact_field_accuracy"],
        "parameter_equivalent": parameter_equivalent,
        "source_location_rate": _source_location_rate(batch),
        "conflict_count": len(batch.get("conflicts") or []),
        "requires_confirmation": bool(batch.get("requires_user_confirmation")),
        "model_compiled": bool(compile_result),
        "model_signature_equivalent": predicted_signature == truth_signature,
        "node_count": (compile_result or {}).get("node_count"),
        "element_count": (compile_result or {}).get("element_count"),
        "ocr_pages": sum(
            int(item.get("ocr_pages") or 0)
            for item in batch.get("files") or []
        ),
        "prompt_tokens": int(usage.get("prompt_tokens") or 0),
        "completion_tokens": int(usage.get("completion_tokens") or 0),
        "total_tokens": int(usage.get("total_tokens") or 0),
        "llm_calls": int(usage.get("calls") or 0),
        "llm_attempts": int(usage.get("attempts") or 0),
        "llm_failures": int(usage.get("failures") or 0),
        "llm_timeouts": int(usage.get("timeouts") or 0),
        "latency_ms": int((time.perf_counter() - started) * 1000),
        "failure_class": _failure_class(batch, parameter_equivalent),
        "parameter_hash": stable_hash(predicted, length=32) if predicted else None,
        "model_signature": predicted_signature,
        "details": {
            "parameter_comparison": metrics,
            "model_input_comparison": model_metrics,
            "compile": compile_result,
        },
    }
    _write_json(condition_dir / "run.json", row)
    return row


def _resolved_payload(resolved: dict[str, Any]) -> dict[str, Any]:
    payload = resolved.get("standard_input_package") or resolved
    payload = _unwrap_fields(payload)
    if "tower_segments" in payload:
        return payload
    sections = payload.get("geometry.tower_sections") or []
    nested: dict[str, Any] = {
        "schema_version": "g3e_active_parameter_payload_v1",
        "structure_family": "wind_tower_shell_with_door",
        "tower_segments": [
            {
                "segment_id": f"SEG_{index + 1:03d}",
                **{key: _coerce_scalar(value) for key, value in section.items()},
            }
            for index, section in enumerate(sections)
        ],
        "door_opening": {},
        "material": {},
        "mesh": {},
        "loads": {},
        "analysis": {},
        "boundary": {},
    }
    prefixes = {
        "geometry.door.": "door_opening",
        "material.": "material",
        "mesh.": "mesh",
        "loads.": "loads",
        "analysis.": "analysis",
        "boundary.": "boundary",
    }
    for key, value in payload.items():
        for prefix, group in prefixes.items():
            if key.startswith(prefix):
                field = key[len(prefix) :]
                nested[group][field] = _coerce_scalar(value)
                break
    if "enabled" in nested["door_opening"]:
        nested["door_opening"]["enabled"] = bool(nested["door_opening"]["enabled"])
    for field in ("static_enabled", "modal_enabled"):
        if field in nested["analysis"]:
            nested["analysis"][field] = bool(nested["analysis"][field])
    enabled_tasks = []
    if nested["analysis"].get("static_enabled") is True:
        enabled_tasks.append("static")
    if nested["analysis"].get("modal_enabled") is True:
        enabled_tasks.append("modal")
    if enabled_tasks:
        nested["analysis_tasks"] = enabled_tasks
    nested["output_requests"] = ["U", "RF", "S", "E", "FREQUENCY"]
    return nested


def _apply_public_compiler_defaults(payload: dict[str, Any]) -> dict[str, Any]:
    completed = json.loads(json.dumps(payload))
    mesh = completed.setdefault("mesh", {})
    mesh.setdefault("element_type", "S4R")
    mesh.setdefault("patch_margin_u_mm", CANONICAL_DEFAULTS["mesh.patch_margin_u_mm"])
    mesh.setdefault("patch_margin_z_mm", CANONICAL_DEFAULTS["mesh.patch_margin_z_mm"])
    mesh.setdefault("hole_layers", CANONICAL_DEFAULTS["mesh.hole_layers"])
    door = completed.setdefault("door_opening", {})
    door.setdefault("theta_center_deg", 0.0)
    analysis = completed.setdefault("analysis", {})
    analysis.setdefault("static_enabled", True)
    analysis.setdefault("modal_enabled", True)
    analysis.setdefault("modal_modes", 20)
    completed.setdefault("boundary", {}).setdefault("name", "BASE_FIXED")
    completed.setdefault("analysis_tasks", ["static", "modal"])
    completed.setdefault("output_requests", ["U", "RF", "S", "E", "FREQUENCY"])
    return completed


def _coerce_scalar(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    text = value.strip()
    if text.lower() == "true":
        return True
    if text.lower() == "false":
        return False
    try:
        return float(text) if any(token in text.lower() for token in (".", "e")) else int(text)
    except ValueError:
        return value


def _semantic_model_signature(payload: dict[str, Any]) -> str:
    return stable_hash(_canonical_numbers(payload), length=64)


def _canonical_numbers(value: Any) -> Any:
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return float(f"{float(value):.12g}")
    if isinstance(value, dict):
        return {key: _canonical_numbers(child) for key, child in value.items()}
    if isinstance(value, list):
        return [_canonical_numbers(child) for child in value]
    return value


def _parameter_set_from_active_payload(
    template: ParameterSet,
    payload: dict[str, Any],
    *,
    project_id: str,
) -> ParameterSet:
    parameter_set = template.model_copy(deep=True)
    parameter_set.project_id = project_id
    values: dict[str, Any] = {}
    for index, section in enumerate(payload["tower_segments"]):
        for field, value in section.items():
            if field != "segment_id":
                values[f"geometry.tower_sections.{index}.{field}"] = value
    for group, prefix in (
        ("door_opening", "geometry.door"),
        ("material", "material"),
        ("mesh", "mesh"),
        ("loads", "loads"),
        ("analysis", "analysis"),
        ("boundary", "boundary"),
    ):
        for field, value in payload[group].items():
            values[f"{prefix}.{field}"] = value
    for key, record in parameter_set.parameters.items():
        record.value = values.get(key)
        record.status = "confirmed" if record.value is not None else "candidate"
        record.sources = []
    parameter_set.analysis_tasks = list(payload["analysis_tasks"])
    parameter_set.output_requests = list(payload["output_requests"])
    validation = validate_parameter_set(parameter_set)
    if validation["status"] != "pass":
        raise ValueError(f"编译候选参数未通过门禁：{validation['issues']}")
    return parameter_set


def _compile_model(
    reference_ir: dict[str, Any],
    parameter_set: ParameterSet,
    root: Path,
    *,
    project_id: str,
) -> dict[str, Any]:
    root.mkdir(parents=True, exist_ok=True)
    ir = apply_parameter_set_to_ir(parameter_set, deepcopy(reference_ir))
    state = ProjectState(
        project_id=project_id,
        root=root,
        excel_path=root / "input.xlsx",
        ir=ir,
        artifacts={},
    )
    report = build_l0(state)
    if report.get("status") != "ready":
        raise RuntimeError(f"确定性模型编译失败：{report.get('errors')}")
    job_hashes = {}
    for job_type, job in sorted(report.get("jobs", {}).items()):
        path = Path(job["path"])
        job_hashes[job_type] = _normalized_inp_sha256(path, project_id)
    mesh_report = report.get("mesh_report") or {}
    node_count = _first_metric(mesh_report, "node_count", "nodes")
    element_count = _first_metric(
        mesh_report, "element_count", "s4r_element_count", "s4r_elements", "elements"
    )
    signature_payload = {
        "job_hashes": job_hashes,
        "node_count": node_count,
        "element_count": element_count,
        "generator": report.get("generator"),
    }
    return {
        "status": "ready",
        "signature": stable_hash(signature_payload, length=64),
        "job_hashes": job_hashes,
        "node_count": node_count,
        "element_count": element_count,
        "generator": report.get("generator"),
    }


def _normalized_inp_sha256(path: Path, project_id: str) -> str:
    return normalized_inp_sha256(path, project_id)


def _first_metric(payload: Any, *keys: str) -> int | None:
    if isinstance(payload, dict):
        for key in keys:
            value = payload.get(key)
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                return int(value)
        for child in payload.values():
            found = _first_metric(child, *keys)
            if found is not None:
                return found
    elif isinstance(payload, list):
        for child in payload:
            found = _first_metric(child, *keys)
            if found is not None:
                return found
    return None


def _unwrap_fields(value: Any) -> Any:
    if isinstance(value, dict):
        if "value" in value and set(value).intersection(
            {"unit", "source_sheet", "source_cell_or_range", "status"}
        ):
            return _unwrap_fields(value["value"])
        return {key: _unwrap_fields(child) for key, child in value.items()}
    if isinstance(value, list):
        return [_unwrap_fields(item) for item in value]
    return value


def _source_location_rate(batch: dict[str, Any]) -> float:
    candidates = batch.get("candidates") or []
    if not candidates:
        return 0.0
    located = sum(
        1
        for item in candidates
        if item.get("evidence_ids") or item.get("source_locations")
    )
    return located / len(candidates)


def _usage_delta(before: dict[str, Any], after: dict[str, Any]) -> dict[str, int]:
    keys = (
        "attempts", "calls", "failures", "timeouts", "prompt_tokens",
        "completion_tokens", "total_tokens", "latency_ms",
    )
    return {
        key: max(0, int(after.get(key) or 0) - int(before.get(key) or 0))
        for key in keys
    }


def _failure_class(batch: dict[str, Any], equivalent: bool) -> str:
    if equivalent:
        return "none"
    status = str(batch.get("status") or "")
    if status in {"failed", "blocked"}:
        return status
    if batch.get("conflicts"):
        return "parameter_conflict"
    if batch.get("requires_user_confirmation"):
        return "confirmation_required"
    return "parameter_mismatch"


def _summarize_results(
    rows: list[dict[str, Any]],
    *,
    mode: str,
    repetitions: int,
) -> dict[str, Any]:
    formats = {}
    for carrier in ENGINEERING_CARRIERS:
        selected = [row for row in rows if row["carrier"] == carrier]
        successes = sum(bool(row["parameter_equivalent"]) for row in selected)
        model_successes = sum(bool(row["model_signature_equivalent"]) for row in selected)
        formats[carrier] = {
            "runs": len(selected),
            "parameter_equivalence": aggregate_binary_metric(successes, len(selected)),
            "model_equivalence": aggregate_binary_metric(model_successes, len(selected)),
            "mean_precision": _mean(selected, "parameter_precision"),
            "mean_recall": _mean(selected, "parameter_recall"),
            "mean_source_location_rate": _mean(selected, "source_location_rate"),
            "mean_latency_ms": _mean(selected, "latency_ms"),
            "stability": stability_metrics(selected),
        }
    return {
        "schema_version": RESULT_SCHEMA_VERSION,
        "experiment_id": "G3-E",
        "mode": mode,
        "repetitions": repetitions,
        "run_count": len(rows),
        "generated_at": _utc_now(),
        "formats": formats,
        "runs": rows,
    }


def _write_result_templates(directory: Path) -> None:
    directory.mkdir(parents=True, exist_ok=True)
    headers = [
        "case_id",
        "carrier",
        "repeat",
        "mode",
        "provider",
        "model",
        "temperature",
        "prompt_version",
        "batch_status",
        "parameter_precision",
        "parameter_recall",
        "parameter_f1",
        "parameter_exact_accuracy",
        "parameter_equivalent",
        "source_location_rate",
        "conflict_count",
        "requires_confirmation",
        "model_compiled",
        "model_signature_equivalent",
        "node_count",
        "element_count",
        "ocr_pages",
        "prompt_tokens",
        "completion_tokens",
        "total_tokens",
        "llm_calls",
        "llm_attempts",
        "llm_failures",
        "llm_timeouts",
        "latency_ms",
        "failure_class",
    ]
    with (directory / "run_results_template.csv").open(
        "w", encoding="utf-8-sig", newline=""
    ) as stream:
        csv.writer(stream).writerow(headers)
    _write_json(
        directory / "summary_template.json",
        {
            "schema_version": RESULT_SCHEMA_VERSION,
            "experiment_id": "G3-E",
            "status": "not_run",
            "formats": {},
        },
    )


def _write_rows_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    if not rows:
        return
    fieldnames = [key for key in rows[0] if key != "details"]
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows:
            writer.writerow({key: row.get(key) for key in fieldnames})


def _write_format_report(path: Path, result: dict[str, Any]) -> None:
    lines = [
        "# G3-E 跨格式解析差异报告",
        "",
        f"- 模式：{result['mode']}",
        f"- 运行数：{result['run_count']}",
        "",
        "| 格式 | 参数等价率 | 模型等价率 | 平均来源定位率 | 平均延迟(ms) |",
        "|---|---:|---:|---:|---:|",
    ]
    for carrier, metrics in result["formats"].items():
        parameter_rate = metrics["parameter_equivalence"]["rate"]
        model_rate = metrics["model_equivalence"]["rate"]
        lines.append(
            f"| {carrier} | {_percent(parameter_rate)} | {_percent(model_rate)} | "
            f"{_percent(metrics['mean_source_location_rate'])} | "
            f"{metrics['mean_latency_ms']:.1f} |"
        )
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def _flatten(value: Any, prefix: str = "") -> Iterable[tuple[str, Any]]:
    if isinstance(value, dict):
        for key, child in value.items():
            path = f"{prefix}.{key}" if prefix else key
            yield from _flatten(child, path)
    elif isinstance(value, list) and value and all(
        isinstance(item, dict) for item in value
    ):
        for index, child in enumerate(value):
            yield from _flatten(child, f"{prefix}[{index}]")
    else:
        yield prefix, value


def _media_type(path: Path) -> str:
    return {
        ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ".csv": "text/csv",
        ".json": "application/json",
        ".yaml": "application/yaml",
        ".txt": "text/plain",
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        ".pdf": "application/pdf",
        ".png": "image/png",
    }.get(path.suffix.lower(), "application/octet-stream")


def _mean(rows: list[dict[str, Any]], key: str) -> float:
    values = [float(row[key]) for row in rows if isinstance(row.get(key), (int, float))]
    return sum(values) / len(values) if values else 0.0


def _percent(value: float | None) -> str:
    return "—" if value is None else f"{value * 100:.1f}%"


def _relative(root: Path, path: Path) -> str:
    return path.relative_to(root).as_posix()


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _write_json(path: Path, payload: Any) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return path


def _utc_now() -> str:
    from datetime import datetime, timezone

    return datetime.now(timezone.utc).isoformat()
