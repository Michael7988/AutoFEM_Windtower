from __future__ import annotations

from pathlib import Path

from windtower_backend.config import EXAMPLE_ROOT

import openpyxl

from windtower_backend.importers.standard_input import normalize_import_to_standard_input, profile_workbook, validate_standard_input
from windtower_backend.importers.windtower_adapter import parse_windtower_workbook
from windtower_backend.modeling import create_project


ROOT = Path(__file__).resolve().parents[1]
TEN_MW = EXAMPLE_ROOT / "synthetic_standard_tower.xlsx"


def test_semantic_workbook_profiles_and_normalizes_to_standard_input() -> None:
    profile = profile_workbook(TEN_MW)

    assert profile["input_format"] == "semantic_standard_workbook"
    assert profile["format_confidence"] >= 0.9

    normalized = normalize_import_to_standard_input(TEN_MW)

    assert normalized["validation_report"]["status"] == "pass"
    assert len(normalized["standard_input_package"]["tower_segments"]) == 10
    assert len(normalized["tower"]) == 11
    assert normalized["parsed_facts"]["door_complete"] is True
    assert normalized["parsed_facts"]["fem_required_keys_missing"] == []
    assert len(normalized["loads_inventory"]["ultimate_cases"]) == 4
    assert normalized["llm_parse_assist"]["total_tokens"] == 0


def test_semantic_workbook_enters_existing_project_flow_without_http_500() -> None:
    response = create_project(TEN_MW.name, TEN_MW.read_bytes(), modeling_goal="静力极限工况和模态分析")

    assert response["input_format"] == "semantic_standard_workbook"
    assert response["validation_report"]["status"] == "pass"
    assert response["summary"]["section_count"] == 10
    assert response["summary"]["tower_height_m"] == 100.0
    assert response["llm_usage"]["status"] == "not_called"


def test_unknown_workbook_returns_diagnostics_instead_of_crashing(tmp_path: Path) -> None:
    path = tmp_path / "unknown.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Random"
    ws.append(["foo", "bar"])
    ws.append([1, 2])
    wb.save(path)

    parsed = parse_windtower_workbook(path)

    assert parsed["detected_format"] == "unknown_llm_assisted"
    assert parsed["confidence_report"]["llm_assist_required"] is True
    assert parsed["validation_report"]["status"] == "blocked"
    assert parsed["llm_parse_assist"]["triggered"] is True


def test_gpt_generated_workbook_profile_is_detected_without_full_workbook_llm_payload(tmp_path: Path) -> None:
    path = tmp_path / "gpt_generated.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "GPT_TowerSegments"
    ws.append(["StandardInputPackage", "towerSegments", "z_bottom_mm", "z_top_mm", "diameter_bottom_mm"])
    ws.append(["S01", "towerSegments", 0, 10000, 6500])
    ws2 = wb.create_sheet("GPT_DoorOpening")
    ws2.append(["doorOpening", "width_mm", "height_mm", "z_bottom_mm"])
    ws2.append(["main", 1200, 2200, 500])
    wb.save(path)

    profile = profile_workbook(path)

    assert profile["input_format"] == "gpt_generated_workbook"
    assert profile["input_family"] == "gpt_generated_workbook"
    assert profile["llm_mapping_contract"]["target"] == "CandidateMappingIR"
    assert profile["llm_mapping_contract"]["canonical_payload"] == "StandardInputPackage"
    assert "numeric_column_stats" in profile["sheets"][0]
    assert "semantic_hints" in profile["sheets"][0]


def test_missing_defaultable_fem_keys_release_as_warning() -> None:
    candidate = {
        "standard_input_package": {
            "tower_segments": [
                {
                    "z_bottom_mm": {"value": 0},
                    "z_top_mm": {"value": 10000},
                    "diameter_bottom_mm": {"value": 6500},
                    "diameter_top_mm": {"value": 6200},
                    "thickness_mm": {"value": 35},
                }
            ],
            "door_opening": {"enabled": {"value": False}},
        },
        "fem_params": {
            "material.elastic_modulus_mpa": 210000.0,
            "material.poisson_ratio": 0.3,
            "section.type": "Shell Section",
            "element.element_type": "S4R",
        },
    }

    report = validate_standard_input(candidate)

    assert report["status"] == "warning"
    assert report["blocked"] is False
    assert report["requires_user_confirmation"] is False
    assert report["defaults_applied"]
    assert {item["severity"] for item in report["missing_items"]} == {"warning"}
