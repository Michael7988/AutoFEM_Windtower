from __future__ import annotations

from pathlib import Path

import openpyxl

from windtower_backend.importers.standard_input import normalize_import_to_standard_input, profile_workbook


def _gpt_style_workbook(tmp_path: Path) -> Path:
    path = tmp_path / "gpt_style_free_headers.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "GPT_TowerSegments"
    ws.append(["StandardInputPackage", "towerSegments", "z_bottom_mm", "z_top_mm", "diameter_bottom_mm", "diameter_top_mm", "thickness_mm"])
    ws.append(["S01", "towerSegments", 0, 12000, 6500, 6200, 38])
    ws.append(["S02", "towerSegments", 12000, 24000, 6200, 5900, 34])
    door = wb.create_sheet("GPT_DoorOpening")
    door.append(["doorOpening", "width_mm", "height_mm", "z_bottom_mm"])
    door.append(["main", 1100, 2200, 500])
    material = wb.create_sheet("GPT_Material")
    material.append(["StandardInputPackage", "elastic_modulus_mpa", "poisson_ratio", "material_name"])
    material.append(["material", 206000, 0.3, "Q345"])
    wb.save(path)
    return path


def test_compact_workbook_profile_exposes_llm_mapping_contract(tmp_path: Path) -> None:
    profile = profile_workbook(_gpt_style_workbook(tmp_path))

    assert profile["input_format"] == "gpt_generated_workbook"
    assert profile["llm_mapping_contract"]["target"] == "CandidateMappingIR"
    assert profile["llm_mapping_contract"]["canonical_payload"] == "StandardInputPackage"
    assert "abaqus_inp" in profile["llm_mapping_contract"]["forbidden_outputs"]
    assert "raw_excel_bytes" not in profile
    assert all("preview" in sheet for sheet in profile["sheets"])


def test_llm_advisor_is_mapping_only_when_unknown_format_needs_help(tmp_path: Path) -> None:
    path = _gpt_style_workbook(tmp_path)

    normalized = normalize_import_to_standard_input(path, enable_llm=True, force_llm=True)
    assist = normalized["llm_parse_assist"]

    assert assist["triggered"] is True
    assert normalized["llm_required"] is True
    assert "llm_standard_input_candidate" not in normalized or normalized.get("llm_standard_input_candidate") is None
    assert normalized["workbook_profile"]["llm_mapping_contract"]["target"] == "CandidateMappingIR"
    assert normalized["candidate_mapping_ir"]["schema_version"] == "candidate_mapping_ir_v1"
    assert normalized["candidate_mapping_ir"]["authority"] == "candidate_only"
    assert set(normalized["workbook_profile"]["llm_mapping_contract"]["forbidden_outputs"]) >= {
        "abaqus_inp",
        "mesh",
        "final_ir",
    }
