from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

from windtower_backend.benchmark.dataset import ALPHAS, DISCLAIMER, _build_anchors, _reference_paths, generate_llm_benchmark
from windtower_backend.importers.candidate_mapping import candidate_mapping_to_standard_input, compare_standard_input_to_ground_truth
from windtower_backend.importers.standard_input import normalize_import_to_standard_input
from windtower_backend.core.pipeline import _resolve_llm_mode
from windtower_backend.l0.implementation import _odb_extraction_script


def test_generate_ten_reproducible_truth_cases(tmp_path: Path) -> None:
    first = generate_llm_benchmark(tmp_path / "one")
    second = generate_llm_benchmark(tmp_path / "two")
    assert first["case_count"] == second["case_count"] == 10
    assert first["alphas"] == second["alphas"] == list(ALPHAS)
    for case_a, case_b in zip(first["cases"], second["cases"]):
        truth_a = _truth(case_a)
        truth_b = _truth(case_b)
        truth_a["benchmark"].pop("workbook", None)
        truth_a["benchmark"].pop("workbook_sha256", None)
        truth_b["benchmark"].pop("workbook", None)
        truth_b["benchmark"].pop("workbook_sha256", None)
        assert truth_a == truth_b
        assert truth_a["benchmark"]["declaration"] == DISCLAIMER


def test_benchmark_invariants_and_disclaimer(tmp_path: Path) -> None:
    manifest = generate_llm_benchmark(tmp_path)
    anchor_a, anchor_b, _ = _build_anchors(_reference_paths())
    for item in manifest["cases"]:
        truth = _truth(item)
        segments = truth["tower_segments"]
        assert len(segments) == 10
        assert segments[0]["z_bottom_mm"] == 0.0
        for left, right in zip(segments, segments[1:]):
            assert left["z_top_mm"] == right["z_bottom_mm"]
        diameters = [segments[0]["diameter_bottom_mm"], *[segment["diameter_top_mm"] for segment in segments]]
        assert all(value > 0 for value in diameters)
        assert all(left >= right for left, right in zip(diameters, diameters[1:]))
        assert all(segment["thickness_mm"] > 0 for segment in segments)
        assert _between(segments[-1]["z_top_mm"], anchor_a["height_mm"], anchor_b["height_mm"])
        for value, left, right in zip(diameters, anchor_a["diameter_mm"], anchor_b["diameter_mm"]):
            assert _between(value, left, right)
        for segment, left, right in zip(segments, anchor_a["thickness_mm"], anchor_b["thickness_mm"]):
            assert _between(segment["thickness_mm"], left, right)
        door = truth["door_opening"]
        assert segments[0]["z_bottom_mm"] <= door["z_bottom_mm"] < segments[0]["z_top_mm"]
        for key in ("width_mm", "height_mm", "z_bottom_mm"):
            assert _between(door[key], anchor_a["door"][key], anchor_b["door"][key])
        fem = truth["fem_settings"]
        assert _between(fem["loads.top_total_force_n"], anchor_a["top_force_n"], anchor_b["top_force_n"])
        assert _between(fem["loads.top_mass_tonne"], anchor_a["top_mass_tonne"], anchor_b["top_mass_tonne"])
        assert _between(fem["loads.wind_pressure_mpa"], anchor_a["wind_pressure_mpa"], anchor_b["wind_pressure_mpa"])
        workbook = load_workbook(item["input"], data_only=True)
        assert any(DISCLAIMER in str(cell.value) for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row if cell.value)


def test_s01_rule_only_is_unknown_and_blocked(tmp_path: Path) -> None:
    manifest = generate_llm_benchmark(tmp_path)
    input_path = Path(manifest["cases"][0]["input"])
    result = normalize_import_to_standard_input(input_path, enable_llm=False)
    assert result["input_format"] == "unknown_llm_assisted"
    assert result["parse_strategy"] == "llm_assisted"
    assert result["validation_report"]["blocked"] is True


def test_candidate_mapping_evidence_conversion_and_truth_gate(tmp_path: Path) -> None:
    manifest = generate_llm_benchmark(tmp_path)
    case = manifest["cases"][0]
    input_path = Path(case["input"])
    truth = _truth(case)
    raw_segments = []
    for segment in truth["tower_segments"]:
        raw_segments.append([
            segment["segment_id"],
            segment["z_bottom_mm"] / 1000.0,
            segment["z_top_mm"] / 1000.0,
            segment["diameter_bottom_mm"] / 10.0,
            segment["diameter_top_mm"] / 10.0,
            segment["thickness_mm"],
        ])
    door = truth["door_opening"]
    material = truth["materials"][0]
    fem = truth["fem_settings"]
    fields = [_field("domain.wind_tower.tower_segments", truth["tower_segments"], "轮廓清单_Profile", "A4:F13", raw_segments)]
    fields.append(_field(
        "domain.wind_tower.door_opening",
        door,
        "设计意图_DesignIntent",
        "A9:C13",
        ["yes", door["shape"], door["width_mm"] / 10.0, door["height_mm"] / 1000.0, door["z_bottom_mm"] / 1000.0],
    ))
    fields.append(_field(
        "generic.materials[0]",
        material,
        "设计意图_DesignIntent",
        "A4:C8",
        [
            material["name"],
            material["elastic_modulus_mpa"],
            material["poisson_ratio"],
            material["density_tonne_mm3"],
            material["yield_strength_mpa"],
        ],
    ))
    fields.append(_field(
        "generic.fem_settings",
        fem,
        "设计意图_DesignIntent",
        "A4:C38",
        [
            material["name"],
            fem["element.element_type"],
            fem["section.type"],
            fem["section.thickness_source"],
            fem["mesh.strategy"],
            fem["sets_surfaces.base_set"],
            fem["sets_surfaces.top_set"],
            fem["sets_surfaces.door_set"],
            fem["sets_surfaces.wind_surface"],
            fem["mesh.global_size_mm"],
            fem["mesh.door_size_mm"],
            fem["loads.top_total_force_n"] / 1000.0,
            fem["loads.top_mass_tonne"],
            fem["loads.wind_pressure_mpa"],
            fem["steps.STEP-1-TOP-REACTION"],
            fem["steps.STEP-2-WIND-PRESSURE"],
            fem["steps.nlgeom"],
            fem["modal.modes"],
        ],
    ))
    mapping = {
        "schema_version": "candidate_mapping_ir_v1",
        "model_family": "wind_tower",
        "source_file": input_path.name,
        "fields": fields,
        "unresolved": [],
        "evidence_summary": {"field_count": len(fields)},
        "issues": [],
    }
    converted = candidate_mapping_to_standard_input(mapping, input_path, require_benchmark_fields=True)
    assert converted["promotion_status"] == "candidate_validated"
    comparison = compare_standard_input_to_ground_truth(converted, truth)
    assert comparison["status"] == "pass"
    assert comparison["accuracy"] == 1.0


def test_candidate_mapping_rejects_forbidden_ir_and_false_evidence(tmp_path: Path) -> None:
    manifest = generate_llm_benchmark(tmp_path)
    input_path = Path(manifest["cases"][0]["input"])
    mapping = {
        "schema_version": "candidate_mapping_ir_v1",
        "model_family": "wind_tower",
        "fields": [_field("domain.wind_tower.tower_segments[0]", {}, "轮廓清单_Profile", "A4:F4", [999999])],
        "final_ir": {"forbidden": True},
    }
    result = candidate_mapping_to_standard_input(mapping, input_path)
    assert result["promotion_status"] == "blocked"
    paths = {item["path"] for item in result["candidate_validation"]["issues"]}
    assert "forbidden_output" in paths
    assert "domain.wind_tower.tower_segments[0]" in paths


def test_llm_mode_boolean_compatibility_and_modal_extraction_scope() -> None:
    assert _resolve_llm_mode(None, enable_llm=False, force_llm=False) == "off"
    assert _resolve_llm_mode(None, enable_llm=True, force_llm=False) == "adaptive"
    assert _resolve_llm_mode(None, enable_llm=True, force_llm=True) == "forced_parse"
    assert _resolve_llm_mode("forced_parse", enable_llm=False, force_llm=False) == "forced_parse"
    script = _odb_extraction_script()
    assert 'write_csv("modal_frequencies.csv", rows,' in script
    assert "rows[:6]" not in script


def _field(path: str, value: object, sheet: str, cell_range: str, raw: object) -> dict[str, object]:
    return {
        "path": path,
        "value": raw,
        "unit": "mixed",
        "status": "observed",
        "evidence": {"sheet": sheet, "cell_range": cell_range, "raw_value": raw},
        "normalization": {"normalized_value": value, "normalized_unit": "mm-N-MPa-tonne"},
        "confidence": 1.0,
        "criticality": "high",
        "mapping_reason": "Directly mapped from the cited synthetic workbook cells.",
    }


def _truth(case: dict[str, object]) -> dict[str, object]:
    return json.loads(Path(str(case["ground_truth"])).read_text(encoding="utf-8"))


def _between(value: float, left: float, right: float) -> bool:
    return min(left, right) - 1e-9 <= value <= max(left, right) + 1e-9
