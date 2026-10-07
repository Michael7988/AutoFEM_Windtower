from __future__ import annotations

from pathlib import Path

from windtower_backend.config import EXAMPLE_ROOT
from typing import Callable

import openpyxl

from windtower_backend.importers.standard_input import normalize_import_to_standard_input


ROOT = Path(__file__).resolve().parents[1]
TEN_MW = EXAMPLE_ROOT / "synthetic_standard_tower.xlsx"


def _variant(tmp_path: Path, name: str, mutate: Callable[[openpyxl.Workbook], None]) -> Path:
    path = tmp_path / name
    wb = openpyxl.load_workbook(TEN_MW)
    mutate(wb)
    wb.save(path)
    return path


def _default_keys(report: dict) -> set[str]:
    return {str(item.get("field") or item.get("key") or item.get("path")) for item in report.get("defaults_applied", [])}


def _issue_paths(report: dict) -> set[str]:
    return {str(item.get("path") or item.get("key") or item.get("item")) for item in report.get("issues", []) + report.get("missing_items", [])}


def test_missing_material_property_warns_and_applies_default(tmp_path: Path) -> None:
    def mutate(wb: openpyxl.Workbook) -> None:
        for row in range(3, 5):
            wb["04_Materials"][f"C{row}"] = None

    path = _variant(tmp_path, "missing_material_e.xlsx", mutate)

    normalized = normalize_import_to_standard_input(path)
    report = normalized["validation_report"]

    assert report["status"] == "warning"
    assert report["blocked"] is False
    assert "material.elastic_modulus_mpa" in _default_keys(report)
    assert normalized["input_quality_scores"]["readiness"] == "warning"


def test_missing_recommended_mesh_size_warns_but_modeling_can_continue(tmp_path: Path) -> None:
    path = _variant(tmp_path, "missing_mesh_size.xlsx", lambda wb: setattr(wb["05_FEM"]["D8"], "value", None))

    normalized = normalize_import_to_standard_input(path)
    report = normalized["validation_report"]

    assert report["status"] == "warning"
    assert report["blocked"] is False
    assert "mesh.global_size_mm" in _default_keys(report)
    assert normalized["input_quality_scores"]["fem_completeness"] < 1.0


def test_disabled_door_allows_missing_door_geometry(tmp_path: Path) -> None:
    def mutate(wb: openpyxl.Workbook) -> None:
        ws = wb["03_Door"]
        ws["D4"] = False
        for cell in ("D8", "D9", "D10"):
            ws[cell] = None

    path = _variant(tmp_path, "door_disabled_missing_geometry.xlsx", mutate)

    normalized = normalize_import_to_standard_input(path)
    report = normalized["validation_report"]

    assert report["blocked"] is False
    assert not any(path.startswith("door_opening.") for path in _issue_paths(report))


def test_missing_tower_thickness_blocks_real_mesh(tmp_path: Path) -> None:
    def mutate(wb: openpyxl.Workbook) -> None:
        for row in range(5, 15):
            wb["02_Segments"][f"G{row}"] = None

    path = _variant(tmp_path, "missing_tower_thickness.xlsx", mutate)

    normalized = normalize_import_to_standard_input(path)
    report = normalized["validation_report"]

    assert report["blocked"] is True
    assert report["status"] == "blocked"
    assert any(path.startswith("tower_segments") for path in _issue_paths(report))


def test_unknown_workbook_enters_llm_assisted_blocked_state_without_crashing(tmp_path: Path) -> None:
    path = tmp_path / "unknown.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Random Inputs"
    ws.append(["foo", "bar", "baz"])
    ws.append([1, 2, 3])
    wb.save(path)

    normalized = normalize_import_to_standard_input(path, enable_llm=True, force_llm=True)

    assert normalized["input_format"] == "unknown_llm_assisted"
    assert normalized["validation_report"]["blocked"] is True
    assert normalized["llm_required"] is True
    assert normalized["input_quality_scores"]["llm_required"] is True
    assert normalized["input_quality_scores"]["readiness"] == "blocked"
