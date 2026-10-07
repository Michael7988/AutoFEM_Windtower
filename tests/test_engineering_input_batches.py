from __future__ import annotations

import io
import json
import time
from pathlib import Path

import fitz
from PIL import Image
from docx import Document
from pptx import Presentation

from windtower_backend.config import load_settings
from windtower_backend.core.pipeline import WindTowerPipeline
from windtower_backend.input_batches.extractors import extract_file
from windtower_backend.input_batches.models import (
    CompleteParametersRequest,
    ConfirmBatchReviewRequest,
    ConfirmFemSpecificationRequest,
    ConfirmIntentRequest,
    ConfirmTemplateRequest,
    ReparseRequest,
    ResolveConflictsRequest,
)
from windtower_backend.input_batches.providers import OpenAIEngineeringProvider, QwenEngineeringProvider, create_engineering_provider
from windtower_backend.input_batches.service import EngineeringInputBatchService


def _wait(service: EngineeringInputBatchService, batch_id: str) -> dict:
    for _ in range(100):
        batch = service.get_batch(batch_id)
        if batch["status"] not in {"queued", "processing"}:
            return batch
        time.sleep(0.05)
    raise AssertionError("batch did not finish")


def _confirm_quick_intent(
    service: EngineeringInputBatchService,
    batch_id: str,
    *,
    tasks: list[str],
    features: list[str] | None = None,
) -> dict:
    service.confirm_template(batch_id, ConfirmTemplateRequest(
        template_id="wind_tower_quick_v1",
        version="1.0.0",
    ))
    service.confirm_intent(batch_id, ConfirmIntentRequest(
        model_family="wind_tower",
        fidelity="quick",
        analysis_tasks=tasks,
        requested_features=features or ["tower_shell"],
        output_targets=["displacement"],
    ))
    return _wait(service, batch_id)


def _confirm_fem_and_build_review(
    service: EngineeringInputBatchService,
    batch_id: str,
) -> ConfirmBatchReviewRequest:
    batch = service.get_batch(batch_id)
    if batch["fem_review"]["status"] != "confirmed":
        batch = service.confirm_fem_specification(
            batch_id,
            ConfirmFemSpecificationRequest(
                review_revision=batch["fem_review_revision"],
                parameters=[],
            ),
        )
    return ConfirmBatchReviewRequest(
        review_revision=batch["review_revision"],
        parameters=[
            {"key": item["key"], "value": item["value"], "unit": item["unit"]}
            for item in batch["review_parameters"]
        ],
    )


def test_local_structured_and_document_extractors_keep_source_locations() -> None:
    json_fragments, _ = extract_file("f1", "parameters.json", b'{"door_opening":{"height_mm":2400}}')
    assert any(item.location.json_path == "$.door_opening.height_mm" and item.text == "2400" for item in json_fragments)

    document = Document()
    document.add_paragraph("Door height 2400 mm")
    document.add_table(rows=1, cols=2).rows[0].cells[0].text = "mesh"
    docx_bytes = io.BytesIO()
    document.save(docx_bytes)
    docx_fragments, _ = extract_file("f2", "notes.docx", docx_bytes.getvalue())
    assert any(item.extraction_method == "local_docx" for item in docx_fragments)
    assert any(item.kind == "table" for item in docx_fragments)

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Wind pressure 0.0012 MPa"
    pptx_bytes = io.BytesIO()
    presentation.save(pptx_bytes)
    pptx_fragments, _ = extract_file("f3", "brief.pptx", pptx_bytes.getvalue())
    assert pptx_fragments[0].location.slide == 1

    yaml_fragments, _ = extract_file("f4", "parameters.yaml", b"mesh:\n  global_size_mm: 450\n")
    assert any(item.location.json_path == "$.mesh.global_size_mm" and item.text == "450" for item in yaml_fragments)


def test_pdf_uses_text_layer_and_only_ocrs_scanned_pages() -> None:
    document = fitz.open()
    text_page = document.new_page()
    text_page.insert_text((72, 72), "Wind tower engineering parameters " * 4)
    document.new_page()
    pdf = document.tobytes()
    calls = []

    def fake_ocr(content: bytes, media_type: str, prompt: str) -> dict:
        calls.append((len(content), media_type, prompt))
        return {"status": "called", "model": "fake-ocr", "text": "Door height 2300 mm", "confidence": 0.9}

    fragments, meta = extract_file("f1", "mixed.pdf", pdf, fake_ocr)
    assert meta["ocr_pages"] == 1
    assert len(calls) == 1
    assert fragments[0].extraction_method == "local_pdf_text"
    assert fragments[1].from_ocr is True
    assert fragments[1].location.page == 2


def test_multi_file_batch_detects_conflicts_and_reuses_extraction_cache(tmp_path: Path) -> None:
    pipeline = WindTowerPipeline(load_settings())
    service = EngineeringInputBatchService(pipeline)
    service.root = tmp_path / "batches"
    service.cache_root = service.root / "_cache"
    service.cache_root.mkdir(parents=True)

    first = service.create_batch([
        ("a.txt", "text/plain", "门洞高度 2300 mm".encode("utf-8")),
        ("b.json", "application/json", b'{"door_opening":{"height_mm":2400}}'),
    ])
    batch = _wait(service, first["batch_id"])
    conflict = next(item for item in batch["conflicts"] if item["key"] == "geometry.door.height_mm")
    assert batch["schema_version"] == "engineering_input_batch_v3"
    assert batch["selected_template_id"] == "wind_tower_quick_v1"
    assert batch["status"] == "blocked"
    assert conflict["blocking"] is False
    assert len(conflict["candidate_ids"]) == 2

    selected = conflict["candidate_ids"][0]
    resolved = service.resolve_conflicts(
        first["batch_id"],
        ResolveConflictsRequest(resolutions=[{"conflict_id": conflict["conflict_id"], "candidate_id": selected}]),
    )
    assert resolved["status"] == "blocked"

    second = service.create_batch([("again.txt", "text/plain", "门洞高度 2300 mm".encode("utf-8"))])
    cached = _wait(service, second["batch_id"])
    assert cached["files"][0]["cache_hits"] == 1


def test_capabilities_are_provider_neutral_and_limits_are_exposed() -> None:
    service = EngineeringInputBatchService(WindTowerPipeline(load_settings()))
    capabilities = service.capabilities()
    assert capabilities["schema_version"] == "engineering_llm_capabilities_v1"
    assert ".pdf" in capabilities["accepted_extensions"]
    assert ".json" in capabilities["accepted_extensions"]
    assert capabilities["limits"]["max_files"] == 20


def test_unavailable_ocr_fails_only_the_image_and_other_files_can_continue(tmp_path: Path) -> None:
    pipeline = WindTowerPipeline(load_settings())
    service = EngineeringInputBatchService(pipeline)
    service.root = tmp_path / "batches"
    service.cache_root = service.root / "_cache"
    service.cache_root.mkdir(parents=True)
    service.provider.extract_image_text = lambda *_args: {
        "status": "not_called", "provider": "none", "model": "none", "text": "", "confidence": 0,
    }
    image = Image.new("RGB", (32, 32), "white")
    image_bytes = io.BytesIO()
    image.save(image_bytes, format="PNG")

    created = service.create_batch([
        ("notes.txt", "text/plain", "门洞高度 2300 mm".encode("utf-8")),
        ("drawing.png", "image/png", image_bytes.getvalue()),
    ])
    batch = _wait(service, created["batch_id"])
    statuses = {item["filename"]: item["status"] for item in batch["files"]}
    assert statuses["notes.txt"] == "completed"
    assert statuses["drawing.png"] == "failed"
    assert batch["selected_template_id"] == "wind_tower_quick_v1"
    assert batch["status"] == "blocked"
    drawing = next(item for item in batch["files"] if item["filename"] == "drawing.png")
    assert drawing["error"]


def test_openai_provider_can_be_selected_without_changing_contract(monkeypatch) -> None:
    monkeypatch.setenv("ENGINEERING_LLM_PROVIDER", "openai")
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    provider = create_engineering_provider(load_settings())
    assert isinstance(provider, OpenAIEngineeringProvider)
    capabilities = provider.capabilities()
    assert capabilities["provider"] == "openai"
    assert capabilities["schema_version"] == "engineering_llm_capabilities_v1"


def test_complete_standard_input_json_can_create_wind_tower_without_excel(tmp_path: Path) -> None:
    pipeline = WindTowerPipeline(load_settings())
    service = EngineeringInputBatchService(pipeline)
    service.root = tmp_path / "batches"
    service.cache_root = service.root / "_cache"
    service.cache_root.mkdir(parents=True)
    package = {
        "schema_version": "standard_input_package_v1",
        "source_file": "tower.json",
        "input_format": "standard_input_json",
        "tower_segments": [{
            "z_bottom_mm": 0,
            "z_top_mm": 30000,
            "diameter_bottom_mm": 4500,
            "diameter_top_mm": 3200,
            "thickness_mm": 28,
        }],
        "door_opening": {"enabled": False},
        "materials": [{
            "name": "Q345",
            "elastic_modulus_mpa": 206000,
            "poisson_ratio": 0.3,
            "density_tonne_mm3": 7.85e-9,
            "yield_strength_mpa": 345,
        }],
        "fem_settings": {
            "material.name": "Q345",
            "material.elastic_modulus_mpa": 206000,
            "material.poisson_ratio": 0.3,
            "section.type": "Shell Section",
            "element.element_type": "S4R",
            "mesh.strategy": "structured_quad_shell",
            "mesh.global_size_mm": 650,
            "loads.top_mass_tonne": 0,
            "loads.top_total_force_n": 0,
            "loads.top_force_direction": "+X",
            "loads.wind_pressure_mpa": 0.0012,
            "analysis.static_enabled": True,
            "analysis.modal_enabled": False,
            "boundary.name": "BASE_FIXED",
        },
    }
    created = service.create_batch(
        [("tower.json", "application/json", json.dumps({"standard_input_package": package}).encode())],
        modeling_goal="建立风塔并进行静力分析",
    )
    batch = _wait(service, created["batch_id"])
    assert batch["status"] == "ready"
    assert batch["selected_template_id"] == "wind_tower_quick_v1"
    assert batch["review_parameters"]
    review = ConfirmBatchReviewRequest(
        review_revision=batch["review_revision"],
        parameters=[
            {"key": item["key"], "value": item["value"], "unit": item["unit"]}
            for item in batch["review_parameters"]
        ],
    )
    engineering_reviewed = service.confirm_engineering_review(created["batch_id"], review)
    assert engineering_reviewed["engineering_review"]["status"] == "confirmed"
    assert engineering_reviewed["fem_review"]["status"] == "pending"
    assert engineering_reviewed["merge_readiness"]["ready"] is False
    fem_reviewed = service.confirm_fem_specification(
        created["batch_id"],
        ConfirmFemSpecificationRequest(
            review_revision=engineering_reviewed["fem_review_revision"],
            parameters=[],
        ),
    )
    assert fem_reviewed["merge_readiness"]["ready"] is True
    project = service.confirm(created["batch_id"])
    assert project["project_id"].startswith("wt-")
    assert project["summary"]["section_count"] == 1
    assert project["input_batch"]["status"] == "confirmed"
    assert project["engineering_review"]["status"] == "confirmed"


def test_deterministic_mode_never_calls_parameter_llm(tmp_path: Path) -> None:
    service = EngineeringInputBatchService(WindTowerPipeline(load_settings()))
    service.root = tmp_path / "batches"
    service.cache_root = service.root / "_cache"
    service.cache_root.mkdir(parents=True)
    calls = []
    service.provider.map_parameters = lambda *_args: calls.append(True) or {}
    created = service.create_batch([("notes.txt", "text/plain", "门洞高度 2300 mm".encode())], parse_mode="deterministic")
    batch = _wait(service, created["batch_id"])
    assert calls == []
    assert batch["parse_mode"] == "deterministic"
    assert batch["provider"] == "deterministic"


def test_llm_mode_requires_configured_provider(tmp_path: Path) -> None:
    service = EngineeringInputBatchService(WindTowerPipeline(load_settings()))
    service.root = tmp_path / "batches"
    service.cache_root = service.root / "_cache"
    service.cache_root.mkdir(parents=True)
    service.provider.capabilities = lambda: {"configured": False, "provider": "qwen", "main_model": "qwen3.7-max"}
    try:
        service.create_batch([("notes.txt", "text/plain", b"data")], parse_mode="llm")
    except RuntimeError as exc:
        assert "QWEN_API_KEY" in str(exc)
    else:
        raise AssertionError("LLM mode must reject an unconfigured provider")


def test_llm_full_discards_local_candidates_and_calls_both_advisors(tmp_path: Path) -> None:
    service = EngineeringInputBatchService(WindTowerPipeline(load_settings()))
    service.root = tmp_path / "batches"
    service.cache_root = service.root / "_cache"
    service.cache_root.mkdir(parents=True)
    service.provider.capabilities = lambda: {
        "configured": True, "provider": "qwen", "main_model": "qwen3.8-max"
    }
    calls: list[tuple[str, tuple[str, ...]]] = []

    def mapped(role: str):
        def invoke(_evidence, _goal, _prompt, target_keys):
            calls.append((role, tuple(target_keys)))
            return {"status": "called", "provider": "qwen", "model": "qwen3.8-max", "fields": [], "request_count": 1}
        return invoke

    service.engineering_advisor.map_parameters = mapped("engineering")
    service.fem_advisor.map_parameters = mapped("fem")
    created = service.create_batch(
        [("notes.txt", "text/plain", "门洞高度 2300 mm".encode())],
        parse_mode="llm_full",
    )
    batch = _wait(service, created["batch_id"])
    assert batch["parse_mode"] == "llm_full"
    assert {role for role, _ in calls} == {"engineering", "fem"}
    assert all(item["extraction_method"] != "deterministic_text" for item in batch["candidates"])
    assert batch["candidate_mapping_ir"]["target_keys"]


def test_llm_batch_combines_evidence_and_creates_project_without_excel(tmp_path: Path) -> None:
    service = EngineeringInputBatchService(WindTowerPipeline(load_settings()))
    service.root = tmp_path / "batches"
    service.cache_root = service.root / "_cache"
    service.cache_root.mkdir(parents=True)
    service.provider.capabilities = lambda: {
        "configured": True, "provider": "qwen", "main_model": "qwen3.7-max",
        "ocr_model": "qwen3.5-ocr", "vision_model": "qwen3.7-max-2026-06-08",
    }
    service.provider.usage = lambda: {"calls": 1, "total_tokens": 100}

    def fake_mapping(evidence: list[dict], _goal: str | None, *_args, **_kwargs) -> dict:
        evidence_id = evidence[0]["evidence_id"]
        common = {"status": "observed", "confidence": 0.95, "criticality": "high", "mapping_reason": "test", "evidence_ids": [evidence_id]}
        return {
            "status": "called", "provider": "qwen", "model": "qwen3.7-max", "schema_version": "candidate_mapping_ir_v1",
            "fields": [
                {**common, "path": "domain.wind_tower.tower_segments", "value": [{"z_bottom_mm": 0, "z_top_mm": 30000, "diameter_bottom_mm": 4500, "diameter_top_mm": 3200, "thickness_mm": 28}]},
                {**common, "path": "domain.wind_tower.door_opening", "value": {"enabled": False}},
                {**common, "path": "generic.materials[0]", "value": {"name": "Q345", "elastic_modulus_mpa": 206000, "poisson_ratio": 0.3, "density_tonne_mm3": 7.85e-9, "yield_strength_mpa": 345}},
                {**common, "path": "generic.fem_settings", "value": {
                    "mesh.global_size_mm": 650, "loads.top_mass_tonne": 0, "loads.top_total_force_n": 0,
                    "loads.top_force_direction": "+X", "loads.wind_pressure_mpa": 0.0012,
                    "analysis.static_enabled": True, "analysis.modal_enabled": False, "boundary.name": "BASE_FIXED",
                }},
            ],
        }

    service.provider.map_parameters = fake_mapping
    created = service.create_batch([
        ("geometry.txt", "text/plain", "塔高30m，底部直径4.5m，顶部直径3.2m，厚度28mm。".encode()),
        ("loads.txt", "text/plain", "材料Q345，风压0.0012MPa。".encode()),
    ], parse_mode="llm")
    _wait(service, created["batch_id"])
    batch = _confirm_quick_intent(service, created["batch_id"], tasks=["static"])
    assert batch["parse_mode"] == "llm"
    assert batch["candidate_mapping_ir"]["schema_version"] == "candidate_mapping_ir_v2"
    assert batch["readiness"]["preview_allowed"] is True
    resolutions = [
        {"conflict_id": conflict["conflict_id"], "candidate_id": conflict["candidate_ids"][0]}
        for conflict in batch["conflicts"]
    ]
    resolved = service.resolve_conflicts(created["batch_id"], ResolveConflictsRequest(resolutions=resolutions))
    assert resolved["readiness"]["preview_allowed"] is True
    review = _confirm_fem_and_build_review(service, created["batch_id"])
    project = service.confirm(created["batch_id"], review)
    assert project["summary"]["section_count"] == 1
    assert project["input_batch"]["provider"] == "qwen"


def test_qwen_provider_uses_responses_api_and_returns_candidate_mapping(monkeypatch) -> None:
    monkeypatch.setenv("QWEN_PROVIDER_PROFILE", "legacy")
    monkeypatch.setenv("QWEN_API_KEY", "test-key")
    monkeypatch.setenv("QWEN_MODEL", "qwen3.7-max")
    monkeypatch.setenv("QWEN_BASE_URL", "https://example.invalid/compatible-mode/v1")
    calls = []

    class FakeResponse:
        def raise_for_status(self) -> None:
            return None

        def json(self) -> dict:
            return {
                "output": [{"content": [{"type": "output_text", "text": json.dumps({
                    "schema_version": "candidate_mapping_ir_v1",
                    "fields": [{"path": "domain.wind_tower.tower_segments", "value": [], "evidence_ids": ["ev-1"]}],
                })}]}],
                "usage": {"input_tokens": 10, "output_tokens": 5, "total_tokens": 15},
            }

    def fake_post(url: str, **kwargs):
        calls.append((url, kwargs))
        return FakeResponse()

    monkeypatch.setattr("windtower_backend.input_batches.providers.requests.post", fake_post)
    provider = QwenEngineeringProvider(load_settings())
    result = provider.map_parameters([{
        "evidence_id": "ev-1", "filename": "tower.txt", "location": {"line": 1}, "text": "tower segment",
    }], None)
    assert calls[0][0] == "https://example.invalid/compatible-mode/v1/responses"
    assert calls[0][1]["json"]["model"] == "qwen3.7-max"
    assert result["schema_version"] == "candidate_mapping_ir_v1"
    assert result["fields"][0]["path"] == "domain.wind_tower.tower_segments"


def test_manual_completion_append_and_reparse_reuse_the_batch(tmp_path: Path) -> None:
    service = EngineeringInputBatchService(WindTowerPipeline(load_settings()))
    service.root = tmp_path / "batches"
    service.cache_root = service.root / "_cache"
    service.cache_root.mkdir(parents=True)
    created = service.create_batch([("notes.txt", "text/plain", b"engineering notes")], parse_mode="deterministic")
    initial = _wait(service, created["batch_id"])
    assert "mesh.global_size_mm" not in initial["resolved_parameters"]
    completed = service.complete_parameters(created["batch_id"], CompleteParametersRequest(
        changes=[{"key": "mesh.global_size_mm", "value": 650, "unit": "mm"}],
    ))
    assert completed["resolved_parameters"]["mesh.global_size_mm"] == 650
    appended = service.append_files(created["batch_id"], [
        ("loads.json", "application/json", b'{"loads":{"wind_pressure_mpa":0.0012}}'),
    ])
    after_append = _wait(service, appended["batch_id"])
    assert len(after_append["files"]) == 2
    reparsed = service.reparse(created["batch_id"], ReparseRequest(parse_mode="deterministic"))
    after_reparse = _wait(service, reparsed["batch_id"])
    assert after_reparse["parse_mode"] == "deterministic"
    assert after_reparse["resolved_parameters"]["mesh.global_size_mm"] == 650
